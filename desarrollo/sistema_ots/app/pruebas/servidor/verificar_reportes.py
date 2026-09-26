"""verificar_reportes.py — Reportes por zona con exportación, y el cronograma que escribe (T2.14.5).

Contra el sitio de pruebas, como administración y como jefe de UIO:

  - reportes.php?zona=UIO con la administración → la cifra de casos coincide con la del jefe de UIO;
  - reporte_exportar.php → xlsx (Excel, ≥ 4 hojas), pdf (%PDF) y pptx (≥ 6 diapositivas), y EXPORTAR_REPORTE;
  - cronograma_importar_cli.php llena ingresos_preventivos; cronograma.php devuelve fuente «tabla», csrf y módulos;
  - cronograma_accion.php: el jefe de UIO no toca un ingreso de CNLJ (403); sobre uno de UIO reagenda (200 y
    fila REAGENDA), actualiza el kit, agenda un local nuevo, lo cierra (CUMPLIDO) y anota; sin X-Csrf → 403.

Lo que la prueba escribe sobre ingresos reales se devuelve a su estado; los ingresos sintéticos (año 2099) se borran.

Se corre desde el PC con la llave del servidor:
    INDUSTEC_LLAVE_SSH=... INDUSTEC_SSH_USER=... PYTHONUTF8=1 python verificar_reportes.py
"""
import io
import json
import re
import sys
import urllib.error
import urllib.request
import zipfile

import verificar_http as vh

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def cabeceras(sesion, ruta):
    req = urllib.request.Request(vh.BASE + ruta, headers={"User-Agent": "verificar_reportes/1.0"})
    try:
        r = sesion.op.open(req, timeout=180)
        return r.status, {k.lower(): v for k, v in r.headers.items()}, r.read()
    except urllib.error.HTTPError as e:
        return e.code, {k.lower(): v for k, v in e.headers.items()}, e.read()


def post_json_sin_csrf(sesion, ruta, cuerpo):
    req = urllib.request.Request(vh.BASE + ruta, data=json.dumps(cuerpo).encode("utf-8"),
                                 headers={"Content-Type": "application/json", "User-Agent": "verificar_reportes/1.0"})
    try:
        r = sesion.op.open(req, timeout=60)
        return r.status, r.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace")


def casos_de(html):
    m = re.search(r'data-casos="(\d+)"', html)
    return int(m.group(1)) if m else None


def main():
    claves = json.loads(vh.ssh("cat ~/respaldos/claves_prueba.json"))
    ids = claves["ids"]
    s = {u: vh.Sesion(u) for u in ("jefe_prueba_uio", "admin_prueba")}
    for u, se in s.items():
        st, loc = se.entrar(claves["claves"][u])
        vh.anotar("ingreso", f"{u} entra", st == 302 and "login.php" not in loc, f"{st} → {loc}")
    jefe, adm = s["jefe_prueba_uio"], s["admin_prueba"]

    print("\n== 1. el tablero por zona y por mes ==")
    st, _, c_adm = adm.pedir("reportes.php?zona=UIO")
    st2, _, c_jefe = jefe.pedir("reportes.php")
    vh.anotar("T2.14.5", "reportes.php?zona=UIO (administración) y reportes.php (jefe UIO) → 200", st == 200 and st2 == 200, f"{st} · {st2}")
    vh.anotar("T2.14.5", "la cifra de casos de UIO coincide entre la administración y el jefe", casos_de(c_adm) is not None and casos_de(c_adm) == casos_de(c_jefe), f"{casos_de(c_adm)} · {casos_de(c_jefe)}")
    vh.anotar("T2.14.5", "hay sección de rendimiento por técnico y de cumplimiento del preventivo", "Rendimiento por técnico" in c_adm and "Cumplimiento del preventivo" in c_adm, "")
    # Desde el vocabulario único (24-sep-2026) el botón de CNLJ se rotula «CUENCA-LOJA».
    vh.anotar("T2.14.5", "hay filtros rápidos de zona y selector de mes para la administración", 'name="mes"' in c_adm and ">CUENCA-LOJA<" in c_adm, "")
    vh.anotar("T2.14.5", "el jefe no ve el filtro de otras zonas", ">CUENCA-LOJA<" not in c_jefe.split("Tablero de servicio")[1][:1500], "")
    st, _, c_mes = adm.pedir("reportes.php?mes=2026-09")
    vh.anotar("T2.14.5", "reportes.php?mes=AAAA-MM → 200 y lo dice", st == 200 and "septiembre de 2026" in c_mes, st)
    st, _, c_todas = adm.pedir("reportes.php")
    vh.anotar("T2.14.5", "el literal «7.069» ya no está escrito a mano", "7.069" not in c_todas, "")

    print("\n== 2. la exportación en tres formatos ==")
    e0 = int(vh.sql("SELECT COUNT(*) n FROM bitacora WHERE accion = 'EXPORTAR_REPORTE' AND usuario_id = ?", [ids["admin_prueba"]])[0]["n"])
    st, cab, cuerpo = cabeceras(adm, "reporte_exportar.php?formato=xlsx&zona=UIO")
    hojas = 0
    try:
        with zipfile.ZipFile(io.BytesIO(cuerpo)) as z:
            hojas = len([n for n in z.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml", n)])
    except Exception:
        hojas = 0
    vh.anotar("T2.14.5", "xlsx → 200, tipo Excel y ≥ 4 hojas", st == 200 and "spreadsheetml" in cab.get("content-type", "") and hojas >= 4, f"{st} · {cab.get('content-type', '')[:40]} · hojas={hojas}")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(cuerpo), read_only=True)
        nombres = wb.sheetnames
        # La hoja de las órdenes abiertas se llama con el término del TOTAL (vocabulario.json).
        vh.anotar("T2.14.5", "openpyxl lo abre: hojas Resumen y TOTAL DE ÓRDENES ABIERTAS", "Resumen" in nombres and "TOTAL DE ÓRDENES ABIERTAS" in nombres, nombres)
        if "TOTAL DE ÓRDENES ABIERTAS" in nombres:
            ws = wb["TOTAL DE ÓRDENES ABIERTAS"]
            cab = [c.value for c in next(ws.iter_rows(min_row=6, max_row=6))]
            vh.anotar("T2.14.5", "la hoja trae el aviso SAP, la clasificación de la tarjeta y el estado del equipo",
                      "AVISO SAP" in cab and "CLASIFICACIÓN" in cab and "ESTADO DEL EQUIPO" in cab, cab)
    except ImportError:
        print("   (openpyxl no está en este PC: se contaron las hojas dentro del zip)")
    except Exception as ex:
        vh.anotar("T2.14.5", "openpyxl lo abre", False, str(ex)[:100])
    st, cab, cuerpo = cabeceras(adm, "reporte_exportar.php?formato=pdf&zona=UIO")
    vh.anotar("T2.14.5", "pdf → 200 application/pdf que empieza por %PDF", st == 200 and cab.get("content-type", "").startswith("application/pdf") and cuerpo.startswith(b"%PDF"), f"{st} · {len(cuerpo)} bytes")
    st, cab, cuerpo = cabeceras(adm, "reporte_exportar.php?formato=pptx&mes=2026-09")
    diapos = 0
    try:
        with zipfile.ZipFile(io.BytesIO(cuerpo)) as z:
            diapos = len([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)])
    except Exception:
        diapos = 0
    vh.anotar("T2.14.5", "pptx → 200, tipo PowerPoint y ≥ 6 diapositivas", st == 200 and "presentationml" in cab.get("content-type", "") and diapos >= 6, f"{st} · diapositivas={diapos}")
    e1 = int(vh.sql("SELECT COUNT(*) n FROM bitacora WHERE accion = 'EXPORTAR_REPORTE' AND usuario_id = ?", [ids["admin_prueba"]])[0]["n"])
    vh.anotar("T2.14.5", "cada exportación queda como EXPORTAR_REPORTE", e1 == e0 + 3, f"{e0} → {e1}")
    st, _, _ = cabeceras(adm, "reporte_exportar.php?formato=doc")
    vh.anotar("T2.14.5", "un formato desconocido → 400", st == 400, st)
    st, _, _ = cabeceras(jefe, "reporte_exportar.php?formato=pdf&zona=CNLJ")
    vh.anotar("T2.14.5", "el jefe de UIO pide el PDF de CNLJ y recibe el suyo (el alcance manda), no un error", st == 200, st)

    print("\n== 3. el cronograma en la base ==")
    salida = vh.ssh(f"cd {vh.D} && php cronograma_importar_cli.php --ensayo")
    print("   " + salida.strip().splitlines()[0][:150])
    salida = vh.ssh(f"cd {vh.D} && php cronograma_importar_cli.php")
    print("   " + "\n   ".join(salida.strip().splitlines()[-2:]))
    n = int(vh.sql("SELECT COUNT(*) n FROM ingresos_preventivos")[0]["n"])
    vh.anotar("T2.14.5", "cronograma_importar_cli.php deja filas en ingresos_preventivos", n > 0, n)
    salida2 = vh.ssh(f"cd {vh.D} && php cronograma_importar_cli.php")
    n2 = int(vh.sql("SELECT COUNT(*) n FROM ingresos_preventivos")[0]["n"])
    vh.anotar("T2.14.5", "correrlo dos veces no duplica", n2 == n, f"{n} → {n2}")
    st, _, c = adm.pedir("cronograma.php")
    j = json.loads(c or "{}")
    vh.anotar("T2.14.5", "cronograma.php devuelve fuente «tabla», csrf y módulos", st == 200 and j.get("fuente") == "tabla" and j.get("csrf") and isinstance(j.get("modulos"), list) and len(j["modulos"]) > 3, f"{st} · fuente={j.get('fuente')}")
    vh.anotar("T2.14.5", "cada ingreso trae ingreso_id y estado_hoy, y no viaja «origen»", bool(j.get("cronograma", {}).get("ingresos")) and "ingreso_id" in j["cronograma"]["ingresos"][0] and "estado_hoy" in j["cronograma"]["ingresos"][0] and "origen" not in j.get("cronograma", {}), "")
    st, _, c = jefe.pedir("cronograma.php")
    jj = json.loads(c or "{}")
    zonas = {i.get("zona") for i in jj.get("cronograma", {}).get("ingresos", [])}
    vh.anotar("T2.14.5", "el jefe de UIO solo recibe ingresos de UIO", st == 200 and zonas <= {"UIO"} and len(zonas) == 1, sorted(zonas))

    print("\n== 4. el cronograma escribe, con alcance y con rastro ==")
    cnlj = vh.sql("SELECT ingreso_id, plan_vigente_inicio FROM ingresos_preventivos WHERE zona = 'CNLJ' AND estado <> 'CUMPLIDO' AND plan_vigente_inicio IS NOT NULL LIMIT 1")
    uio = vh.sql("SELECT ingreso_id, local_codigo, plan_vigente_inicio, plan_vigente_fin, plan_original_inicio, kit_estado, kit_fecha, kit_nota, estado "
                 "FROM ingresos_preventivos WHERE zona = 'UIO' AND estado <> 'CUMPLIDO' AND plan_vigente_inicio IS NOT NULL AND actualizado_por IS NULL LIMIT 1")
    if not cnlj or not uio:
        vh.anotar("T2.14.5", "hay ingresos de CNLJ y de UIO para probar", False, f"cnlj={len(cnlj)} uio={len(uio)}")
    else:
        idc, idu = int(cnlj[0]["ingreso_id"]), int(uio[0]["ingreso_id"])
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "reagendar", "ingreso_id": idc, "inicio": "2026-12-01", "fin": "2026-12-02", "motivo": "Pedido de Grupo KFC"})
        sigue = vh.sql("SELECT plan_vigente_inicio FROM ingresos_preventivos WHERE ingreso_id = ?", [idc])[0]["plan_vigente_inicio"]
        vh.anotar("T2.14.5", "el jefe de UIO no reagenda un ingreso de CNLJ → 403 y nada cambia", st == 403 and sigue == cnlj[0]["plan_vigente_inicio"], f"{st} · {sigue}")
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "reagendar", "ingreso_id": idu, "inicio": "2026-12-03", "fin": "2026-12-04", "motivo": ""})
        vh.anotar("T2.14.5", "reagendar sin motivo → 400", st == 400, st)
        n0 = int(vh.sql("SELECT COUNT(*) n FROM cronograma_novedades WHERE ingreso_id = ? AND tipo = 'REAGENDA'", [idu])[0]["n"])
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "reagendar", "ingreso_id": idu, "inicio": "2026-12-03", "fin": "2026-12-04", "motivo": "Falta el kit de mantenimiento", "detalle": "PRUEBA T2.14.5"})
        f = vh.sql("SELECT plan_vigente_inicio, plan_original_inicio, actualizado_por FROM ingresos_preventivos WHERE ingreso_id = ?", [idu])[0]
        nov = vh.sql("SELECT fecha_antes, fecha_despues, motivo FROM cronograma_novedades WHERE ingreso_id = ? AND tipo = 'REAGENDA' ORDER BY novedad_id DESC LIMIT 1", [idu])
        vh.anotar("T2.14.5", "sobre uno de UIO → 200: se mueve el vigente, NO el original, y queda la novedad REAGENDA",
                  st == 200 and f["plan_vigente_inicio"] == "2026-12-03" and f["plan_original_inicio"] == uio[0]["plan_original_inicio"]
                  and int(vh.sql("SELECT COUNT(*) n FROM cronograma_novedades WHERE ingreso_id = ? AND tipo = 'REAGENDA'", [idu])[0]["n"]) == n0 + 1
                  and nov and nov[0]["fecha_antes"] == uio[0]["plan_vigente_inicio"] and nov[0]["fecha_despues"] == "2026-12-03", f"{st} · {f} · {nov}")
        vh.anotar("T2.14.5", "bitácora CRONOGRAMA_REAGENDA", int(vh.sql("SELECT COUNT(*) n FROM bitacora WHERE accion = 'CRONOGRAMA_REAGENDA' AND referencia = ?", [str(idu)])[0]["n"]) >= 1, "")
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "kit", "ingreso_id": idu, "estado": "CONFIRMADO", "fecha": "2026-09-13", "nota": "PRUEBA T2.14.5"})
        k = vh.sql("SELECT kit_estado, kit_fecha FROM ingresos_preventivos WHERE ingreso_id = ?", [idu])[0]
        vh.anotar("T2.14.5", "actualizar el kit → 200 y queda CONFIRMADO con fecha", st == 200 and k["kit_estado"] == "CONFIRMADO" and k["kit_fecha"] == "2026-09-13", f"{st} · {k}")
        st, _ = post_json_sin_csrf(jefe, "cronograma_accion.php", {"accion": "nota", "ingreso_id": idu, "nota": "sin token"})
        vh.anotar("T2.14.5", "POST sin X-Csrf → 403", st == 403, st)
        salida3 = vh.ssh(f"cd {vh.D} && php cronograma_importar_cli.php")
        f2 = vh.sql("SELECT plan_vigente_inicio, kit_estado FROM ingresos_preventivos WHERE ingreso_id = ?", [idu])[0]
        vh.anotar("T2.14.5", "volver a importar el JSON no pisa lo editado desde la app", f2["plan_vigente_inicio"] == "2026-12-03" and f2["kit_estado"] == "CONFIRMADO", f2)
        # se devuelve el ingreso real a como estaba
        vh.ejecutar("DELETE FROM cronograma_novedades WHERE ingreso_id = ? AND (detalle = 'PRUEBA T2.14.5' OR motivo LIKE 'Kit %')", [idu])
        vh.ejecutar("UPDATE ingresos_preventivos SET plan_vigente_inicio = ?, plan_vigente_fin = ?, kit_estado = ?, kit_fecha = ?, kit_nota = ?, "
                    "actualizado_por = NULL, actualizado_en = NULL WHERE ingreso_id = ?",
                    [uio[0]["plan_vigente_inicio"], uio[0]["plan_vigente_fin"], uio[0]["kit_estado"], uio[0]["kit_fecha"], uio[0]["kit_nota"], idu])
        vh.anotar("T2.14.5", "y se restaura", vh.sql("SELECT plan_vigente_inicio FROM ingresos_preventivos WHERE ingreso_id = ?", [idu])[0]["plan_vigente_inicio"] == uio[0]["plan_vigente_inicio"], "")

    # Un ingreso sintético (año 2099) para agendar, anotar y cerrar sin tocar datos reales.
    local_uio = uio[0]["local_codigo"] if uio else vh.sql("SELECT local_codigo FROM ingresos_preventivos WHERE zona = 'UIO' LIMIT 1")[0]["local_codigo"]
    vh.ejecutar("DELETE n FROM cronograma_novedades n JOIN ingresos_preventivos i ON i.ingreso_id = n.ingreso_id WHERE i.anio = 2099")
    vh.ejecutar("DELETE FROM ingresos_preventivos WHERE anio = 2099")
    st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "agendar", "local": local_uio, "numero": 4, "anio": 2099, "inicio": "2099-01-10", "fin": "2099-01-11", "kit_confirmado": False})
    j = json.loads(c or "{}")
    sint = vh.sql("SELECT ingreso_id, zona, plan_original_inicio, plan_vigente_inicio, kit_estado, estado FROM ingresos_preventivos WHERE anio = 2099 AND local_codigo = ?", [local_uio])
    vh.anotar("T2.14.5", "agendar un local de UIO → 200, con el plan original fijado y sin kit", st == 200 and sint and sint[0]["zona"] == "UIO" and sint[0]["plan_original_inicio"] == "2099-01-10" and sint[0]["kit_estado"] == "SIN_KIT", f"{st} · {sint}")
    if sint:
        sid = int(sint[0]["ingreso_id"])
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "agendar", "local": local_uio, "numero": 4, "anio": 2099, "inicio": "2099-02-10", "fin": "2099-02-11"})
        vh.anotar("T2.14.5", "agendarlo otra vez → 409 (ya está agendado: se reagenda)", st == 409, st)
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "nota", "ingreso_id": sid, "tipo": "Acceso negado", "nota": "PRUEBA T2.14.5: nota"})
        vh.anotar("T2.14.5", "anotar una novedad → 200 y fila NOTA", st == 200 and int(vh.sql("SELECT COUNT(*) n FROM cronograma_novedades WHERE ingreso_id = ? AND tipo = 'NOTA'", [sid])[0]["n"]) == 1, st)
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "cerrar", "ingreso_id": sid, "real_inicio": "2099-01-10", "real_fin": "2099-01-12", "ot": "OT-99999-PRUEBA-UIO", "nota": "PRUEBA"})
        f = vh.sql("SELECT estado, real_fin, ot_ids FROM ingresos_preventivos WHERE ingreso_id = ?", [sid])[0]
        nov = vh.sql("SELECT motivo FROM cronograma_novedades WHERE ingreso_id = ? AND tipo = 'CIERRE'", [sid])
        vh.anotar("T2.14.5", "cerrar → CUMPLIDO con fecha real, la OT anotada y la novedad dice «tarde» (fin real después del plan)",
                  st == 200 and f["estado"] == "CUMPLIDO" and f["real_fin"] == "2099-01-12" and "OT-99999" in (f["ot_ids"] or "") and nov and "tarde" in nov[0]["motivo"], f"{st} · {f} · {nov}")
        st, _, c = jefe.pedir("cronograma_accion.php", cuerpo_json={"accion": "reagendar", "ingreso_id": sid, "inicio": "2099-03-01", "motivo": "Otro"})
        vh.anotar("T2.14.5", "un ingreso cumplido no se reagenda → 409", st == 409, st)
        st, _, c = adm.pedir("cronograma.php")
        j = json.loads(c or "{}")
        el = [i for i in j.get("cronograma", {}).get("ingresos", []) if i.get("ingreso_id") == sid]
        vh.anotar("T2.14.5", "cronograma.php lo devuelve cumplido, con sus novedades", el and el[0].get("estado_hoy") == "cumplido" and len(el[0].get("novedades", [])) >= 3, el[0].get("estado_hoy") if el else None)
        vh.ejecutar("DELETE FROM cronograma_novedades WHERE ingreso_id = ?", [sid])
        vh.ejecutar("DELETE FROM ingresos_preventivos WHERE ingreso_id = ?", [sid])
    st, _, c = adm.pedir("cronograma_accion.php", cuerpo_json={"accion": "agendar", "local": "ZZZ999", "numero": 1, "anio": 2099, "inicio": "2099-01-10"})
    vh.anotar("T2.14.5", "agendar un local que no está en el maestro → 400", st == 400, st)

    for se in s.values():
        se.pedir("salir.php", form={})
    fallas = [r for r in vh.resultados if not r["ok"]]
    print(f"\n{len(vh.resultados) - len(fallas)} de {len(vh.resultados)} comprobaciones pasan; {len(fallas)} fallan")
    return 1 if fallas else 0


if __name__ == "__main__":
    sys.exit(main())
