"""
T2.27.2 — La respuesta de los miércoles al «REPORTE 2026 SEMANA NN - MANTENIMIENTO CORRECTIVO» de KFC.

QUÉ PIDE KFC
Cada lunes Erika Zambrano (KFC) manda a los cuatro proveedores el export de SAP
de todo el año (~44.000 filas, 16 MB) y pide «la validación y cierre de las
órdenes correspondientes al reporte adjunto hasta el miércoles a las 13h00».
El jueves se revisa en la reunión semanal y KFC publica el porcentaje de
cierre de cada proveedor: en la semana 37 INDUSTEC quedó en 47,62 %, el más
rezagado junto con Megaservicios. La administración devuelve el mismo Excel
con una hoja RESUMEN —los correctivos de INDUSTEC abiertos o en tratamiento—
donde la columna «Estatus 2 del Aviso» se reemplaza por «ESTATUS IND», el
estado de gestión que INDUSTEC le da a cada uno (CERRADO, INFORME TECNICO,
PENDIENTE OK OP´S, IMPORTACION...). Hoy lo hace a mano y sale después del
plazo: casi siempre el miércoles por la tarde.

QUÉ HACE ESTE GENERADOR
1. Toma el Excel de KFC TAL CUAL y le AGREGA tres hojas —RESUMEN, RESUMEN IND
   (el conteo por ESTATUS IND que va en el correo) y ND INDUSTEC (las órdenes sin
   proveedor que KFC le asigna a INDUSTEC)— escribiéndolas directamente en el
   paquete del archivo. Las hojas de KFC y sus tablas dinámicas no se tocan: se
   copian byte a byte. Abrirlo y guardarlo con openpyxl las habría perdido.
2. PROPONE el ESTATUS IND de cada aviso con su evidencia. No se inventa: lo que
   solo sabe la administración (una importación, un despacho coordinado por
   teléfono) sale del STATUS_PENDIENTES que ella mandó el martes y de lo que
   ella misma puso la semana anterior. La hoja REVISION dice de dónde sale cada uno.
3. Deja el texto del correo con las cifras.

Uso:
    .venv/Scripts/python.exe scripts/t2_27_respuesta_kfc.py            # baja del correo lo de esta semana
    .venv/Scripts/python.exe scripts/t2_27_respuesta_kfc.py --sap <REPORTE SEMANA 38>.xlsx \\
        --status <STATUS_PENDIENTES del martes>.xlsx --anterior <respuesta de la semana 37>.xlsx \\
        --corte 2026-09-16 --comparar <respuesta real de la semana 38>.xlsx
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import t2_27_fuentes as F  # noqa: E402
from comun import corto, termino  # noqa: E402

# Vocabulario: los valores de ESTATUS IND (CERRADO, INFORME TECNICO, PENDIENTE OK OP´S…),
# las hojas RESUMEN, RESUMEN IND y ND INDUSTEC y las columnas del libro de KFC son contrato
# y se escriben tal cual. Lo que no es de KFC ni de Isabel —la evidencia de cada propuesta,
# la hoja LEEME, la consola— nombra los estados con el diccionario único (comun.termino).

# El RESUMEN de ella tiene las columnas de FILTRO con «Estatus 2 del Aviso» reemplazada
# por «ESTATUS IND» en la misma posición (semana 38, Sent UID 10231).
COLUMNA_REEMPLAZADA = "Estatus 2 del Aviso"


def traducir_responsable(k: str) -> str | None:
    """RESPONSABLES del STATUS del martes -> vocabulario de ESTATUS IND que ella usa con KFC."""
    t = (k or "").upper()
    for clave, estado in (("OK OP", "PENDIENTE OK OP´S"), ("OK MARJORIE", "PENDIENTE OK OP´S"), ("IMPORTACION", "IMPORTACION"),
                          ("DESPACHO", "DESPACHO BOD"), ("MIGUEL", "GESTION MIGUEL V."), ("RONALD", "GESTION RONALD V."),
                          ("BODEGA", "GESTION BODEGA"), ("PROVEEDOR", "GESTION PROVEEDORES"), ("COTIZ", "COTIZACION PENDIENTE")):
        if clave in t:
            return estado
    return None


def leer_status(ruta: Path | None) -> dict:
    if not ruta:
        return {}
    ws = openpyxl.load_workbook(ruta, data_only=True)["ORDENES"]
    return {str(int(r[1])) if isinstance(r[1], (int, float)) else str(r[1]).strip(): r
            for r in ws.iter_rows(min_row=2, max_col=13, values_only=True) if r[1] not in (None, "")}


def leer_respuesta(ruta: Path | None) -> dict:
    """{aviso: ESTATUS IND} de una respuesta ya enviada (su hoja RESUMEN)."""
    if not ruta:
        return {}
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if "RESUMEN" not in wb.sheetnames:
        return {}
    filas = wb["RESUMEN"].iter_rows(values_only=True)
    enc = [str(h or "").strip() for h in next(filas)]
    if "ESTATUS IND" not in enc:
        return {}
    i = enc.index("ESTATUS IND")
    return {str(r[0]).strip(): str(r[i] or "").strip() for r in filas if r and r[0]}


def proponer(aviso: str, ots: dict, status: dict, previa: dict) -> tuple[str, str]:
    """(ESTATUS IND propuesto, de dónde sale). En orden de solidez de la evidencia."""
    mias = ots.get(aviso, [])
    cierre = [o for o in mias if o["estado_ot"] == "CERRADA"]
    if cierre:
        return "CERRADO", f"{termino('OT_CIERRE')} {cierre[-1]['id_industec']} del {cierre[-1]['fecha_atencion']:%d/%m}"
    if aviso in status:
        k = status[aviso][10]
        t = traducir_responsable(str(k or ""))
        if t:
            return t, f"STATUS del martes: RESPONSABLES «{k}»"
    if previa.get(aviso):
        return previa[aviso], "lo que se reportó la semana anterior (sin evidencia nueva)"
    if mias:
        return "INFORME TECNICO", (f"{termino('ABIERTA')}: hay visita ({mias[-1]['id_industec']} del "
                                   f"{mias[-1]['fecha_atencion']:%d/%m}) pero no {termino('OT_CIERRE')}")
    return "INFORME TECNICO", (f"{termino('ESPERA_INFORME')}: no hay ninguna {termino('OT_INDUSTEC')} con este "
                               f"{termino('AVISO_SAP')} (falta la visita o su OT)")


# ---------------------------------------------------------------------------
#  Escribir hojas nuevas dentro del paquete del Excel de KFC, sin tocar lo demás
# ---------------------------------------------------------------------------

def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _excel_serial(d: dt.date) -> float:
    return (dt.datetime.combine(d, dt.time()) - dt.datetime(1899, 12, 30)).days if not isinstance(d, dt.datetime) \
        else (d - dt.datetime(1899, 12, 30)).total_seconds() / 86400


def hoja_xml(filas: list[list], estilo_enc: str, estilo_fecha: str, anchos: dict[int, float] | None = None) -> str:
    """Una hoja con cadenas en línea (no toca sharedStrings) y fechas con el estilo de fecha del libro."""
    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
           '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
           '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>',
           '<sheetFormatPr defaultRowHeight="15"/>']
    if anchos:
        out.append("<cols>" + "".join(f'<col min="{c}" max="{c}" width="{w}" customWidth="1"/>' for c, w in sorted(anchos.items())) + "</cols>")
    out.append("<sheetData>")
    for i, fila in enumerate(filas, start=1):
        celdas = []
        for j, v in enumerate(fila, start=1):
            ref = f"{_col(j)}{i}"
            est = f' s="{estilo_enc}"' if i == 1 and estilo_enc else ""
            if v is None or v == "":
                continue
            if isinstance(v, (dt.date, dt.datetime)):
                celdas.append(f'<c r="{ref}" s="{estilo_fecha}"><v>{_excel_serial(v)}</v></c>')
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                celdas.append(f'<c r="{ref}"{est}><v>{v}</v></c>')
            elif isinstance(v, dt.time):
                celdas.append(f'<c r="{ref}"{est} t="inlineStr"><is><t>{v:%H:%M}</t></is></c>')
            else:
                # El texto del gerente en «Circunstancia» trae a veces caracteres de control que
                # no caben en XML: uno solo haría que Excel pida reparar el libro entero.
                texto = escape(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(v))).replace("\r", "")
                celdas.append(f'<c r="{ref}"{est} t="inlineStr"><is><t xml:space="preserve">{texto}</t></is></c>')
        out.append(f'<row r="{i}">' + "".join(celdas) + "</row>")
    out.append("</sheetData>")
    if len(filas) > 1:
        out.append(f'<autoFilter ref="A1:{_col(len(filas[0]))}{len(filas)}"/>')
    out.append('<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>')
    return "".join(out)


def agregar_hojas(origen: Path, destino: Path, hojas: list[tuple[str, list[list], dict]]) -> None:
    """Copia el libro de KFC y le agrega hojas al final. Todo lo que ya tenía se copia sin tocar."""
    with zipfile.ZipFile(origen) as z:
        partes = {i.filename: (i, z.read(i.filename)) for i in z.infolist()}
    libro = partes["xl/workbook.xml"][1].decode("utf-8")
    rels = partes["xl/_rels/workbook.xml.rels"][1].decode("utf-8")
    tipos = partes["[Content_Types].xml"][1].decode("utf-8")
    nombres = re.findall(r'<sheet [^>]*name="([^"]+)"', libro)
    # Estilos para reutilizar: el del encabezado de FILTRO (A1) y el de su columna de fecha (B2).
    rid_filtro = re.search(r'<sheet [^>]*name="FILTRO"[^>]*r:id="([^"]+)"', libro).group(1)
    target = re.search(rf'<Relationship [^>]*Id="{rid_filtro}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{rid_filtro}"', rels)
    xml_filtro = partes["xl/" + target.group(1).lstrip("/").replace("xl/", "", 1)][1][:400_000].decode("utf-8", "ignore")
    enc = re.search(r'<c r="A1" s="(\d+)"', xml_filtro)
    fec = re.search(r'<c r="B2" s="(\d+)"', xml_filtro)
    estilo_enc, estilo_fecha = (enc.group(1) if enc else ""), (fec.group(1) if fec else "0")
    ids_hoja = [int(x) for x in re.findall(r'sheetId="(\d+)"', libro)]
    ids_rel = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
    n_parte = max([int(x) for x in re.findall(r"worksheets/sheet(\d+)\.xml", " ".join(partes))] + [0])
    nuevas = {}
    for k, (nombre, filas, anchos) in enumerate(hojas, start=1):
        if nombre in nombres:
            raise SystemExit(f"ABORTADO: el libro de KFC ya tiene una hoja «{nombre}».")
        parte = f"xl/worksheets/sheet{n_parte + k}.xml"
        rid = f"rId{max(ids_rel) + k}"
        nuevas[parte] = hoja_xml(filas, estilo_enc, estilo_fecha, anchos).encode("utf-8")
        libro = libro.replace("</sheets>", f'<sheet name="{escape(nombre)}" sheetId="{max(ids_hoja) + k}" r:id="{rid}"/></sheets>')
        rels = rels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{n_parte + k}.xml"/></Relationships>')
        tipos = tipos.replace("</Types>", f'<Override PartName="/{parte}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
    tmp = destino.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, (info, datos) in partes.items():
            if nombre == "xl/workbook.xml":
                datos = libro.encode("utf-8")
            elif nombre == "xl/_rels/workbook.xml.rels":
                datos = rels.encode("utf-8")
            elif nombre == "[Content_Types].xml":
                datos = tipos.encode("utf-8")
            z.writestr(info, datos)
        for parte, datos in nuevas.items():
            z.writestr(parte, datos)
    tmp.replace(destino)


# ---------------------------------------------------------------------------

def generar(sap_ruta: Path, status_ruta: Path | None, anterior_ruta: Path | None, corte: dt.date, salida: Path) -> dict:
    semana = F.semana_del_reporte_kfc(sap_ruta)
    sap = F.leer_sap_semanal(sap_ruta)
    # Las filas de FILTRO completas, en su orden y con todas sus columnas, para el RESUMEN.
    wb = openpyxl.load_workbook(sap_ruta, read_only=True, data_only=True)
    filas_it = wb["FILTRO"].iter_rows(values_only=True)
    enc = [str(h or "").strip() for h in next(filas_it)]
    ix = {h: i for i, h in enumerate(enc)}
    for h in ("AVISO", "DESCRIPCIÓN", "ESTATUS A", "PROVEEDOR", COLUMNA_REEMPLAZADA):
        if h not in ix:
            raise SystemExit(f"{sap_ruta.name}: FILTRO no tiene la columna {h!r}.")
    mias, vistos = [], set()
    for r in filas_it:
        if not r or r[ix["AVISO"]] in (None, ""):
            continue
        a = str(r[ix["AVISO"]]).strip()
        if (r[ix["PROVEEDOR"]] == "INDUSTEC" and r[ix["DESCRIPCIÓN"]] == "Mant. Correctivo"
                and r[ix["ESTATUS A"]] in ("ABIERTO", "TRATAMIENTO") and a not in vistos):
            vistos.add(a)
            mias.append(list(r))
    hojas_kfc = list(wb.sheetnames)
    nd = []
    if "#O_ND" in wb.sheetnames:
        it = wb["#O_ND"].iter_rows(values_only=True)
        enc_nd = list(next(it))
        i_prov = [i for i, h in enumerate(enc_nd) if str(h or "").strip().upper() == "PROVEEDOR"]
        nd = [enc_nd] + [list(r) for r in it if r and i_prov and r[i_prov[0]] == "INDUSTEC"]
    wb.close()

    cnx = F.conectar()
    ots = F.ordenes_por_aviso(cnx, corte)
    cnx.close()
    status = leer_status(status_ruta)
    previa = leer_respuesta(anterior_ruta)

    i_rep = ix[COLUMNA_REEMPLAZADA]
    enc_res = enc[:i_rep] + ["ESTATUS IND"] + enc[i_rep + 1:]
    resumen, revision, conteo = [enc_res], [], collections.Counter()
    for r in mias:
        a = str(r[ix["AVISO"]]).strip()
        est, porque = proponer(a, ots, status, previa)
        conteo[est] += 1
        resumen.append(r[:i_rep] + [est] + r[i_rep + 1:])
        revision.append([a, r[ix["ESTATUS A"]], sap[a].get("local") if a in sap else "", est, porque, previa.get(a, ""),
                         (status[a][10] if a in status else "")])
    orden_estados = sorted(conteo.items(), key=lambda x: (x[0] != "CERRADO", -x[1], x[0]))
    tabla_ind = [["ESTATUS IND", "Cuenta de AVISO"]] + [[e, n] for e, n in orden_estados] + [["Total general", sum(conteo.values())]]

    salida.mkdir(parents=True, exist_ok=True)
    base = re.sub(r"^(\d{4}-\d{2}-\d{2} INBOX |\d+_)", "", sap_ruta.name)
    destino = salida / (Path(base).stem.strip() + " (respuesta INDUSTEC, generado agente).xlsx")
    hojas = [("RESUMEN", resumen, {1: 11, 2: 12, 5: 28, 9: 20}), ("RESUMEN IND", tabla_ind, {1: 26, 2: 16})]
    if len(nd) > 1:
        hojas.append(("ND INDUSTEC", nd, {1: 11, 4: 18}))
    agregar_hojas(sap_ruta, destino, hojas)

    # Verificación cruzada (I-10): se relee el archivo y el RESUMEN tiene que tener exactamente los
    # avisos del filtro, y la tabla IND tiene que sumar lo mismo. Si no, se borra y se aborta.
    chk = openpyxl.load_workbook(destino, read_only=True, data_only=True)
    avisos_res = [str(x[0]) for x in chk["RESUMEN"].iter_rows(min_row=2, values_only=True) if x and x[0]]
    total_ind = [x for x in chk["RESUMEN IND"].iter_rows(values_only=True)][-1][1]
    # Las hojas de KFC tienen que seguir ahí, en su orden, antes de las nuevas.
    hojas_ok = chk.sheetnames == hojas_kfc + [h[0] for h in hojas]
    chk.close()
    if sorted(avisos_res) != sorted(str(r[ix["AVISO"]]).strip() for r in mias) or total_ind != len(mias) or not hojas_ok:
        destino.unlink()
        raise SystemExit(f"ABORTADO: el archivo no cuadra (RESUMEN {len(avisos_res)} vs {len(mias)}, IND {total_ind}). Se borró.")
    return dict(destino=destino, semana=semana, mias=mias, conteo=conteo, orden=orden_estados, revision=revision,
                nd=nd[1:], sap=sap_ruta, status=status_ruta, anterior=anterior_ruta, corte=corte, ix=ix)


def escribir_revision(res: dict) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEEME"
    for i, t in enumerate([
        f"Respuesta al REPORTE SEMANA {res['semana']} de KFC — propuesta del agente, para revisar antes de enviar",
        "",
        f"Excel de KFC: {res['sap'].name}",
        f"STATUS del martes (responsables): {res['status'].name if res['status'] else 'no se usó'}",
        f"Respuesta de la semana anterior (ESTATUS IND previo): {res['anterior'].name if res['anterior'] else 'no se usó'}",
        f"{termino('OT_INDUSTEC')}: base local, hasta el {res['corte']:%d/%m/%Y}.",
        "",
        f"RESUMEN: {len(res['mias'])} {termino('ORDEN', len(res['mias']))} correctivas de INDUSTEC abiertas o en tratamiento "
        "en SAP (mismo filtro que la hoja RESUMEN que se devuelve a KFC).",
        f"ESTATUS IND es una PROPUESTA. En orden de solidez: CERRADO si hay {termino('OT_CIERRE')}; lo que dice el STATUS del martes;",
        f"lo que se puso la semana anterior; e INFORME TECNICO si no hay {termino('OT_CIERRE')}. Revise la columna «De dónde sale».",
        "Lo que solo sabe la administración (un despacho coordinado por teléfono, una importación) hay que ajustarlo a mano.",
        "",
        "El libro de KFC se devuelve intacto: sus hojas y tablas dinámicas no se tocaron; se agregaron RESUMEN, RESUMEN IND y ND INDUSTEC.",
    ], start=1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 140
    h = wb.create_sheet("PROPUESTA")
    h.append(["AVISO", "ESTATUS A (SAP)", "Local", "ESTATUS IND propuesto", "De dónde sale", "Semana anterior", "STATUS martes (RESPONSABLES)"])
    for r in res["revision"]:
        h.append(r)
    for c, w in zip("ABCDEFG", (11, 14, 9, 24, 70, 22, 36)):
        h.column_dimensions[c].width = w
    h.freeze_panes = "A2"
    destino = res["destino"].with_name(f"REVISION RESPUESTA SEMANA {res['semana']} (generado agente).xlsx")
    wb.save(destino)
    return destino


def cuerpo_correo(res: dict) -> str:
    total = len(res["mias"])
    cer = res["conteo"].get("CERRADO", 0)
    abiertas = [(e, n) for e, n in res["orden"] if e != "CERRADO"]
    lineas = [
        "Estimados,",
        # Las zonas con su rótulo del diccionario, como en el STATUS del martes y en B.IA.
        "Se envía el estatus de las órdenes correspondientes a las tres zonas que maneja Industec "
        f"({corto('ZONA_UIO')}, {corto('ZONA_LARB')} y {corto('ZONA_CNLJ')}).",
        f"Actualmente se registra un total de {total} órdenes, de las cuales {cer} se encuentran en estado CERRADO"
        + (f" y {total - cer} continúan en gestión:" if abiertas else "."),
    ]
    lineas += [f"  *   {e}: {n}" for e, n in abiertas]
    if res["nd"]:
        lineas.append(f"De las {termino('ND_SIN_PROVEEDOR', 2)}, {len(res['nd'])} corresponden a Industec.")
    lineas += ["Se adjunta el reporte con la hoja RESUMEN.", "Saludos cordiales,"]
    return "\n".join(lineas)


def comparar(generado: Path, real: Path) -> None:
    """ESTATUS IND propuesto contra el que ella puso, aviso por aviso."""
    g, r = leer_respuesta(generado), leer_respuesta(real)
    comunes = sorted(set(g) & set(r))
    iguales = [a for a in comunes if g[a].upper().replace("´", "'") == r[a].upper().replace("´", "'")]
    por_real = collections.Counter((r[a], g[a]) for a in comunes if a not in iguales)
    print(f"\n== ESTATUS IND contra {real.name}")
    print(f"  Avisos: generado {len(g)} · real {len(r)} · en ambos {len(comunes)}")
    print(f"  Idénticos: {len(iguales)} de {len(comunes)} ({len(iguales) / max(1, len(comunes)) * 100:.1f} %)")
    print(f"  Diferencias (real -> propuesto): {dict(por_real.most_common())}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--sap", help="Excel semanal de KFC; por omisión, el último de Recibidos")
    ap.add_argument("--status", help="STATUS_PENDIENTES del martes; por omisión, el último de Enviados")
    ap.add_argument("--anterior", help="respuesta enviada la semana previa; por omisión, la última de Enviados")
    ap.add_argument("--corte", help="OT INDUSTEC hasta esta fecha (AAAA-MM-DD); por omisión, hoy")
    ap.add_argument("--comparar", help="respuesta real de la misma semana, para medir la propuesta")
    a = ap.parse_args()
    corte = dt.date.fromisoformat(a.corte) if a.corte else dt.date.today()
    sap = Path(a.sap) if a.sap else (F.bajar_ultimo_adjunto("INBOX", r"REPORTE.*SEMANA.*CORRECTIVO.*\.xlsx$", dias=10, antes_de=corte + dt.timedelta(days=1)) or [None])[0]
    if not sap:
        raise SystemExit("No encontré en Recibidos el REPORTE ... SEMANA NN ... CORRECTIVO de KFC. Pásalo con --sap.")
    status = Path(a.status) if a.status else (F.bajar_ultimo_adjunto("Sent", r"STATUS_PENDIENTES.*\.xlsx$", dias=10, antes_de=corte + dt.timedelta(days=1)) or [None])[0]
    semana = F.semana_del_reporte_kfc(sap)
    if a.anterior:
        anterior = Path(a.anterior)
    else:
        b = F.bajar_ultimo_adjunto("Sent", rf"REPORTE.*SEMANA\s+{(semana or 0) - 1}\b.*\.xlsx$", dias=14, antes_de=corte + dt.timedelta(days=1))
        anterior = b[0] if b else None
    print(f"Respuesta a la semana {semana}\n  KFC: {sap}\n  STATUS: {status}\n  anterior: {anterior}")
    res = generar(sap, status, anterior, corte, F.SALIDAS / f"{corte:%Y-%m-%d}")
    rev = escribir_revision(res)
    correo = res["destino"].with_name(f"CORREO RESPUESTA SEMANA {res['semana']} (generado agente).txt")
    correo.write_text(cuerpo_correo(res), encoding="utf-8")
    print(f"\n  {len(res['mias'])} {termino('ORDEN', len(res['mias']))} correctivas de INDUSTEC abiertas o en tratamiento "
          f"en SAP · {termino('ND_SIN_PROVEEDOR', 2)} de INDUSTEC: {len(res['nd'])}")
    print(f"  ESTATUS IND propuesto: {dict(res['orden'])}")
    print(f"  Respuesta: {res['destino']}\n  Revisión:  {rev}\n  Correo:    {correo}")
    if a.comparar:
        comparar(res["destino"], Path(a.comparar))


if __name__ == "__main__":
    main()
