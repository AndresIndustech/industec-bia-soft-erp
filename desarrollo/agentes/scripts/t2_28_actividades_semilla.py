"""
T2.28.12a — Biblioteca de actividades: minería del histórico, solo lectura en la estación
(observación 5 de INDUSTEC: "guía de actividades por técnico, preventivo y correctivo").

QUÉ ES
Hoy el técnico redacta "actividades" y la descripción del preventivo desde cero, en texto
libre. Este script mina lo que YA se hizo en 7.452 órdenes reales (2.972 descripciones de
equipo en preventivo, 6.692 textos de "actividades" en correctivo — cifras de esta corrida,
no las del documento de especificación: ver más abajo) y arma un BORRADOR determinista de
la biblioteca de pasos por familia de equipo, para que el jefe técnico lo apruebe o corrija
en la pantalla `actividades.php` (T2.28.12c, fuera de esta subtarea).

SOLO LECTURA. No escribe en ninguna base, no toca el servidor, no aplica ninguna migración.
Es exactamente la Fase 1 de T2.28.12 (`T2_28_OBSERVACIONES_INDUSTEC.md` líneas 738-750):
--analizar es el único modo que existe en esta fase; 12b (migración 018 + carga, tras la
puerta D3) y 12c (la pantalla) son de otra conversación.

DE DÓNDE SALE CADA COSA
- `familias_equipo` y `diagnosticos` NO existen en la base local `industec_ots` (esa base es
  el corpus histórico de OTs, no la app del sitio de pruebas): se leen tal cual del propio
  SQL de semilla `desarrollo/sistema_ots/app/sql/009_semilla_diagnosticos.sql`, sin tocar
  Hostinger. Es dato estático ya versionado en el repo.
- `ot_equipos.descripcion` (preventivo) y `ots.actividades` (correctivo) salen de la base
  local `industec_ots` (MariaDB de la estación), filtrando `en_cuarentena = 0` y
  `correlativo < 90000` (sin las órdenes sintéticas de prueba), como en
  `t2_27_fuentes.ordenes_por_aviso` y en la skill `industec-agentes-y-entregables` §A3.

DOS DECISIONES QUE EL CÓDIGO REAL OBLIGÓ A TOMAR, DISTINTAS DE LA LETRA DEL DOCUMENTO
1. El documento dice "en correctivo, el primer equipo (orden=1) de la orden". En el esquema
   real `ot_equipos.orden` EMPIEZA EN 0: las 6.785 órdenes con un solo equipo lo tienen en
   orden=0, y toda orden con varios equipos también arranca en 0 (medido: orden 0 → 7.478
   filas, orden 1 → 693, ...). orden=1 es el SEGUNDO equipo, no el primero. Este script usa
   orden=0, igual que ya hace `t2_27_fuentes.ordenes_por_aviso` (`e.orden = 0`). Gana el
   código (regla del documento, cabecera de `T2_28_OBSERVACIONES_INDUSTEC.md`).
2. El mapa sustantivo→imperativo del documento (13 entradas) no cubre 3 de los 16 sustantivos
   que el propio patrón de corte usa: "Mantenimiento", "Descalcificación", "Desinfección".
   Se extendieron con el mismo criterio morfológico que el resto (Descalcificación→Descalcifica,
   Desinfección→Desinfecta, igual que Calibración→Calibra; Mantenimiento→"Realiza
   mantenimiento de", igual que Sujeción→"Ajusta la sujeción de"). Cada fila que usa una de
   estas tres queda marcada `mapa_extendido=True` en el JSON y con una nota en "Ayuda" del
   Excel, para que el jefe técnico las mire primero.

Uso:
    .venv/Scripts/python.exe scripts/t2_28_actividades_semilla.py --analizar
"""
from __future__ import annotations

import argparse
import collections
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import json
import mysql.connector
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comun import RAIZ, SALIDAS, cargar_env  # noqa: E402

SQL_SEMILLA = RAIZ / "desarrollo" / "sistema_ots" / "app" / "sql" / "009_semilla_diagnosticos.sql"

# --------------------------------------------------------------------------------------
#  Normalización de texto
# --------------------------------------------------------------------------------------

def _sin_tildes(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c))


def normalizar_mayus(s: str) -> str:
    """Mayúsculas, sin tildes, espacios colapsados — para resolver familia (spec 12a)."""
    return re.sub(r"\s+", " ", _sin_tildes(s or "").upper()).strip()


def normalizar_min(s: str) -> str:
    """Minúsculas, sin tildes, espacios colapsados — la clave de agrupación de actividades."""
    return re.sub(r"\s+", " ", _sin_tildes(s or "").lower()).strip(" .;:,-")


STOPWORDS = {"de", "la", "el", "en", "por", "con", "sin", "del", "al", "y", "o", "a", "que",
             "los", "las", "un", "una", "no", "se", "su", "para", "segun", "tras"}


def palabras_clave(txt: str) -> set[str]:
    return {p for p in (normalizar_min(w) for w in re.findall(r"[A-Za-zÁÉÍÓÚáéíóúÑñ]+", txt or ""))
            if p and p not in STOPWORDS and len(p) > 2}


# --------------------------------------------------------------------------------------
#  familias_equipo y diagnosticos: estáticos, del SQL de semilla ya versionado (no se
#  consulta ningún servidor: esas tablas ni existen en la base local de la estación).
# --------------------------------------------------------------------------------------

def cargar_familias() -> list[tuple[str, str, int]]:
    texto = SQL_SEMILLA.read_text(encoding="utf-8")
    bloque = re.search(r"INSERT INTO familias_equipo.*?VALUES\s*(.*?)\nON DUPLICATE", texto, re.S)
    if not bloque:
        raise SystemExit(f"ABORTADO: no encontré el INSERT de familias_equipo en {SQL_SEMILLA}. "
                          "¿Cambió el archivo? (gana el código: reviso a mano antes de seguir)")
    filas = re.findall(r"\('([^']+)',\s*'([^']+)',\s*(\d+),\s*(\d+)\)", bloque.group(1))
    if len(filas) != 23:
        raise SystemExit(f"ABORTADO: esperaba 23 familias_equipo (el comentario del propio SQL lo dice), "
                          f"parseé {len(filas)}. Reviso el parseo antes de confiar en el resto.")
    familias = sorted(((fam, patron, int(orden)) for fam, patron, orden, activo in filas), key=lambda f: f[2])
    return familias


def cargar_diagnosticos() -> list[dict]:
    texto = SQL_SEMILLA.read_text(encoding="utf-8")
    bloque = re.search(r"INSERT INTO diagnosticos.*?VALUES\s*(.*?)\nON DUPLICATE", texto, re.S)
    if not bloque:
        raise SystemExit(f"ABORTADO: no encontré el INSERT de diagnosticos en {SQL_SEMILLA}.")
    filas = re.findall(
        r"\('([^']*)',\s*'([^']*)',\s*'([^']*)',\s*'([^']*)',\s*(?:NULL|'[^']*'),\s*(\d+)\)",
        bloque.group(1))
    if len(filas) != 159:
        raise SystemExit(f"ABORTADO: esperaba 159 diagnosticos (comentario del propio SQL), "
                          f"parseé {len(filas)}. Reviso el parseo antes de confiar en el resto.")
    diags = []
    for codigo, familia, titulo, _texto, _orden in filas:
        diags.append({"codigo": codigo, "familia": familia, "titulo": titulo,
                       "_palabras": palabras_clave(titulo)})
    return diags


def resolver_familia(equipo: str | None, familias: list[tuple[str, str, int]]) -> str | None:
    norm = normalizar_mayus(equipo or "")
    if not norm:
        return None
    for fam, patron, _orden in familias:
        if re.search(patron, norm):
            return fam
    return None


def prefijo_familia(familia: str, diagnosticos: list[dict]) -> str:
    for d in diagnosticos:
        if d["familia"] == familia:
            return d["codigo"].split("-")[0]
    return normalizar_mayus(familia)[:3]  # no debería pasar: las 23 familias tienen diagnósticos semilla


def diagnostico_asociado(familia: str, texto_norm: str, diagnosticos: list[dict]) -> str | None:
    palabras_frase = palabras_clave(texto_norm)
    mejor, mejor_n = None, 1  # se exige >= 2 (spec): arranco en 1 para que 2 lo supere
    for d in diagnosticos:
        if d["familia"] != familia:
            continue
        comunes = len(palabras_frase & d["_palabras"])
        if comunes >= 2 and comunes > mejor_n:
            mejor, mejor_n = d["codigo"], comunes
    return mejor


# --------------------------------------------------------------------------------------
#  Partir en actividades atómicas (preventivo) / frases (correctivo)
# --------------------------------------------------------------------------------------

SUSTANTIVOS = ("Limpieza", "Revisi[oó]n", "Ajuste", "Lubricaci[oó]n", "Cambio", "Verificaci[oó]n",
               "Medici[oó]n", "Desarme", "Armado", "Sujeci[oó]n", "Calibraci[oó]n", "Inspecci[oó]n",
               "Pruebas?", "Mantenimiento", "Descalcificaci[oó]n", "Desinfecci[oó]n")
PATRON_SUSTANTIVOS = "|".join(SUSTANTIVOS)
DIVISOR_FINO = re.compile(rf"(?=\b(?:{PATRON_SUSTANTIVOS})\b)")
DIVISOR_GRUESO = re.compile(r"[\n.;]+")
SUSTANTIVO_INICIAL = re.compile(rf"^({PATRON_SUSTANTIVOS})\b\s*(.*)$", re.I)

MAPA_IMPERATIVO = {
    # sustantivo normalizado -> (imperativo de tú, artículo, nombre canónico para texto_informe)
    # El nombre canónico va SIEMPRE en singular: la fuente trae "Pruebas" (plural) y "Prueba"
    # (singular) para el mismo caso, y "la pruebas" (artículo singular + sustantivo plural) es
    # un descuerdo de número que un texto_informe no puede llevar.
    "limpieza": ("Limpia", "la", "limpieza"),
    "revision": ("Revisa", "la", "revisión"),
    "ajuste": ("Ajusta", "el", "ajuste"),
    "lubricacion": ("Lubrica", "la", "lubricación"),
    "cambio": ("Cambia", "el", "cambio"),
    "verificacion": ("Verifica", "la", "verificación"),
    "medicion": ("Mide", "la", "medición"),
    "desarme": ("Desarma", "el", "desarme"),
    "armado": ("Arma", "el", "armado"),
    "sujecion": ("Ajusta la sujeción de", "la", "sujeción"),
    "calibracion": ("Calibra", "la", "calibración"),
    "inspeccion": ("Inspecciona", "la", "inspección"),
    "prueba": ("Prueba", "la", "prueba"),
    "pruebas": ("Prueba", "la", "prueba"),
}
# Los 3 sustantivos que el patrón de corte usa pero el mapa del documento no cubre
# (ver docstring del módulo, punto 2). Aparte del mapa de arriba para poder marcarlos.
MAPA_IMPERATIVO_EXTENDIDO = {
    "mantenimiento": ("Realiza mantenimiento de", "el", "mantenimiento"),
    "descalcificacion": ("Descalcifica", "la", "descalcificación"),
    "desinfeccion": ("Desinfecta", "la", "desinfección"),
}


def atomizar_preventivo(texto: str) -> list[str]:
    atomos = []
    for grueso in DIVISOR_GRUESO.split(texto or ""):
        grueso = grueso.strip()
        if not grueso:
            continue
        for fino in DIVISOR_FINO.split(grueso):
            fino = fino.strip(" ,:-")
            if fino:
                atomos.append(fino)
    return atomos


def frasear_correctivo(texto: str) -> list[str]:
    frases = []
    for f in DIVISOR_GRUESO.split(texto or ""):
        f = f.strip(" ,:-")
        if f and len(f) >= 4:
            frases.append(f)
    return frases


def redactar(atomo_original: str) -> tuple[str, str, bool]:
    """(instruccion, texto_informe, mapa_extendido) a partir de UN átomo/frase original."""
    m = SUSTANTIVO_INICIAL.match(atomo_original.strip())
    base = atomo_original.strip()
    if not m:
        return base, f"Se realizó {base[:1].lower()}{base[1:]}", False
    sust, resto = m.group(1), m.group(2).strip(" ,:-")
    clave = normalizar_min(sust)
    extendido = clave in MAPA_IMPERATIVO_EXTENDIDO
    imperativo, articulo, nombre = MAPA_IMPERATIVO.get(clave) or MAPA_IMPERATIVO_EXTENDIDO.get(clave, (None, None, None))
    if imperativo is None:  # no debería pasar: los 16 sustantivos del patrón están en uno de los dos mapas
        return base, f"Se realizó {base[:1].lower()}{base[1:]}", False
    resto_min = resto.lower()
    instruccion = f"{imperativo} {resto_min}".strip()
    texto_informe = f"Se realizó {articulo} {nombre} {resto_min}".strip()
    return instruccion, texto_informe, extendido


# --------------------------------------------------------------------------------------
#  Base de datos local (industec_ots, la de la estación — no Hostinger)
# --------------------------------------------------------------------------------------

def conectar():
    env = cargar_env(("DB_HOST", "DB_PORT", "DB_USER", "DB_PASSWORD", "DB_NAME"))
    return mysql.connector.connect(host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
                                    password=env["DB_PASSWORD"], database=env["DB_NAME"])


def leer_preventivo(cnx) -> list[tuple[str, str, str]]:
    cur = cnx.cursor()
    cur.execute("""
        SELECT e.id_industec, e.equipo, e.descripcion
          FROM ot_equipos e JOIN ots o ON o.id_industec = e.id_industec
         WHERE o.modulo = 'PREVENTIVO' AND o.en_cuarentena = 0 AND o.correlativo < 90000
           AND e.descripcion IS NOT NULL AND TRIM(e.descripcion) <> ''
    """)
    return cur.fetchall()


def leer_correctivo(cnx) -> list[tuple[str, str]]:
    cur = cnx.cursor()
    cur.execute("""
        SELECT o.id_industec, o.actividades
          FROM ots o
         WHERE o.modulo = 'CORRECTIVO' AND o.en_cuarentena = 0 AND o.correlativo < 90000
           AND o.actividades IS NOT NULL AND TRIM(o.actividades) <> ''
    """)
    return cur.fetchall()


def leer_primer_equipo(cnx) -> dict[str, str]:
    """{id_industec: equipo} del equipo orden=0 — el primero de la orden (ver docstring, punto 1)."""
    cur = cnx.cursor()
    cur.execute("SELECT id_industec, equipo FROM ot_equipos WHERE orden = 0")
    return {r[0]: r[1] for r in cur.fetchall()}


# --------------------------------------------------------------------------------------
#  Minería
# --------------------------------------------------------------------------------------

def minar_preventivo(filas_prev, familias):
    grupos = collections.defaultdict(lambda: {"conteo": 0, "posiciones": [], "ejemplos": []})
    equipos_por_familia = collections.Counter()
    sin_familia = []
    for id_industec, equipo, descripcion in filas_prev:
        familia = resolver_familia(equipo, familias)
        if familia is None:
            sin_familia.append((id_industec, equipo, descripcion[:150]))
            continue
        equipos_por_familia[familia] += 1
        atomos = atomizar_preventivo(descripcion)
        n = len(atomos)
        for i, atomo in enumerate(atomos):
            clave = normalizar_min(atomo)
            if not clave or len(clave) < 4:
                continue
            pos = (i / (n - 1)) if n > 1 else 0.0
            g = grupos[(familia, clave)]
            g["conteo"] += 1
            g["posiciones"].append(pos)
            if len(g["ejemplos"]) < 3 and atomo not in g["ejemplos"]:
                g["ejemplos"].append(atomo)

    resultado = []
    for (familia, clave), g in grupos.items():
        total = equipos_por_familia[familia]
        umbral = max(3, 0.05 * total)
        if g["conteo"] < umbral:
            continue
        resultado.append({
            "familia": familia, "texto_norm": clave, "conteo": g["conteo"],
            "pos_media": sum(g["posiciones"]) / len(g["posiciones"]), "ejemplos": g["ejemplos"],
        })
    resultado.sort(key=lambda r: (r["familia"], r["pos_media"]))
    return resultado, equipos_por_familia, sin_familia


def minar_correctivo(filas_cor, primer_equipo, familias):
    grupos = collections.defaultdict(lambda: {"conteo": 0, "ejemplos": []})
    ordenes_por_familia = collections.Counter()
    sin_familia = []
    for id_industec, actividades in filas_cor:
        equipo0 = primer_equipo.get(id_industec)
        familia = resolver_familia(equipo0, familias) if equipo0 else None
        if familia is None:
            sin_familia.append((id_industec, equipo0, (actividades or "")[:150]))
            continue
        ordenes_por_familia[familia] += 1
        for frase in frasear_correctivo(actividades):
            clave = normalizar_min(frase)
            if not clave or len(clave) < 4:
                continue
            g = grupos[(familia, clave)]
            g["conteo"] += 1
            if len(g["ejemplos"]) < 3 and frase not in g["ejemplos"]:
                g["ejemplos"].append(frase)

    por_familia = collections.defaultdict(list)
    for (familia, clave), g in grupos.items():
        por_familia[familia].append({"familia": familia, "texto_norm": clave,
                                      "conteo": g["conteo"], "ejemplos": g["ejemplos"]})
    resultado = []
    for familia, items in por_familia.items():
        items.sort(key=lambda x: (-x["conteo"], x["texto_norm"]))
        resultado.extend(items[:25])
    resultado.sort(key=lambda r: (r["familia"], -r["conteo"]))
    return resultado, ordenes_por_familia, sin_familia


PASO_SEGURIDAD = "Desconecta el equipo de la energía y cierra el gas antes de abrirlo."


def armar_filas(minado: list[dict], modo: str, diagnosticos, con_diagnostico: bool) -> list[dict]:
    """Une el paso de seguridad + las actividades minadas, con código, paso e instrucción."""
    prefijo_modo = "PRE" if modo == "PREVENTIVO" else "COR"
    por_familia = collections.defaultdict(list)
    for item in minado:
        por_familia[item["familia"]].append(item)

    filas = []
    for familia in sorted(por_familia):
        items = por_familia[familia]
        pref_fam = prefijo_familia(familia, diagnosticos)
        paso = 1
        filas.append({
            "codigo": f"{prefijo_modo}-{pref_fam}-{paso:03d}", "familia": familia, "modo": modo,
            "diagnostico_codigo": None, "paso": paso, "frecuencia": None,
            "fuente": "práctica estándar de seguridad, no sale del histórico",
            "ejemplos": [], "instruccion": PASO_SEGURIDAD,
            "texto_informe": "Se verificó la desconexión de energía y el cierre de gas antes de intervenir el equipo.",
            "mapa_extendido": False,
        })
        for item in items:
            paso += 1
            instruccion, texto_informe, extendido = redactar(item["ejemplos"][0] if item["ejemplos"] else item["texto_norm"])
            diag = diagnostico_asociado(familia, item["texto_norm"], diagnosticos) if con_diagnostico else None
            filas.append({
                "codigo": f"{prefijo_modo}-{pref_fam}-{paso:03d}", "familia": familia, "modo": modo,
                "diagnostico_codigo": diag, "paso": paso, "frecuencia": item["conteo"],
                "fuente": "histórico" if modo == "CORRECTIVO" else "histórico (posición media {:.2f})".format(item.get("pos_media", 0)),
                "ejemplos": item["ejemplos"], "instruccion": instruccion, "texto_informe": texto_informe,
                "mapa_extendido": extendido,
            })
    return filas


# --------------------------------------------------------------------------------------
#  Salida: Excel (Preventivo/Correctivo) + JSON para cargar (12b, tras D3)
# --------------------------------------------------------------------------------------

ENCABEZADOS = ["Código propuesto", "Familia", "Modo", "Diagnóstico (código)", "Paso", "Frecuencia",
               "Fuente", "Ejemplo original 1", "Ejemplo original 2", "Ejemplo original 3",
               "Instrucción", "Texto de informe", "Ayuda", "Unidad de medición", "Mínimo", "Máximo",
               "DECISION"]


def escribir_hoja(wb, nombre, filas):
    ws = wb.create_sheet(nombre)
    for j, v in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=1, column=j, value=v)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", start_color="FF2F75B5")
    for i, f in enumerate(filas, start=2):
        ejemplos = (f["ejemplos"] + ["", "", ""])[:3]
        ayuda = ("Sustantivo sin imperativo en el mapa del documento: extendido por el agente "
                 "con el mismo criterio morfológico. Revisar primero.") if f["mapa_extendido"] else ""
        valores = [f["codigo"], f["familia"], f["modo"], f["diagnostico_codigo"] or "", f["paso"],
                   f["frecuencia"] if f["frecuencia"] is not None else "", f["fuente"],
                   ejemplos[0], ejemplos[1], ejemplos[2], f["instruccion"], f["texto_informe"],
                   ayuda, "", "", "", ""]
        for j, v in enumerate(valores, start=1):
            ws.cell(row=i, column=j, value=v)
    anchos = [16, 20, 12, 18, 6, 10, 26, 40, 40, 40, 46, 50, 40, 12, 8, 8, 30]
    for col, w in zip("ABCDEFGHIJKLMNOPQ", anchos):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws


def escribir_resumen(wb, prev_por_fam, cor_por_fam, sin_fam_prev, sin_fam_cor, familias):
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = "T2.28.12a — Biblioteca de actividades, minería del histórico (solo lectura)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Generado {datetime.now():%Y-%m-%d %H:%M}"
    enc = ["Familia", "Equipos preventivo (con descripción)", "Órdenes correctivo (con actividades)"]
    for j, v in enumerate(enc, start=1):
        c = ws.cell(row=4, column=j, value=v)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", start_color="FF2F75B5")
    for i, (fam, _patron, _orden) in enumerate(familias, start=1):
        ws.cell(row=4 + i, column=1, value=fam)
        ws.cell(row=4 + i, column=2, value=prev_por_fam.get(fam, 0))
        ws.cell(row=4 + i, column=3, value=cor_por_fam.get(fam, 0))
    fila = 4 + len(familias) + 2
    ws.cell(row=fila, column=1, value="Sin familia resuelta (preventivo)").font = Font(bold=True)
    ws.cell(row=fila, column=2, value=len(sin_fam_prev))
    ws.cell(row=fila + 1, column=1, value="Sin familia resuelta (correctivo, por el equipo orden=0)").font = Font(bold=True)
    ws.cell(row=fila + 1, column=2, value=len(sin_fam_cor))
    for col, w in zip("ABC", (24, 34, 34)):
        ws.column_dimensions[col].width = w


def escribir_sin_familia(wb, sin_fam_prev, sin_fam_cor):
    ws = wb.create_sheet("Sin familia")
    ws.append(["Modo", "id_industec", "equipo (texto real)", "texto (primeros 150 car.)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", start_color="FFC00000")
    for id_industec, equipo, texto in sin_fam_prev[:500]:
        ws.append(["PREVENTIVO", id_industec, equipo, texto])
    for id_industec, equipo, texto in sin_fam_cor[:500]:
        ws.append(["CORRECTIVO", id_industec, equipo, texto])
    for col, w in zip("ABCD", (14, 34, 30, 70)):
        ws.column_dimensions[col].width = w


def filas_a_json(filas: list[dict]) -> list[dict]:
    return [{
        "codigo": f["codigo"], "familia": f["familia"], "modo": f["modo"],
        "diagnostico_codigo": f["diagnostico_codigo"], "paso": f["paso"],
        "instruccion": f["instruccion"], "texto_informe": f["texto_informe"],
        "ayuda": None, "medicion_unidad": None, "medicion_min": None, "medicion_max": None,
        "obligatoria": 1, "estado": "BORRADOR", "origen": "SEMILLA",
        "frecuencia_historico": f["frecuencia"], "fuente": f["fuente"],
        "ejemplos_originales": f["ejemplos"], "mapa_extendido": f["mapa_extendido"],
    } for f in filas]


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--analizar", action="store_true",
                     help="Único modo de esta fase: mina el histórico y escribe el borrador (solo lectura).")
    a = ap.parse_args()
    if not a.analizar:
        raise SystemExit("Esta fase de T2.28.12a solo implementa --analizar (Fase 1, solo lectura). "
                          "La carga (12b) y la pantalla (12c) son de otra conversación, tras la puerta D3.")

    familias = cargar_familias()
    diagnosticos = cargar_diagnosticos()
    print(f"familias_equipo: {len(familias)} (de {SQL_SEMILLA.relative_to(RAIZ)})")
    print(f"diagnosticos: {len(diagnosticos)} (mismo archivo)")

    cnx = conectar()
    try:
        filas_prev = leer_preventivo(cnx)
        filas_cor = leer_correctivo(cnx)
        primer_equipo = leer_primer_equipo(cnx)
    finally:
        cnx.close()
    print(f"ot_equipos.descripcion de PREVENTIVO (en_cuarentena=0, correlativo<90000, no vacías): {len(filas_prev)}")
    print(f"ots.actividades de CORRECTIVO (mismos filtros): {len(filas_cor)}")

    minado_prev, equipos_prev_por_fam, sin_fam_prev = minar_preventivo(filas_prev, familias)
    minado_cor, ordenes_cor_por_fam, sin_fam_cor = minar_correctivo(filas_cor, primer_equipo, familias)

    filas_prev_final = armar_filas(minado_prev, "PREVENTIVO", diagnosticos, con_diagnostico=False)
    filas_cor_final = armar_filas(minado_cor, "CORRECTIVO", diagnosticos, con_diagnostico=True)

    SALIDAS.mkdir(parents=True, exist_ok=True)
    destino_xlsx = SALIDAS / "BIBLIOTECA DE ACTIVIDADES - BORRADOR (generado agente).xlsx"
    destino_json = SALIDAS / "BIBLIOTECA DE ACTIVIDADES - BORRADOR (generado agente).json"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    escribir_hoja(wb, "Preventivo", filas_prev_final)
    escribir_hoja(wb, "Correctivo", filas_cor_final)
    escribir_sin_familia(wb, sin_fam_prev, sin_fam_cor)
    escribir_resumen(wb, equipos_prev_por_fam, ordenes_cor_por_fam, sin_fam_prev, sin_fam_cor, familias)
    wb.save(destino_xlsx)

    payload = {
        "generado_en": datetime.now().isoformat(timespec="seconds"),
        "generado_por": "t2_28_actividades_semilla.py --analizar",
        "advertencia": "BORRADOR sin cargar. La carga real (12b) requiere la puerta D3 y "
                        "actividades_cargar_cli.php --ejecutar; este JSON no se ha aplicado a ninguna base.",
        "actividades": filas_a_json(filas_prev_final) + filas_a_json(filas_cor_final),
    }
    destino_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # --- Informe por consola ---------------------------------------------------------
    print("\n=== Conteo por familia (preventivo: equipos con descripción · correctivo: órdenes con actividades) ===")
    for fam, _patron, _orden in familias:
        pv, cv = equipos_prev_por_fam.get(fam, 0), ordenes_cor_por_fam.get(fam, 0)
        n_prev_act = sum(1 for f in filas_prev_final if f["familia"] == fam and f["frecuencia"] is not None)
        n_cor_act = sum(1 for f in filas_cor_final if f["familia"] == fam and f["frecuencia"] is not None)
        print(f"  {fam:<22} equipos_prev={pv:<5} actividades_prev={n_prev_act:<4} "
              f"ordenes_cor={cv:<5} actividades_cor={n_cor_act:<4}")
    print(f"\nSin familia resuelta — preventivo: {len(sin_fam_prev)} de {len(filas_prev)} "
          f"({100 * len(sin_fam_prev) / max(1, len(filas_prev)):.1f}%)")
    print(f"Sin familia resuelta — correctivo: {len(sin_fam_cor)} de {len(filas_cor)} "
          f"({100 * len(sin_fam_cor) / max(1, len(filas_cor)):.1f}%)")
    ext_prev = sum(1 for f in filas_prev_final if f["mapa_extendido"])
    ext_cor = sum(1 for f in filas_cor_final if f["mapa_extendido"])
    print(f"Filas con mapa imperativo EXTENDIDO por el agente (Mantenimiento/Descalcificación/"
          f"Desinfección, fuera del mapa del documento): {ext_prev} preventivo, {ext_cor} correctivo")
    print(f"\nFilas totales -> Preventivo: {len(filas_prev_final)} (incl. {sum(1 for f in familias if any(x['familia']==f[0] for x in filas_prev_final))} pasos de seguridad)")
    print(f"Filas totales -> Correctivo: {len(filas_cor_final)}")
    print(f"\n  {destino_xlsx}")
    print(f"  {destino_json}")


if __name__ == "__main__":
    main()
