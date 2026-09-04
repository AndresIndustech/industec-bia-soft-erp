"""Compara celda a celda el plan generado por el agente contra el plan real
llevado a mano, siguiendo el criterio de aceptacion del plan (T1.11): >=95%
identicas antes de que la administracion deje de generarlo manualmente.

La comparacion es por CONTENIDO DE FILA (matcheando por '# OT', que es el
aviso SAP -- la clave de negocio real), no por posicion de fila, ya que el
orden exacto de filas puede diferir sin que eso sea un error.
"""
import openpyxl
from pathlib import Path
from openpyxl.utils import get_column_letter

REAL = Path(r"D:\RESPALDOS\_ORIGEN_DRIVE\GESTION DE OTS INDUSTEC\2026\PLANES SEMANALES\PLAN DE TRABAJO _ ZONA UIO\PLAN SEGUIMIENTO OTS UIO _ SEPTIEMBRE.xlsx")
GENERADO = Path(r"D:\INDUSTECH IA\SALIDAS IA\MANTENIMIENTO\MANTENIMIENTOS CORRECTIVOS\ZONA UIO\PLANES SEMANALES\PLAN SEGUIMIENTO OTS UIO _ SEPTIEMBRE (generado agente).xlsx")

COLS = ["#", "# OT", "TECNICO EVALUACION", "TECNICO CIERRE", "LOCAL",
        "FECHA DE INICIO", "EQUIPO", "MARCA", "TRABAJO REALIZADO EVALUACION",
        "REPUESTO", "ESTATUS SAP", "PRESUPUESTO", "TRABAJO REALIZADO CIERRE",
        "OBSERVACIONES", "ESTATUS DEL EQUIPO", "#OT INDUSTEC EVALUACION",
        "FECHA EVALUACION", "#OT INDUSTEC CIERRE", "FECHA CIERRE",
        "REQUERIMIENTO A TIEMPO", "CALIFICACION SATISFACCIÓN", "ESTADO"]


def leer_filas(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Hoja1"]
    filas = {}
    for r in range(2, ws.max_row + 1):
        fila = {}
        for j, nombre in enumerate(COLS, start=1):
            v = ws.cell(r, j).value
            if v is not None:
                v = str(v).strip()
            fila[nombre] = v
        aviso = fila.get("# OT")
        if aviso:
            filas[str(aviso)] = fila
    return filas


def normalizar(v):
    if v is None:
        return ""
    s = str(v).strip()
    return s


def main():
    reales = leer_filas(REAL)
    generadas = leer_filas(GENERADO)

    print(f"Filas en el real: {len(reales)}  Filas en lo generado: {len(generadas)}")

    solo_en_real = set(reales) - set(generadas)
    solo_en_generado = set(generadas) - set(reales)
    en_ambos = set(reales) & set(generadas)

    print(f"Avisos solo en el REAL (el agente no los genero): {len(solo_en_real)}")
    print(f"Avisos solo en lo GENERADO (el agente agrego de mas): {len(solo_en_generado)}")
    print(f"Avisos en ambos (comparables celda a celda): {len(en_ambos)}")

    total_celdas = 0
    celdas_iguales = 0
    diffs_por_columna = {}
    for aviso in en_ambos:
        fr, fg = reales[aviso], generadas[aviso]
        for col in COLS:
            if col in ("#",):  # el numero de fila no es contenido de negocio
                continue
            total_celdas += 1
            vr, vg = normalizar(fr.get(col)), normalizar(fg.get(col))
            if vr == vg:
                celdas_iguales += 1
            else:
                diffs_por_columna.setdefault(col, []).append((aviso, vr[:40], vg[:40]))

    pct = (celdas_iguales / total_celdas * 100) if total_celdas else 0
    print(f"\nCeldas comparadas (solo avisos en ambos): {total_celdas}")
    print(f"Celdas identicas: {celdas_iguales} ({pct:.1f}%)")
    print(f"\n=== Diferencias por columna ===")
    for col, difs in sorted(diffs_por_columna.items(), key=lambda x: -len(x[1])):
        print(f"  {col}: {len(difs)} diferencias")
        for aviso, vr, vg in difs[:3]:
            print(f"      aviso {aviso}: real={vr!r}  generado={vg!r}")

    print(f"\nAvisos solo en REAL (muestra): {list(solo_en_real)[:10]}")
    print(f"Avisos solo en GENERADO (muestra): {list(solo_en_generado)[:10]}")


if __name__ == "__main__":
    main()
