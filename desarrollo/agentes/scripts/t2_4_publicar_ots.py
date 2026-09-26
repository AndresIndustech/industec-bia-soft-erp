"""
T2.4.4 - Publicacion del catalogo de OTs para INDUSTEC.

Deja en SALIDAS IA\\OTS un indice buscable de todas las ordenes de trabajo, con
la ruta exacta del PDF de cada una. Es lo que libera a la administracion de la
clasificacion manual: en vez de abrir carpetas hasta dar con un documento, se
filtra por local, zona, aviso o fecha y la columna ARCHIVO dice donde esta.

QUE PRODUCE, todo con la marca "(generado agente)" para que nunca se confunda
con un archivo llevado a mano (I-3, no se sobrescribe nada de la administracion):

  CATALOGO OTS (generado agente).xlsx
      Una hoja por zona mas una de resumen. Toda orden activa, con su local
      resuelto contra el maestro, su aviso, el estatus que declara SAP y la
      ruta del PDF. Encabezado congelado y autofiltro puesto.

  NOVEDADES {fecha} (generado agente).xlsx
      Solo lo que entro desde la publicacion anterior. Es el archivo que la
      administracion mira a diario; el catalogo completo es para buscar.

  LEEME.md
      De donde sale cada columna y que significa cada cifra.

POR QUE SE PUBLICA EL ESTATUS DE SAP Y NO EL DEL SISTEMA:
`estatus_general` de SAP es el unico criterio de cerrado del proyecto (I-10).
El `estado_ot` que trae el PDF es lo que el tecnico marco en el formulario, y
tomarlo como criterio de backlog fue el error que sobre-conto el pendiente 8
veces en T1.11. Aqui se publican los dos, en columnas distintas y con nombres
que no se prestan a confusion.

USO:
    .venv/Scripts/python.exe scripts/t2_4_publicar_ots.py
    .venv/Scripts/python.exe scripts/t2_4_publicar_ots.py --desde 2026-09-01
"""
import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import mysql.connector
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comun import termino  # noqa: E402  (diccionario unico, vocabulario.json)

ENV_PATH = Path(r"D:\INDUSTECH IA\desarrollo\agentes\config\.env")
SALIDA = Path(r"D:\INDUSTECH IA\SALIDAS IA\OTS")
MARCA = SALIDA / ".ultima_publicacion.json"

AZUL = "1F4E79"
GRIS = "F2F2F2"

COLUMNAS = [
    ("OT INDUSTEC", "id_industec", 30),
    ("ZONA", "zona", 8),
    ("TIPO", "modulo", 12),
    ("LOCAL", "local_codigo", 10),
    ("NOMBRE DEL LOCAL", "_local_nombre", 30),
    ("CADENA", "_cadena", 20),
    ("# OT (AVISO SAP)", "aviso", 15),
    ("ESTATUS SAP", "_estatus_sap", 14),
    ("FECHA DE ATENCION", "fecha_atencion", 14),
    ("TECNICO", "tecnico_nombre", 24),
    ("EQUIPO (SAP)", "_equipo_sap", 30),
    ("FASE", "fase", 12),
    ("DIA", "dia_intervencion", 6),
    ("MARCO EL TECNICO", "estado_ot", 16),
    ("A TIEMPO", "atiempo", 9),
    ("CALIFICACION", "satisfaccion", 12),
    ("FOTOS", "fotos_cantidad", 7),
    ("FIRMA", "_firma", 7),
    ("ARCHIVO", "ruta_pdf", 70),
]


def cargar_env():
    env = {}
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


def leer_ots(cnx, desde=None):
    """Une la orden con el maestro de locales y con el catalogo SAP.

    LEFT JOIN en los dos, deliberadamente: una orden sin aviso SAP es un caso de
    negocio legitimo (el tecnico ya estaba en sitio y el aviso todavia no
    existia, regla del cliente del 2026-09-04). Con INNER JOIN esas ordenes
    desaparecerian del catalogo sin que nadie se entere.
    """
    sql = """
        SELECT o.id_industec, o.zona, o.modulo, o.local_codigo, o.aviso, o.fase,
               o.dia_intervencion, o.fecha_atencion, o.tecnico_nombre,
               o.estado_ot, o.atiempo, o.satisfaccion, o.fotos_cantidad,
               o.firma_presente, o.ruta_pdf, o.creado_en,
               l.nombre  AS _local_nombre,
               l.cadena  AS _cadena,
               s.estatus_general     AS _estatus_sap,
               s.equipo_denominacion AS _equipo_sap
        FROM ots o
        LEFT JOIN locales    l ON l.local_codigo = o.local_codigo
        LEFT JOIN avisos_sap s ON s.aviso        = o.aviso
        WHERE o.en_cuarentena = 0
    """
    params = []
    if desde:
        sql += " AND o.creado_en >= %s"
        params.append(desde)
    sql += " ORDER BY o.zona, o.modulo, o.correlativo"
    cur = cnx.cursor(dictionary=True)
    cur.execute(sql, params)
    filas = cur.fetchall()
    cur.close()
    for f in filas:
        f["_firma"] = "Si" if f.get("firma_presente") else "No"
        # Si no hay dato, se dice: celda vacia, nunca un cero ni un "N/A" que
        # despues alguien lea como si fuera un valor medido (I-6).
        if not f.get("aviso"):
            f["aviso"] = None
        if not f.get("_estatus_sap"):
            f["_estatus_sap"] = "(sin aviso en SAP)" if not f.get("aviso") else "(no consta)"
    return filas


def _hoja(wb, titulo, filas):
    ws = wb.create_sheet(titulo[:31])
    ws.append([c[0] for c in COLUMNAS])
    for i, (_, _, ancho) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    col_fecha = next(i for i, c in enumerate(COLUMNAS, start=1)
                     if c[1] == "fecha_atencion")
    for f in filas:
        ws.append([f.get(clave) for _, clave, _ in COLUMNAS])
        # Sin formato explicito Excel muestra "2025-09-29 00:00:00" en una
        # columna que es una fecha, no un instante. La administracion filtra
        # por fecha todos los dias: la hora en cero solo estorba.
        ws.cell(ws.max_row, col_fecha).number_format = "YYYY-MM-DD"

    # Encabezado fijo y autofiltro: sin esto, un catalogo de 7.000 filas es
    # ilegible y la administracion vuelve a abrir carpetas a mano, que es
    # justo lo que este archivo viene a evitar.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(ws.max_row, 2)}"
    return ws


def _resumen(wb, filas):
    ws = wb.create_sheet("RESUMEN", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 60

    def titulo(t):
        ws.append([t])
        c = ws.cell(ws.max_row, 1)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")

    # Lo que se cataloga son OT INDUSTEC, el documento que emite el tecnico; «orden» es el
    # trabajo que pide KFC (vocabulario.json). Los encabezados de columna «# OT (AVISO SAP)»
    # y «ESTATUS SAP» siguen el plan de Isabel y el export de KFC.
    ws.append(["Catalogo de OT INDUSTEC"])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append([f"Generado el {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])

    titulo("Por zona y tipo")
    ws.append(["ZONA / TIPO", "OT INDUSTEC", ""])
    for (z, m), n in sorted(Counter((f["zona"], f["modulo"]) for f in filas).items()):
        ws.append([f"{z} - {m}", n, ""])
    ws.append(["TOTAL", len(filas), ""])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    ws.append([])

    titulo("Estado segun SAP")
    ws.append(["ESTATUS", "OT INDUSTEC", "COMO LEERLO"])
    # Las claves son los valores de SAP (ABIERTO/TRATAMIENTO/CERRADO, contrato) y los dos
    # rotulos que pone leer_ots(); las notas nombran los estados con el diccionario unico.
    notas = {
        "CERRADO": f"La orden esta {termino('CERRADA_SAP')}. Es el unico criterio de cierre del proyecto.",
        "TRATAMIENTO": "La orden esta en tratamiento en SAP.",
        "ABIERTO": "La orden esta abierta en SAP.",
        "(sin aviso en SAP)": f"{termino('SIN_AVISO_SAP')[:1].upper() + termino('SIN_AVISO_SAP')[1:]}: "
                              "es normal en la operacion, no un error.",
        "(no consta)": "La OT INDUSTEC trae aviso SAP, pero ese aviso no aparece en el export de SAP disponible.",
    }
    for est, n in Counter(f["_estatus_sap"] for f in filas).most_common():
        ws.append([est, n, notas.get(est, "")])
    ws.append([])

    titulo("Sin PDF localizable")
    sin = [f for f in filas if not f.get("ruta_pdf")]
    ws.append(["OT INDUSTEC cuyo PDF no esta en el arbol", len(sin),
               "Si es mayor que cero, revisar antes de purgar nada del servidor."])
    return ws


def escribir(filas, destino, con_resumen=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if con_resumen:
        _resumen(wb, filas)
    for zona in ("UIO", "LARB", "CNLJ", "OTRA"):
        de_zona = [f for f in filas if f["zona"] == zona]
        if de_zona:
            _hoja(wb, zona, de_zona)
    if not wb.sheetnames:
        _hoja(wb, "SIN DATOS", [])
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def main():
    ap = argparse.ArgumentParser(description="Publica el catalogo de OTs para INDUSTEC")
    ap.add_argument("--desde", help="fecha ISO para el archivo de novedades "
                                    "(por defecto, la ultima publicacion)")
    args = ap.parse_args()

    env = cargar_env()
    cnx = mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"])

    todas = leer_ots(cnx)
    print(f"OT INDUSTEC activas en la base: {len(todas)}")

    catalogo = SALIDA / "CATALOGO OTS (generado agente).xlsx"
    escribir(todas, catalogo)
    print(f"  catalogo -> {catalogo.name}")

    desde = args.desde
    if not desde and MARCA.exists():
        desde = json.loads(MARCA.read_text(encoding="utf-8")).get("hasta")
    novedades = leer_ots(cnx, desde) if desde else []
    if desde:
        sello = datetime.now().strftime("%Y-%m-%d")
        arch = SALIDA / f"NOVEDADES {sello} (generado agente).xlsx"
        escribir(novedades, arch)
        print(f"  novedades desde {desde}: {len(novedades)} OT INDUSTEC -> {arch.name}")
    else:
        print("  (primera publicacion: no hay novedades que separar)")

    MARCA.write_text(json.dumps({
        "hasta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "ordenes_publicadas": len(todas),
    }, indent=2), encoding="utf-8")

    sin_pdf = sum(1 for f in todas if not f.get("ruta_pdf"))
    sin_aviso = sum(1 for f in todas if not f.get("aviso"))
    (SALIDA / "LEEME.md").write_text(f"""# Catálogo de OT INDUSTEC

Generado el **{datetime.now():%Y-%m-%d %H:%M}** desde la base del proyecto.
Todo archivo aquí lleva la marca *(generado agente)*: ninguno reemplaza ni
sobrescribe un archivo llevado a mano por la administración.

## Qué hay

| Archivo | Para qué |
|---|---|
| `CATALOGO OTS (generado agente).xlsx` | **Buscar.** Las {len(todas):,} OT INDUSTEC, una hoja por zona. La columna `ARCHIVO` dice dónde está el PDF |
| `NOVEDADES {{fecha}} (generado agente).xlsx` | **Revisar a diario.** Solo lo que entró desde la publicación anterior |
| `NORMALIZACION_*.csv` | Lo que no se pudo clasificar solo, con el motivo escrito |

## Cómo se usa

El encabezado está congelado y el autofiltro puesto. Filtra por `LOCAL`,
`ZONA`, `# OT (AVISO SAP)` o `FECHA DE ATENCION`, copia la ruta de `ARCHIVO` y
pégala en el explorador. **Ya no hace falta abrir carpetas hasta encontrar un
documento.**

## Dos columnas de estado, y no significan lo mismo

| Columna | Qué es |
|---|---|
| `ESTATUS SAP` | Lo que declara SAP. **Es el único criterio de cierre.** Para backlog, SLA o reportes a KFC, esta |
| `MARCÓ EL TÉCNICO` | Lo que el técnico eligió en el formulario. Es información complementaria, nunca criterio |

Tomar la señal del sistema como criterio de cierre fue lo que sobre-contó el
pendiente 8 veces en la Fase 1. Por eso van separadas y con estos nombres.

## Cifras de esta publicación

| | |
|---|---|
| OT INDUSTEC publicadas | {len(todas):,} |
| Sin número de aviso SAP | {sin_aviso:,} — **no es un error**: la OT INDUSTEC nace sin aviso SAP cuando el técnico ya estaba en sitio |
| Sin PDF localizable | {sin_pdf:,} |

Cuando `Sin PDF localizable` sea mayor que cero, hay que resolverlo **antes** de
purgar nada del servidor de Hostinger.
""", encoding="utf-8")

    print(f"\nPublicado en {SALIDA}")
    print(f"  sin aviso SAP    : {sin_aviso}")
    print(f"  sin PDF ubicable : {sin_pdf}")
    cnx.close()


if __name__ == "__main__":
    main()
