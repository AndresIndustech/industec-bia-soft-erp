"""
T2.27.5 — El pedido de kits de mantenimiento preventivo a la bodega de KFC.

QUÉ ES
Antes de cada vuelta de preventivo la administración escribe a Edgar Armero
(bodega KFC) y a Miguel Vásquez en el hilo «SOLICITUD DE KITS MTO PREVENTIVO
_UIO» con una tabla pegada en el cuerpo: #, LOCAL, UBICACIÓN, CIUDAD, FECHA
INGRESO N (siete pedidos del 6-ago al 21-sep-2026). Si el kit no llega, el
técnico va al local y no puede hacer el mantenimiento.

POR QUÉ ES UNA PROPUESTA Y NO EL PEDIDO
El cronograma del servidor (ingresos_preventivos) conserva las fechas del plan
anual: la ejecución real va dos a tres semanas detrás y nadie reprogramó
ninguna (0 reagendas en el año). El 14-sep ella pidió kits para G001–G006 del
18-sep al 2-oct; la tabla los tiene para el 2 al 15-sep. Y el 21-sep pidió kit
para K062 y V074, que la tabla da por ejecutados (estado CUMPLIDO). Con esa
tabla sola, el pedido saldría con locales y fechas equivocados.

Por eso el generador arma la lista de lo que el cronograma dice que falta
—ingresos atrasados sin ejecutar y los de las próximas dos semanas—, en el
formato exacto del correo, con la FECHA DE INGRESO en blanco: la fecha real la
pone quien programa la visita. Y lista las contradicciones que encontró.

Los estados del preventivo se nombran con el diccionario único (vocabulario.json,
vía comun.termino): EJECUTADO · PENDIENTE · ATRASADO, la leyenda del jefe técnico.
«vencida» es solo el plazo de 48 h de una solicitud; un ingreso que pasó su fecha
está «atrasado». Los valores de la base (CUMPLIDO, CANCELADO) se comparan tal cual.

Uso:
    .venv/Scripts/python.exe scripts/t2_27_kits_preventivo.py                 # las tres zonas, próximas 2 semanas
    .venv/Scripts/python.exe scripts/t2_27_kits_preventivo.py --zona UIO --dias 21
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import t2_27_fuentes as F  # noqa: E402
from comun import corto, de_estado, termino  # noqa: E402

ORDINAL = {1: "1", 2: "2", 3: "3", 4: "4"}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--zona", choices=["UIO", "LARB", "CNLJ"])
    ap.add_argument("--dias", type=int, default=14, help="horizonte hacia adelante (por omisión 14)")
    ap.add_argument("--hoy", help="fecha de referencia AAAA-MM-DD (por omisión, hoy)")
    a = ap.parse_args()
    hoy = dt.date.fromisoformat(a.hoy) if a.hoy else dt.date.today()
    hasta = hoy + dt.timedelta(days=a.dias)
    filas = F.consultar_servidor(
        "SELECT local_codigo, zona, anio, numero, plan_original_inicio, plan_vigente_inicio, plan_vigente_fin, "
        "real_inicio, estado, kit_estado FROM ingresos_preventivos WHERE anio = ? ORDER BY plan_vigente_inicio",
        [hoy.year])
    cnx = F.conectar()
    cur = cnx.cursor(dictionary=True)
    cur.execute("SELECT local_codigo, nombre, cadena, zona FROM locales")
    maestro = {r["local_codigo"]: r for r in cur.fetchall()}
    cnx.close()

    pendientes, contradicciones = [], []
    por_local = {}
    for f in filas:
        por_local.setdefault(f["local_codigo"], []).append(f)
    for f in filas:
        if a.zona and f["zona"] != a.zona:
            continue
        if f["estado"] in ("CUMPLIDO", "CANCELADO") or not f["plan_vigente_inicio"]:
            continue
        ini = dt.date.fromisoformat(f["plan_vigente_inicio"])
        if ini > hasta:
            continue
        # Solo el ingreso más antiguo sin cumplir de cada local: nadie pide el kit del 4 antes del 3.
        anteriores = [x for x in por_local[f["local_codigo"]] if x["numero"] < f["numero"] and x["estado"] not in ("CUMPLIDO", "CANCELADO")]
        if anteriores:
            contradicciones.append((f["local_codigo"], f"el ingreso {f['numero']} está pendiente, pero también el "
                                    f"{', '.join(str(x['numero']) for x in anteriores)}, que es anterior"))
            continue
        m = maestro.get(f["local_codigo"], {})
        pendientes.append({"local": f["local_codigo"], "zona": f["zona"], "numero": f["numero"],
                           "ubicacion": m.get("nombre") or "", "cadena": m.get("cadena") or "",
                           "plan": ini, "vencido_dias": (hoy - ini).days if ini < hoy else 0, "kit": f["kit_estado"]})
    reprogramados = sum(1 for f in filas if f["plan_original_inicio"] != f["plan_vigente_inicio"])
    if reprogramados == 0:
        contradicciones.append(("(todos)", f"ninguno de los {len(filas)} ingresos de {hoy.year} fue reprogramado: "
                                "las fechas son las del plan anual, no las de la ejecución real"))

    salida = F.SALIDAS / f"{hoy:%Y-%m-%d}"
    salida.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PEDIDO DE KITS"
    ws["A1"] = (f"Ingresos de preventivo {termino('PREV_PENDIENTE', 2)} según el cronograma — "
                f"{termino('ATRASADO', 2)} y próximos {a.dias} días (al {hoy:%d/%m/%Y})")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Propuesta del agente: la FECHA DE INGRESO real la pone quien programa la visita. El cronograma no está reprogramado."
    ws["A2"].font = Font(italic=True, color="FFC00000")
    enc = ["#", "LOCAL", "UBICACIÓN", "CADENA", "ZONA", "INGRESO N.º", "FECHA PLAN (cronograma)", "DÍAS DE ATRASO", "KIT (según el sistema)", "FECHA INGRESO"]
    for j, v in enumerate(enc, start=1):
        c = ws.cell(row=4, column=j, value=v)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", start_color="FF2F75B5")
    for i, p in enumerate(sorted(pendientes, key=lambda p: (p["zona"], p["plan"])), start=1):
        # La zona con su rótulo corto del diccionario (CNLJ se lee CUENCA-LOJA), como en B.IA.
        zona = corto(de_estado(p["zona"], "zona")) if p["zona"] in ("UIO", "LARB", "CNLJ", "OTRA") else p["zona"]
        for j, v in enumerate([i, p["local"], p["ubicacion"], p["cadena"], zona, p["numero"], p["plan"],
                               p["vencido_dias"] or None, p["kit"], None], start=1):
            c = ws.cell(row=4 + i, column=j, value=v)
            if isinstance(v, dt.date):
                c.number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFGHIJ", (4, 9, 34, 18, 7, 9, 14, 10, 14, 16)):
        ws.column_dimensions[col].width = w
    h = wb.create_sheet("CONTRADICCIONES")
    h.append(["Local", "Qué no cuadra en el cronograma"])
    for c in contradicciones:
        h.append(list(c))
    h.column_dimensions["A"].width, h.column_dimensions["B"].width = 12, 110
    destino = salida / f"PEDIDO KITS PREVENTIVO {a.zona or 'TRES ZONAS'} (generado agente).xlsx"
    wb.save(destino)

    # La tabla para pegar en el correo, con las columnas que ella usa.
    html = ['<table border="1" cellspacing="0" cellpadding="4" style="border-collapse:collapse;font-family:Calibri;font-size:11pt">',
            "<tr style='background:#2F75B5;color:#fff'><th>#</th><th>LOCAL</th><th>UBICACIÓN</th><th>CIUDAD</th><th>FECHA INGRESO</th></tr>"]
    for i, p in enumerate(sorted(pendientes, key=lambda p: (p["zona"], p["plan"])), start=1):
        html.append(f"<tr><td>{i}</td><td>{escape(p['local'])}</td><td>{escape(p['ubicacion'])}</td><td></td><td></td></tr>")
    html.append("</table>")
    (salida / f"PEDIDO KITS PREVENTIVO {a.zona or 'TRES ZONAS'} (tabla para el correo, generado agente).html").write_text("\n".join(html), encoding="utf-8")
    print(f"Ingresos {termino('PREV_PENDIENTE', 2)} según el cronograma: {len(pendientes)} "
          f"({termino('ATRASADO', 2)}: {sum(1 for p in pendientes if p['vencido_dias'])})")
    print(f"Contradicciones: {len(contradicciones)}")
    for c in contradicciones[:8]:
        print("  -", c[0], c[1])
    print(f"  {destino}")


if __name__ == "__main__":
    main()
