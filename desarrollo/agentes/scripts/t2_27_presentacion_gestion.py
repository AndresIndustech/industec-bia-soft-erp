"""
T2.27.4 — «RESUMEN DE GESTIÓN INDUSTEC», la presentación por zona para las reuniones con Grupo KFC.

QUÉ ES
Para las reuniones de indicadores (Ronald Valero, 15-abr-2026: «revisión de
gestión de mantenimiento; cronogramas preventivos; revisión de equipos críticos
paralizados … favor tener lista la información») la administración arma a mano
un PowerPoint por zona: órdenes por zona, por franquicia, por local y por
equipo, mes a mes, y cierra con CONCLUSIONES y RECOMENDACIONES. El de
octubre–noviembre de 2025 se tecleó entero: trae un local que no existe
(«K131») y cifras que no cuadran con ninguna fuente.

QUÉ HACE
La misma secuencia de diapositivas, con tablas y gráficos nativos (editables en
PowerPoint), contada desde el archivo canónico de órdenes (la base local:
órdenes activas, reales, por fecha de atención). La cadena sale del maestro de
locales, nunca del prefijo del código. Suma una diapositiva de INDICADORES DE
SERVICIO que antes no existía: tiempo de la notificación a la primera visita.

LO QUE NO INVENTA
Las causas técnicas de las conclusiones («termocupla y kit de encendido», «cal
en las máquinas de hielo») vienen del conocimiento de los jefes técnicos, no de
un conteo. El generador deja esas viñetas marcadas en rojo para que las
complete una persona antes de presentar.

Uso:
    .venv/Scripts/python.exe scripts/t2_27_presentacion_gestion.py --desde 2026-07 --hasta 2026-08
    .venv/Scripts/python.exe scripts/t2_27_presentacion_gestion.py --desde 2026-08 --hasta 2026-09 --zona LARB --sap <REPORTE SEMANA NN>.xlsx
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import re
import statistics
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import t2_27_fuentes as F  # noqa: E402

AZUL = RGBColor(0x17, 0x36, 0x5D)
AZUL2 = RGBColor(0x2F, 0x75, 0xB5)
GRIS = RGBColor(0x59, 0x59, 0x59)
ROJO = RGBColor(0xC0, 0x00, 0x00)
NOMBRE_ZONA = {"UIO": "ZONA QUITO", "LARB": "ZONA LARB", "CNLJ": "ZONA CUENCA LOJA"}
MES_ES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Familias de equipo para contar como las cuenta la administración («FREIDORAS», «HORNOS»...).
# El técnico escribe «Freidora de presión», «Freidora abierta», «Freidora HP6»: son la misma
# familia en la diapositiva. Lo que no calza con ninguna queda con su nombre en mayúsculas.
FAMILIAS = [
    ("FREIDORA", "FREIDORAS"), ("HORNO", "HORNOS"), ("HIELO", "MÁQUINA DE HIELO"), ("HELAD", "MÁQUINA DE HELADOS"),
    ("HOLDING", "HOLDING CABINET"), ("CUARTO FR", "CUARTO FRÍO"), ("CONGELADOR", "CONGELADOR"),
    ("REFRIGERADOR HORIZONTAL", "REFRIGERADOR HORIZONTAL"), ("MESA REFRIG", "MESA REFRIGERADA"), ("BASE REFRIG", "MESA REFRIGERADA"),
    ("REFRIGERA", "REFRIGERADOR"), ("NEVERA", "REFRIGERADOR"), ("EXHIBID", "EXHIBIDOR"), ("VITRINA", "EXHIBIDOR"),
    ("LICUADORA", "LICUADORA"), ("CAFE", "MÁQUINA DE CAFÉ"), ("CAFÉ", "MÁQUINA DE CAFÉ"), ("TOST", "TOSTADORA"),
    ("PLANCHA", "PLANCHA"), ("PARRILLA", "PARRILLA"), ("COCINA", "COCINA"), ("LAVAVAJ", "LAVAVAJILLAS"),
    ("APAN", "MESA DE APANADO"), ("EXTRACTOR", "EXTRACTOR"), ("CAMPANA", "CAMPANA"), ("MICROONDAS", "MICROONDAS"),
    ("DISPENSADOR", "DISPENSADOR"), ("LAMPARA", "LÁMPARA DE CALOR"), ("LÁMPARA", "LÁMPARA DE CALOR"),
]


def familia(nombre: str | None) -> str:
    t = (nombre or "").upper().strip()
    if not t:
        return "SIN DATO"
    for clave, fam in FAMILIAS:
        if clave in t:
            return fam
    return re.sub(r"\s+", " ", t)[:30]


def meses_entre(desde: str, hasta: str) -> list[str]:
    a, b = dt.date.fromisoformat(desde + "-01"), dt.date.fromisoformat(hasta + "-01")
    out = []
    while a <= b:
        out.append(f"{a:%Y-%m}")
        a = (a.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    return out


def etiqueta_mes(m: str) -> str:
    return MES_ES[int(m[5:]) - 1]


# ---------------------------------------------------------------------------

def datos(meses: list[str], zonas: list[str], sap_ruta: Path | None) -> dict:
    cnx = F.conectar()
    cur = cnx.cursor(dictionary=True)
    cur.execute("""
        SELECT o.id_industec, o.zona, o.modulo, o.aviso, o.local_codigo, o.fecha_atencion, o.estado_ot,
               DATE_FORMAT(o.fecha_atencion, '%Y-%m') AS mes, l.cadena, l.nombre AS local_nombre, e.equipo
          FROM ots o
          LEFT JOIN locales l ON l.local_codigo = o.local_codigo
          LEFT JOIN ot_equipos e ON e.id_industec = o.id_industec AND e.orden = 0
         WHERE o.en_cuarentena = 0 AND o.correlativo < 90000
           AND DATE_FORMAT(o.fecha_atencion, '%Y-%m') BETWEEN %s AND %s""", (meses[0], meses[-1]))
    filas = [r for r in cur.fetchall() if r["zona"] in zonas]
    cur.execute("SELECT local_codigo, zona, cadena, nombre FROM locales WHERE activo = 1 ORDER BY zona, cadena, local_codigo")
    locales = cur.fetchall()
    # Tiempos de respuesta: del aviso en SAP a la primera orden. Con el Excel semanal de KFC si
    # se da (cubre todo 2026 al día); si no, con avisos_sap de la base (cubre hasta agosto).
    notif = {}
    if sap_ruta:
        for a, r in F.leer_sap_semanal(sap_ruta).items():
            if r["fecha_notificacion"]:
                notif[a] = r["fecha_notificacion"]
        fuente_t = sap_ruta.name
    else:
        cur.execute("SELECT aviso, fecha_notificacion FROM avisos_sap WHERE fecha_notificacion IS NOT NULL")
        notif = {str(r["aviso"]): r["fecha_notificacion"] for r in cur.fetchall()}
        cur.execute("SELECT MIN(fecha_notificacion) mn, MAX(fecha_notificacion) mx FROM avisos_sap")
        c = cur.fetchone()
        fuente_t = f"avisos_sap de la base ({c['mn']} a {c['mx']})"
    cnx.close()
    primera = {}
    for r in sorted(filas, key=lambda r: r["fecha_atencion"]):
        if r["modulo"] == "CORRECTIVO" and r["aviso"]:
            primera.setdefault(str(r["aviso"]).strip(), (r["fecha_atencion"], r["zona"], r["mes"]))
    tiempos = collections.defaultdict(list)
    for a, (f, z, m) in primera.items():
        if a in notif and f >= notif[a]:
            tiempos[(z, m)].append((f - notif[a]).days)
    return dict(filas=filas, locales=locales, tiempos=tiempos, fuente_t=fuente_t)


# ---------------------------------------------------------------------------
#  Diapositivas
# ---------------------------------------------------------------------------

def _titulo(s, texto: str, sub: str = ""):
    barra = s.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(0.9))
    barra.fill.solid()
    barra.fill.fore_color.rgb = AZUL
    barra.line.fill.background()
    tf = barra.text_frame
    tf.text = texto
    p = tf.paragraphs[0]
    p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(24), True, RGBColor(255, 255, 255), "Calibri"
    if sub:
        cuadro = s.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.4))
        cuadro.text_frame.text = sub
        q = cuadro.text_frame.paragraphs[0]
        q.font.size, q.font.italic, q.font.color.rgb = Pt(12), True, GRIS


def _pie(s, texto: str):
    c = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35))
    c.text_frame.text = texto
    c.text_frame.paragraphs[0].font.size = Pt(9)
    c.text_frame.paragraphs[0].font.color.rgb = GRIS


def _tabla(s, filas: list[list], x, y, ancho, alto, anchos=None, tam=11):
    t = s.shapes.add_table(len(filas), len(filas[0]), x, y, ancho, alto).table
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = "" if v is None else (f"{v:,}".replace(",", ".") if isinstance(v, int) else str(v))
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(tam)
            p.font.bold = i == 0 or str(fila[0]).upper() == "TOTAL"
            if j > 0:
                p.alignment = PP_ALIGN.CENTER
    if anchos:
        for j, w in enumerate(anchos):
            t.columns[j].width = w
    return t


def _grafico(s, tipo, categorias, series: dict, x, y, ancho, alto, titulo=""):
    cd = CategoryChartData()
    cd.categories = categorias
    for nombre, valores in series.items():
        cd.add_series(nombre, valores)
    g = s.shapes.add_chart(tipo, x, y, ancho, alto, cd).chart
    g.has_legend = len(series) > 1
    if g.has_legend:
        g.legend.position = XL_LEGEND_POSITION.BOTTOM
        g.legend.include_in_layout = False
    if titulo:
        g.has_title = True
        g.chart_title.text_frame.text = titulo
        g.chart_title.text_frame.paragraphs[0].font.size = Pt(12)
    plot = g.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9)
    # Sin esto PowerPoint usa 18 pt en los ejes y los nombres de cadena se tragan el gráfico.
    g.category_axis.tick_labels.font.size = Pt(10)
    g.value_axis.tick_labels.font.size = Pt(9)
    if g.has_legend:
        g.legend.font.size = Pt(10)
    return g


def _vinetas(s, lineas: list[tuple[str, bool]], y=Inches(1.4)):
    c = s.shapes.add_textbox(Inches(0.6), y, Inches(12.1), Inches(5.5))
    tf = c.text_frame
    tf.word_wrap = True
    for i, (t, pendiente) in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + t
        p.font.size = Pt(16)
        p.space_after = Pt(10)
        if pendiente:
            p.font.color.rgb = ROJO
            p.font.italic = True


def construir(meses: list[str], zonas: list[str], d: dict, destino: Path) -> dict:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vacia = prs.slide_layouts[6]
    periodo = " – ".join(etiqueta_mes(m) for m in (meses[0], meses[-1])) if len(meses) > 1 else etiqueta_mes(meses[0])
    periodo += f" {meses[-1][:4]}"
    fuente = ("Fuente: archivo canónico de órdenes de INDUSTEC (órdenes activas por fecha de atención); "
              "la cadena sale del maestro de locales. Generado por el agente.")
    cifras = {}

    # Portada
    s = prs.slides.add_slide(vacia)
    fondo = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = AZUL
    fondo.line.fill.background()
    c = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2.4)).text_frame
    c.text = "RESUMEN DE GESTIÓN INDUSTEC"
    c.paragraphs[0].font.size, c.paragraphs[0].font.bold, c.paragraphs[0].font.color.rgb = Pt(40), True, RGBColor(255, 255, 255)
    p = c.add_paragraph()
    p.text = (" · ".join(NOMBRE_ZONA[z] for z in zonas)) + f"  |  {periodo}"
    p.font.size, p.font.color.rgb = Pt(20), RGBColor(0xDD, 0xE6, 0xF0)

    filas = d["filas"]
    por_zona_mes = collections.Counter((r["zona"], r["mes"]) for r in filas)

    # Resumen por zona
    s = prs.slides.add_slide(vacia)
    _titulo(s, "RESUMEN DE OTS POR ZONA", f"Órdenes de trabajo emitidas por mes · {periodo}")
    tabla = [["ZONA"] + [etiqueta_mes(m) for m in meses] + ["TOTAL"]]
    for z in zonas:
        v = [por_zona_mes[(z, m)] for m in meses]
        tabla.append([NOMBRE_ZONA[z]] + v + [sum(v)])
    tabla.append(["TOTAL"] + [sum(por_zona_mes[(z, m)] for z in zonas) for m in meses] + [sum(por_zona_mes[(z, m)] for z in zonas for m in meses)])
    _tabla(s, tabla, Inches(0.5), Inches(1.5), Inches(5.6), Inches(0.4) * len(tabla))
    _grafico(s, XL_CHART_TYPE.COLUMN_CLUSTERED, [NOMBRE_ZONA[z] for z in zonas],
             {etiqueta_mes(m): [por_zona_mes[(z, m)] for z in zonas] for m in meses},
             Inches(6.4), Inches(1.4), Inches(6.5), Inches(5.3), "OTS POR ZONA")
    _pie(s, fuente)
    cifras["por_zona"] = {z: [por_zona_mes[(z, m)] for m in meses] for z in zonas}

    for z in zonas:
        fz = [r for r in filas if r["zona"] == z]
        s = prs.slides.add_slide(vacia)
        fondo = s.shapes.add_shape(1, 0, Inches(3), prs.slide_width, Inches(1.5))
        fondo.fill.solid()
        fondo.fill.fore_color.rgb = AZUL2
        fondo.line.fill.background()
        fondo.text_frame.text = NOMBRE_ZONA[z]
        fondo.text_frame.paragraphs[0].font.size = Pt(36)
        fondo.text_frame.paragraphs[0].font.bold = True

        # Por franquicia
        cad = collections.Counter((r["cadena"] or "SIN CADENA EN EL MAESTRO", r["mes"]) for r in fz)
        cadenas = sorted({k for k, _ in cad}, key=lambda k: -sum(cad[(k, m)] for m in meses))
        s = prs.slides.add_slide(vacia)
        _titulo(s, "RESUMEN DE OTS POR FRANQUICIA", f"{NOMBRE_ZONA[z]} · {periodo}")
        t = [["FRANQUICIA"] + [etiqueta_mes(m) for m in meses] + ["TOTAL"]]
        for k in cadenas:
            v = [cad[(k, m)] for m in meses]
            t.append([k] + v + [sum(v)])
        t.append(["TOTAL"] + [sum(cad[(k, m)] for k in cadenas) for m in meses] + [len(fz)])
        _tabla(s, t, Inches(0.5), Inches(1.5), Inches(5.8), Inches(0.36) * len(t), tam=10)
        _grafico(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cadenas, {etiqueta_mes(m): [cad[(k, m)] for k in cadenas] for m in meses},
                 Inches(6.6), Inches(1.4), Inches(6.3), Inches(5.3), "OTS POR FRANQUICIA")
        _pie(s, fuente)

        # Por local (los 15 con más órdenes)
        loc = collections.Counter((r["local_codigo"], r["mes"]) for r in fz)
        nombres = {r["local_codigo"]: r["local_nombre"] or "" for r in fz}
        top = sorted({k for k, _ in loc}, key=lambda k: -sum(loc[(k, m)] for m in meses))[:15]
        s = prs.slides.add_slide(vacia)
        _titulo(s, "RESUMEN DE OTS POR LOCAL", f"{NOMBRE_ZONA[z]} · los 15 locales con más órdenes · {periodo}")
        t = [["LOCAL", "NOMBRE"] + [etiqueta_mes(m) for m in meses] + ["TOTAL"]]
        for k in top:
            v = [loc[(k, m)] for m in meses]
            t.append([k, (nombres.get(k) or "")[:34]] + v + [sum(v)])
        _tabla(s, t, Inches(0.5), Inches(1.45), Inches(12.3), Inches(0.33) * len(t), tam=10,
               anchos=[Inches(1.3), Inches(5.0)] + [Inches(1.6)] * (len(meses) + 1))
        _pie(s, fuente)

        # Por equipo (familias), un gráfico por el periodo completo
        eq = collections.Counter(familia(r["equipo"]) for r in fz if r["modulo"] == "CORRECTIVO")
        top_eq = [k for k, _ in eq.most_common(10)]
        eq_mes = collections.Counter((familia(r["equipo"]), r["mes"]) for r in fz if r["modulo"] == "CORRECTIVO")
        s = prs.slides.add_slide(vacia)
        _titulo(s, "RESUMEN DE OTS POR EQUIPOS", f"{NOMBRE_ZONA[z]} · correctivos, los 10 tipos de equipo con más órdenes · {periodo}")
        t = [["EQUIPO"] + [etiqueta_mes(m) for m in meses] + ["TOTAL"]]
        for k in top_eq:
            v = [eq_mes[(k, m)] for m in meses]
            t.append([k] + v + [sum(v)])
        _tabla(s, t, Inches(0.5), Inches(1.5), Inches(5.8), Inches(0.36) * len(t), tam=10)
        _grafico(s, XL_CHART_TYPE.BAR_CLUSTERED, list(reversed(top_eq)), {"Correctivos": [eq[k] for k in reversed(top_eq)]},
                 Inches(6.6), Inches(1.4), Inches(6.3), Inches(5.3), "OTS POR EQUIPO")
        _pie(s, fuente)

        # Indicadores de servicio (nuevo)
        s = prs.slides.add_slide(vacia)
        _titulo(s, "INDICADORES DE SERVICIO", f"{NOMBRE_ZONA[z]} · de la notificación del aviso en SAP a la primera visita de INDUSTEC")
        t = [["MES", "CORRECTIVOS", "PREVENTIVOS", "AVISOS CON VISITA", "MEDIANA (DÍAS)", "VISITA ≤ 1 DÍA", "VISITA ≤ 2 DÍAS"]]
        for m in meses:
            dias = d["tiempos"].get((z, m), [])
            corr = sum(1 for r in fz if r["mes"] == m and r["modulo"] == "CORRECTIVO")
            prev = sum(1 for r in fz if r["mes"] == m and r["modulo"] == "PREVENTIVO")
            t.append([etiqueta_mes(m), corr, prev, len(dias),
                      f"{statistics.median(dias):g}".replace(".", ",") if dias else "sin dato",
                      f"{sum(1 for x in dias if x <= 1) / len(dias) * 100:.0f} %" if dias else "sin dato",
                      f"{sum(1 for x in dias if x <= 2) / len(dias) * 100:.0f} %" if dias else "sin dato"])
        _tabla(s, t, Inches(0.5), Inches(1.6), Inches(12.3), Inches(0.45) * len(t), tam=12)
        _pie(s, f"Tiempos: {d['fuente_t']}. Días calendario. Un mes «sin dato» es un mes fuera de la cobertura de esa fuente, no un cero.")

        # Conclusiones y recomendaciones
        tot_cad = collections.Counter(r["cadena"] or "SIN CADENA" for r in fz)
        top3 = [k for k, _ in tot_cad.most_common(3)]
        eq3 = [k.lower() for k in top_eq[:3]]
        mv = [por_zona_mes[(z, m)] for m in meses]
        cambio = ""
        if len(mv) >= 2 and mv[-2]:
            dif = (mv[-1] - mv[-2]) / mv[-2] * 100
            cambio = (f"Las órdenes de {etiqueta_mes(meses[-1]).lower()} {'subieron' if dif > 0 else 'bajaron'} "
                      f"{abs(dif):.0f} % frente a {etiqueta_mes(meses[-2]).lower()} ({mv[-2]} → {mv[-1]}).")
        s = prs.slides.add_slide(vacia)
        _titulo(s, "CONCLUSIONES", NOMBRE_ZONA[z])
        lineas = [(f"Las franquicias con mayor cantidad de incidencias en la zona son {', '.join(top3[:-1])} y {top3[-1]}." if len(top3) > 1
                   else f"La franquicia con mayor cantidad de incidencias en la zona es {top3[0]}." if top3 else "Sin órdenes en el periodo.", False)]
        if cambio:
            lineas.append((cambio, False))
        # La plantilla de 2025 decía que el sistema «se alinea con los tiempos esperados por KFC».
        # El generador no afirma eso: pone el dato medido y que lo juzgue quien lee.
        todos = [x for m in meses for x in d["tiempos"].get((z, m), [])]
        if todos:
            lineas.append((f"De {len(todos)} avisos correctivos con visita en el periodo, el "
                           f"{sum(1 for x in todos if x <= 1) / len(todos) * 100:.0f} % recibió la primera visita el mismo día "
                           f"o al día siguiente de la notificación (mediana: {statistics.median(todos):g} días).", False))
        if eq3:
            lineas.append((f"Los tipos de equipos que presentan mayor cantidad de incidencias son: {', '.join(eq3)}.", False))
            for e in eq3:
                lineas.append((f"[COMPLETAR ANTES DE PRESENTAR — jefe técnico de zona] Causa técnica principal de las incidencias en {e}.", True))
        _vinetas(s, lineas)
        s = prs.slides.add_slide(vacia)
        _titulo(s, "RECOMENDACIONES", NOMBRE_ZONA[z])
        _vinetas(s, [
            (f"Para las franquicias con mayor número de incidencias ({', '.join(top3)}): mantener un análisis mensual de recurrencia "
             "de fallas por tienda para priorizar recursos y acciones correctivas específicas en cada una.", False),
            ("Sobre los equipos con mayor incidencia: implementar un monitoreo específico por tipo de equipo para identificar patrones, "
             "fallas repetitivas y tiempos promedio de resolución.", False),
            ("Respecto al Sistema de Gestión de OTs: continuar con el seguimiento diario y capacitar periódicamente a los técnicos y "
             "administradores para garantizar el uso adecuado del sistema y evitar retrasos en el flujo de información.", False),
            ("[COMPLETAR ANTES DE PRESENTAR] Recomendación específica por cada causa técnica identificada en las conclusiones.", True),
        ])

        # Lista de locales
        loc_z = [l for l in d["locales"] if l["zona"] == z]
        s = prs.slides.add_slide(vacia)
        _titulo(s, f"LISTA DE LOCALES {NOMBRE_ZONA[z].replace('ZONA ', '')}", f"TOTAL LOCALES: {len(loc_z)} · maestro de locales")
        porc = collections.defaultdict(list)
        for l in loc_z:
            porc[l["cadena"]].append(f"{l['local_codigo']}  {(l['nombre'] or '')[:30]}")
        col, y, x = 0, Inches(1.4), Inches(0.4)
        cuadro = s.shapes.add_textbox(x, y, Inches(12.5), Inches(5.6)).text_frame
        cuadro.word_wrap = True
        primero = True
        for cadena, ls in sorted(porc.items(), key=lambda kv: -len(kv[1])):
            p = cuadro.paragraphs[0] if primero else cuadro.add_paragraph()
            primero = False
            p.text = f"{cadena} ({len(ls)}): " + " · ".join(ls)
            p.font.size = Pt(10)
            p.space_after = Pt(6)
        cifras[z] = {"total": len(fz), "top_franquicias": top3, "top_equipos": top_eq[:3], "locales": len(loc_z)}

    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destino)
    cifras["diapositivas"] = len(prs.slides)
    return cifras


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    hoy = dt.date.today()
    ant = (hoy.replace(day=1) - dt.timedelta(days=1))
    ap.add_argument("--desde", default=f"{(ant.replace(day=1) - dt.timedelta(days=1)):%Y-%m}", help="primer mes AAAA-MM (por omisión, dos meses atrás)")
    ap.add_argument("--hasta", default=f"{ant:%Y-%m}", help="último mes AAAA-MM (por omisión, el mes pasado)")
    ap.add_argument("--zona", choices=["UIO", "LARB", "CNLJ"], help="una sola zona; por omisión, las tres")
    ap.add_argument("--sap", help="Excel semanal de KFC para los tiempos (cubre todo el año al día)")
    a = ap.parse_args()
    meses = meses_entre(a.desde, a.hasta)
    zonas = [a.zona] if a.zona else ["UIO", "LARB", "CNLJ"]
    d = datos(meses, zonas, Path(a.sap) if a.sap else None)
    nombre = f"RESUMEN GESTION INDUSTEC {'GENERAL' if not a.zona else NOMBRE_ZONA[a.zona]} {meses[0]} a {meses[-1]} (generado agente).pptx"
    destino = F.SALIDAS / "PRESENTACIONES" / nombre
    cifras = construir(meses, zonas, d, destino)
    print(f"Presentación: {destino}\n  {cifras['diapositivas']} diapositivas · órdenes por zona y mes: {cifras['por_zona']}")
    for z in zonas:
        print(f"  {z}: {cifras[z]}")


if __name__ == "__main__":
    main()
