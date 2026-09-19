# -*- coding: utf-8 -*-
"""
Coteja el export de ordenes abiertas de SAP contra todo lo que tenemos.

Por que existe: el buzon no trae todos los casos. Se midio el 2026-09-10 contra
`ORDENES ABIERTAS EN SAP.xlsx`: de 62 ordenes abiertas en SAP, el buzon habia
capturado 31. Y la causa no era la que se suponia — los 213 avisos originales de
SIR de la ventana llevan servicioalcliente@industec.me en el destinatario, el
100%. Lo que pasa es que SAP registra ordenes que SIR nunca notifica por correo.
INDUSTEC igual las trabaja (27 de esas 31 ya tenian OT emitida), pero el sistema
no las ve. Por eso el export es la segunda fuente, no un control opcional.

SOLO LECTURA en todo: no escribe en ninguna base, ni local ni de Hostinger, y no
toca el Drive. Lo unico que produce es el Excel de cotejo en SALIDAS IA, a nombre
nuevo (I-4).

Uso:
    .venv/Scripts/python.exe scripts/t2_12_cotejo_sap_abiertas.py [ruta al xlsx]

Sin argumento toma el del Drive. El archivo se copia a un temporal antes de
leerlo, porque la administracion suele tenerlo abierto en Excel.
"""
import collections
import datetime
import email
import email.header
import imaplib
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import mysql.connector
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(r"D:\INDUSTECH IA\desarrollo\agentes")
ENV_PATH = BASE / "config" / ".env"
CATALOGOS = Path(r"D:\INDUSTECH IA\SALIDAS IA\OTS\catalogos")
SALIDA = Path(r"D:\INDUSTECH IA\SALIDAS IA\OTS")
XLSX_DRIVE = Path(r"G:\Mi unidad\INDUSTEC IA\ORDENES ABIERTAS EN SAP.xlsx")

# Mapeo letra -> indice 0-based, sondeado sobre el archivo real del 2026-09-10.
# La columna A no trae encabezado y es el numero de aviso.
COL_AVISO, COL_FECHA, COL_TIPO = 0, 1, 3          # A, B, D
COL_CLASE, COL_CENTRO, COL_AREA = 4, 5, 6         # E, F, G
COL_EQUIPO, COL_FECHA_REF, COL_ORDEN = 7, 8, 9    # H, I, J
COL_UBICACION, COL_MODIF = 11, 12                 # L, M
COL_MINIMA = 13                                   # hasta 'Modificado por'


def cargar_env() -> dict:
    env = {}
    for linea in ENV_PATH.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#") or "=" not in linea:
            continue
        clave, valor = linea.split("=", 1)
        env[clave.strip()] = valor.strip()
    return env


def leer_export(ruta: Path) -> list:
    """Lee el export de SAP. Aborta si cambio el formato en vez de perder columnas."""
    tmp = Path(tempfile.gettempdir()) / "_cotejo_sap.xlsx"
    shutil.copy2(ruta, tmp)          # el original suele estar abierto en Excel
    try:
        wb = openpyxl.load_workbook(tmp, data_only=True)
        ws = wb[wb.sheetnames[0]]
        if ws.max_column < COL_MINIMA:
            sys.exit("ABORTADO: cambio el formato del export de SAP. Se esperaban al "
                     f"menos {COL_MINIMA} columnas y hay {ws.max_column}")

        filas, descartadas = [], 0
        vistos = set()
        for celdas in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            v = [c.value for c in celdas]
            if all(x is None or str(x).strip() == "" for x in v):
                continue
            crudo = v[COL_AVISO]
            # openpyxl devuelve los numeros como float: str() daria '10352817.0'
            aviso = str(int(crudo)) if isinstance(crudo, float) else str(crudo).strip()
            if not aviso.isdigit():
                descartadas += 1
                continue
            if aviso in vistos:          # se deduplica aqui y se reporta, no en la base
                descartadas += 1
                continue
            vistos.add(aviso)
            fecha = v[COL_FECHA]
            if not isinstance(fecha, (datetime.date, datetime.datetime)):
                fecha = None             # solo fechas tipadas por Excel; el resto va vacio
            elif isinstance(fecha, datetime.datetime):
                fecha = fecha.date()
            filas.append({
                "aviso": aviso, "fecha": fecha,
                "tipo": (v[COL_TIPO] or "").strip(), "clase": (v[COL_CLASE] or "").strip(),
                "local": (v[COL_CENTRO] or "").strip(), "area": (v[COL_AREA] or "").strip(),
                "equipo": (v[COL_EQUIPO] or "").strip(), "orden_sap": v[COL_ORDEN] or "",
                "ubicacion": (v[COL_UBICACION] or "").strip(), "modif": (v[COL_MODIF] or "").strip(),
            })
        if descartadas:
            print(f"filas descartadas del export (sin aviso valido o repetidas): {descartadas}")
        return filas
    finally:
        # Es una copia de datos del cliente fuera de las carpetas del proyecto
        # (riesgo LOPDP): no debe sobrevivir a la corrida, ni siquiera si abortamos.
        tmp.unlink(missing_ok=True)


def cargar_catalogo_buzon() -> tuple:
    ruta = CATALOGOS / "casos_sap.json"
    if not ruta.exists():
        sys.exit(f"ABORTADO: falta {ruta}. Corre antes t2_6_imap_avisos.py")
    d = json.loads(ruta.read_text(encoding="utf-8"))
    # El catalogo se regenera cada 3 horas: todo cotejo declara la hora que uso,
    # porque con la foto de las 13:30 el conteo daba 33 y con la de las 16:44, 31.
    return {c["aviso"]: c for c in d["datos"]}, d["generado"]


def cargar_base(env: dict) -> tuple:
    cn = mysql.connector.connect(host=env["DB_HOST"], port=int(env["DB_PORT"]),
                                 user=env["DB_USER"], password=env["DB_PASSWORD"],
                                 database=env["DB_NAME"])
    cur = cn.cursor(dictionary=True)
    cur.execute("SELECT aviso, estatus_general FROM avisos_sap")
    avisos = {str(r["aviso"]): r["estatus_general"] for r in cur.fetchall()}
    cur.execute("SELECT MIN(fecha_notificacion) AS fmin, MAX(fecha_notificacion) AS fmax "
                "FROM avisos_sap")
    cobertura = cur.fetchone()
    cur.execute("SELECT aviso, id_industec FROM ots WHERE en_cuarentena=0 AND aviso IS NOT NULL")
    ots = collections.defaultdict(list)
    for r in cur.fetchall():
        ots[str(r["aviso"])].append(r["id_industec"])
    cur.execute("SELECT local_codigo, nombre, zona, cadena FROM locales")
    locales = {r["local_codigo"]: r for r in cur.fetchall()}
    cur.close()
    cn.close()
    return avisos, cobertura, ots, locales


RE_OT_ASUNTO = re.compile(r"OT-\d{3,4}-[A-Z]+\d+-(\d{8})(?:-D\d+)?-(?:UIO|LARB|CNLJ)", re.I)


def normalizar_ot(ot: str) -> str:
    """'OT-2418-K073EC-10351653-CNLJ' y 'OT-2418-K073-10351653-CNLJ' son la misma."""
    return re.sub(r"([A-Z]+\d+)EC", r"\1", (ot or "").upper())


def ots_emitidas_por_correo(env: dict, desde: str) -> dict:
    """Que avisos tienen ya una OT emitida, segun el asunto de los correos de envio.

    Hace falta porque la tabla `ots` local solo tiene lo ya ingestado al arbol
    canonico: el 2026-09-10 tenia 3 de las 33 OT que en realidad existian. El
    asunto trae el aviso ('ORDEN DE TRABAJO INDUSTEC - OT-2466-V093-10352936-CNLJ'),
    asi que basta con las cabeceras y no hay que bajar un solo cuerpo.

    SOLO LECTURA: select(readonly=True) -> EXAMINE, y BODY.PEEK.
    """
    emisor = env.get("EMISOR_OT", "reclutamiento@industec.me")
    salida = {}
    try:
        M = imaplib.IMAP4_SSL(env["IMAP_HOST"], int(env["IMAP_PORT"]))
        M.login(env["IMAP_USER"], env["IMAP_PASSWORD"])
        M.select("INBOX", readonly=True)
        # Las frases con espacios van entre comillas o el servidor corta la sesion.
        ok, d = M.search(None, "SINCE", desde, "FROM", '"%s"' % emisor)
        ids = d[0].split() if ok == "OK" and d and d[0] else []
        for i in range(0, len(ids), 25):
            lote = b",".join(ids[i:i + 25])
            ok, dd = M.fetch(lote, "(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE)])")
            if ok != "OK" or not dd:
                continue
            for parte in dd:
                if not isinstance(parte, tuple):
                    continue
                msg = email.message_from_bytes(parte[1])
                try:
                    asunto = str(email.header.make_header(
                        email.header.decode_header(msg.get("Subject", ""))))
                except Exception:
                    asunto = str(msg.get("Subject", ""))
                m = RE_OT_ASUNTO.search(asunto)
                if m:
                    salida.setdefault(m.group(1), asunto.split(" - ")[-1].strip())
        M.logout()
    except Exception as e:
        # Sin buzon el cotejo sigue sirviendo: se dice que falta esa via y no se
        # inventa nada (I-7).
        print(f"AVISO: no se pudo leer el buzon para las OT emitidas ({str(e)[:60]}).")
        print("       La columna 'OT ya emitida' quedara solo con el archivo historico.")
    return salida


SSH_HOST, SSH_PUERTO = "82.25.73.181", 65002
LLAVE = BASE / "config" / "clave_hostinger"
RUTA_REMOTA = "domains/darkviolet-armadillo-872352.hostingersite.com/public_html/ot"

# Lee casos_gestion en la base operativa. Se pasa por stdin para no dejar ningun
# archivo en el servidor, y es un SELECT: no escribe nada.
PHP_GESTION = """<?php
$cfg = require getenv('HOME') . '/%s/nucleo/config.php';
$pdo = new PDO(sprintf('mysql:host=%%s;port=%%d;dbname=%%s;charset=utf8mb4',
      $cfg['db_host'], $cfg['db_port'], $cfg['db_name']), $cfg['db_user'], $cfg['db_pass'],
      [PDO::ATTR_ERRMODE => PDO::ERRMODE_EXCEPTION, PDO::ATTR_DEFAULT_FETCH_MODE => PDO::FETCH_ASSOC]);
$q = $pdo->query("SELECT g.aviso, g.estado, g.ot_cierre, u.nombre AS tecnico
                  FROM casos_gestion g LEFT JOIN usuarios u ON u.usuario_id = g.asignado_a");
echo json_encode($q->fetchAll(), JSON_UNESCAPED_UNICODE);
""" % RUTA_REMOTA


def estado_de_gestion(env: dict) -> dict:
    """Como esta cada caso en el sistema: asignado, atendido, cerrado.

    Hace falta porque `estado_gestion` del catalogo del buzon **siempre dice
    NUEVO** — es un valor fijo del generador, no el estado real. Sin esto el
    cotejo le decia 'asignar tecnico' a un caso ya atendido que solo esperaba
    que la administradora confirmara el cierre en SAP.
    """
    import subprocess
    if not LLAVE.exists():
        print("AVISO: no hay llave SSH; el cotejo no sabra el estado de gestion.")
        return {}
    usuario = env.get("SSH_USER", "").strip()
    if not usuario:
        print("AVISO: falta SSH_USER en config/.env; sin estado de gestion.")
        return {}
    try:
        r = subprocess.run(
            ["ssh", "-i", str(LLAVE), "-p", str(SSH_PUERTO), "-o", "BatchMode=yes",
             "-o", "StrictHostKeyChecking=accept-new", f"{usuario}@{SSH_HOST}", "php"],
            input=PHP_GESTION, capture_output=True, text=True, timeout=90)
        if r.returncode != 0 or not r.stdout.strip().startswith("["):
            print(f"AVISO: no se pudo leer casos_gestion ({(r.stderr or r.stdout)[:70]}).")
            return {}
        return {str(x["aviso"]): x for x in json.loads(r.stdout)}
    except Exception as e:
        print(f"AVISO: no se pudo leer casos_gestion ({str(e)[:70]}).")
        return {}


def main():
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX_DRIVE
    if not ruta.exists():
        sys.exit(f"ABORTADO: no existe {ruta}")
    env = cargar_env()
    sap = leer_export(ruta)
    buzon, generado = cargar_catalogo_buzon()
    db_avisos, cobertura, ots, locales = cargar_base(env)
    hoy = datetime.date.today()

    # La ventana arranca una semana antes del aviso mas viejo del export, para
    # alcanzar la OT aunque se haya emitido apenas llegado el caso.
    fechas = [s["fecha"] for s in sap if s["fecha"]]
    desde = (min(fechas) - datetime.timedelta(days=7)) if fechas else (hoy - datetime.timedelta(days=30))
    emitidas = ots_emitidas_por_correo(env, desde.strftime("%d-%b-%Y"))
    print(f"OT emitidas halladas en el buzon desde {desde}: {len(emitidas)}")
    gestion = estado_de_gestion(env)

    filas = []
    for s in sap:
        av = s["aviso"]
        b = buzon.get(av)
        lm = locales.get(s["local"]) or {}
        # El archivo historico y el correo son dos vias para lo mismo; se unen sin
        # repetir, porque una OT recien emitida aun no esta ingestada al arbol.
        # Se comparan normalizadas: el arbol escribe 'OT-2418-K073EC-...' y el
        # correo 'OT-2418-K073-...', que es la misma orden con y sin sufijo.
        vistas = list(ots.get(av, []))
        if emitidas.get(av) and normalizar_ot(emitidas[av]) not in {normalizar_ot(x) for x in vistas}:
            vistas.append(emitidas[av])
        dias = (hoy - s["fecha"]).days if s["fecha"] else None
        filas.append({
            **s,
            "dias": dias,
            "local_nombre": lm.get("nombre") or (b or {}).get("local_nombre") or "",
            "zona": lm.get("zona") or (b or {}).get("zona") or "",
            "cadena": lm.get("cadena") or (b or {}).get("cadena") or "",
            "en_buzon": "SI" if b else "NO",
            "historico": ", ".join(vistas),
            "estatus_export_anterior": db_avisos.get(av) or "",
            "estado_gestion": (gestion.get(av) or {}).get("estado") or "",
            "tecnico": (gestion.get(av) or {}).get("tecnico") or "",
            "ot_cierre": (gestion.get(av) or {}).get("ot_cierre") or "",
        })

    # en_buzon="SI"/"NO" son las dos ramas de un mismo if (linea 277): siempre
    # suman len(filas), sobre la misma lista. No es una compuerta I-10 -no hay
    # fuente independiente que confirme el total- asi que no se hacia sys.exit(1)
    # por algo que nunca puede fallar: era una garantia de papel.
    en_buzon = sum(1 for f in filas if f["en_buzon"] == "SI")
    fuera = sum(1 for f in filas if f["en_buzon"] == "NO")

    print("=" * 74)
    print(f"COTEJO DE {ruta.name} AL {hoy}")
    print("=" * 74)
    print(f"catalogo del buzon usado: foto del {generado} ({len(buzon)} casos vivos)")
    print(f"base local avisos_sap   : cobertura {cobertura['fmin']} .. {cobertura['fmax']}")
    print(f"\nordenes abiertas en SAP : {len(filas)}")
    print(f"  capturadas por el buzon : {en_buzon}")
    print(f"  NO capturadas ......... : {fuera}")

    # I-12: lo que cae fuera de la cobertura del catalogo no es un error
    if cobertura["fmax"]:
        tarde = [f for f in filas if f["fecha"] and f["fecha"] > cobertura["fmax"]]
        print(f"  posteriores al corte de la base ({cobertura['fmax']}): {len(tarde)}"
              " — fuera de cobertura, no es error")

    con_ot = [f for f in filas if f["historico"]]
    print(f"\ncon OT en el archivo historico y SAP abierto: {len(con_ot)}"
          " — el trabajo se hizo, falta el cierre en SAP")

    print("\npor zona:")
    for z, n in collections.Counter(f["zona"] or "(sin resolver)" for f in filas).most_common():
        print(f"  {z:<22} {n}")

    # Que clases de trabajo trae: si viene una sola, el export esta filtrado y
    # no se puede concluir nada del resto (I-12).
    clases = collections.Counter(f["tipo"] for f in filas)
    print("\nclases de trabajo en el export:")
    for c, n in clases.most_common():
        print(f"  {c:<22} {n}")
    if len(clases) == 1:
        print("  AVISO: el export trae una sola clase de trabajo. Muy probablemente")
        print("  esta filtrado: este cotejo NO dice nada de las demas clases.")

    # El dato compartido debe coincidir, o hay que mirarlo una por una
    difs = []
    for f in filas:
        b = buzon.get(f["aviso"])
        if not b:
            continue
        csap = f["local"].upper().replace("EC", "")
        cbuz = str(b["centro_coste_sap"] or "").upper().replace("EC", "")
        if csap != cbuz and f["local"].upper() != str(b["local"] or "").upper():
            difs.append((f["aviso"], "local", f["local"], b["local"]))
        if f["fecha"] and str(f["fecha"]) != b["fecha_creacion"]:
            difs.append((f["aviso"], "fecha", str(f["fecha"]), b["fecha_creacion"]))
        if f["tipo"] != b["caso"]:
            difs.append((f["aviso"], "tipo", f["tipo"], b["caso"]))
    print(f"\ndiferencias de dato entre SAP y el buzon: {len(difs)}")
    for d in difs:
        print("  aviso %s  %s: SAP=%r  buzon=%r" % d)

    # La OT emitida deberia ser del mismo local que dice SAP
    malos = []
    for f in filas:
        for ot in f["historico"].split(", "):
            m = re.match(r"OT-\d+-([A-Z]+\d+)", ot or "")
            if m and m.group(1) != f["local"].upper().replace("EC", ""):
                malos.append((f["aviso"], f["local"], m.group(1) + "EC", ot))
    if malos:
        print(f"\nOT emitidas contra un local distinto al de SAP: {len(malos)}")
        for m in malos:
            print("  aviso %s  SAP=%s  la OT dice %s  (%s)" % m)

    # Mismo local y mismo equipo con dos avisos abiertos: posible duplicado en SAP
    pares = collections.defaultdict(list)
    for f in filas:
        if f["equipo"]:
            pares[(f["local"], f["equipo"])].append(f["aviso"])
    reps = {k: v for k, v in pares.items() if len(v) > 1}
    if reps:
        print(f"\nposibles duplicados en SAP (mismo local y equipo): {len(reps)}")
        for (loc, eq), avs in reps.items():
            print(f"  {loc:<8} {eq[:40]:<40} {', '.join(avs)}")
        print("  No se corrigen: puede ser reincidencia real. Decide la administracion.")

    escribir_excel(filas, ruta, generado)
    print(f"\n{len(filas)} filas procesadas.")


def escribir_excel(filas: list, origen: Path, generado: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ABIERTAS EN SAP"
    enc = ["AVISO", "FECHA AVISO", "DIAS", "LOCAL", "NOMBRE DEL LOCAL", "ZONA", "CADENA",
           "EQUIPO", "ORDEN SAP", "LO TENIAMOS?", "OT YA EMITIDA", "ESTADO EN EL SISTEMA",
           "TECNICO ASIGNADO", "QUE HAY QUE HACER"]
    ws.append(enc)
    for i, _ in enumerate(enc, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rojo = PatternFill("solid", fgColor="FCE4E4")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    verde = PatternFill("solid", fgColor="E2EFDA")
    for f in sorted(filas, key=lambda x: (-(x["dias"] or 0), x["aviso"])):
        # El orden importa: lo que ya se atendio no se manda a asignar, y lo que
        # nadie esta viendo va primero que cualquier tramite.
        if f["ot_cierre"]:
            accion = "Confirmar el cierre en SAP: el trabajo ya se hizo"
            relleno = verde
        elif f["historico"]:
            accion = "Verificar: hay OT emitida y SAP sigue abierto"
            relleno = amarillo if f["en_buzon"] == "NO" else verde
        elif f["estado_gestion"] == "ASIGNADO":
            accion = "En curso, con tecnico asignado"
            relleno = verde
        elif f["en_buzon"] == "NO":
            accion = "REVISAR: no entro por el buzon, nadie lo esta viendo"
            relleno = rojo
        else:
            accion = "Asignar tecnico"
            relleno = verde
        ws.append([f["aviso"], str(f["fecha"] or ""), f["dias"], f["local"], f["local_nombre"],
                   f["zona"], f["cadena"], f["equipo"], f["orden_sap"], f["en_buzon"],
                   f["historico"], f["estado_gestion"] or "sin gestion", f["tecnico"], accion])
        for col in range(1, len(enc) + 1):
            ws.cell(row=ws.max_row, column=col).fill = relleno

    for i, a in enumerate([11, 12, 6, 9, 30, 7, 16, 34, 12, 13, 30, 20, 26, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("DE DONDE SALE")
    for fila in [
        ["Generado por", "t2_12_cotejo_sap_abiertas.py"],
        ["Fecha del cotejo", str(datetime.date.today())],
        ["Export de SAP", str(origen)],
        ["Catalogo del buzon", f"casos_sap.json, foto del {generado}"],
        ["Que NO dice este archivo", "Si el export de SAP trae filtro de clase o de fecha. "
                                     "Confirmarlo con quien lo genero."],
        ["El Drive", "Solo se leyo, sobre una copia temporal. No se modifico nada."],
    ]:
        ws2.append(fila)
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 90
    for i in range(1, ws2.max_row + 1):
        ws2.cell(row=i, column=1).font = Font(bold=True)
        ws2.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

    # I-4: nombre nuevo, con marca de origen. Nunca se sobrescribe el de la administracion.
    destino = SALIDA / "ORDENES ABIERTAS EN SAP - cotejo (generado agente).xlsx"
    try:
        wb.save(destino)
    except PermissionError:
        sys.exit(f"ABORTADO: {destino.name} esta abierto en Excel. Cierralo y repite.")
    print(f"\nEscrito: {destino}")


if __name__ == "__main__":
    main()
