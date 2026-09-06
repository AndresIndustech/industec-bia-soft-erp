"""
T2.5.3 - Padron de tecnicos reconstruido desde las ordenes de trabajo.

La tabla `tecnicos` es una foto: 19 personas, todas marcadas activas, con fecha
de ingreso y SIN fecha de salida. Pero INDUSTEC tiene alta rotacion, asi que esa
foto no dice quien trabajaba cuando -- y las ordenes historicas estan firmadas
por gente que ya no esta. Eso no es un error de datos: es historia real.

Las ordenes SI llevan ese registro. Cada una tiene fecha y nombre de quien la
firmo, asi que la primera y la ultima orden de cada persona dibujan su ventana
de actividad. Este script la reconstruye y la entrega para que INDUSTEC confirme
quien sigue vigente.

POR QUE HACE FALTA ANTES DE PONER EL DESPLEGABLE DE TECNICOS:
El formulario unico ofrece una lista cerrada. Si esa lista son los 19 de la
tabla, deja fuera a gente que hoy trabaja (aparecen 2.389 ordenes firmadas por
alguien que no esta en la tabla) y mete a gente que ya se fue. Cualquiera de las
dos rompe la captura el primer dia.

COMO AGRUPA LAS 406 GRAFIAS EN PERSONAS:
Por tokens del nombre, sin tildes. Dos grafias son la misma persona si el
conjunto de tokens de una esta contenido en el de la otra Y comparten al menos
dos tokens. Eso une "Anthony Jumbo" con "ANTHONY MEDARDO JUMBO ROJANO" y
"Diego Melendrez" con "Diego Meléndrez", pero NO une a dos personas que solo
comparten el nombre de pila.

NO DECIDE QUIEN SIGUE Y QUIEN NO (I-6). Calcula dias desde la ultima orden y lo
presenta ordenado; la vigencia la declara INDUSTEC, no un umbral inventado por
un script.

USO:
    .venv/Scripts/python.exe scripts/t2_5_padron_tecnicos.py
"""
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import mysql.connector
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENV_PATH = Path(r"D:\INDUSTECH IA\desarrollo\agentes\config\.env")
SALIDA = Path(r"D:\INDUSTECH IA\SALIDAS IA\OTS")

AZUL = "1F4E79"

# Separadores con que los tecnicos escriben "fuimos dos". Aparecen 1.113 veces.
RE_SEPARA = re.compile(r"\s*[,/;]\s*|\s+y\s+", re.IGNORECASE)


def tokens(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return frozenset(t for t in re.split(r"[^A-Za-z]+", s.upper()) if len(t) > 2)


def cargar_env():
    env = {}
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


def agrupar(grafias):
    """Une las grafias que son la misma persona.

    Se ordena de mas tokens a menos y cada grafia corta se absorbe en la primera
    larga que la contenga. Exigir 2 tokens comunes es lo que impide fusionar a
    dos personas distintas que solo comparten el nombre de pila -- y eso pasa:
    hay mas de un Diego y mas de un Luis en la operacion.
    """
    orden = sorted(grafias, key=lambda g: (-len(tokens(g)), g))
    grupos = []
    for g in orden:
        t = tokens(g)
        if len(t) < 2:
            grupos.append({"semilla": t, "grafias": [g]})
            continue
        for gr in grupos:
            # Se compara contra la SEMILLA -- la grafia mas larga del grupo -- y
            # no contra una union que va creciendo. Con la union, absorber
            # "Anthony Jumbo Perez" metia PEREZ en el conjunto y a partir de ahi
            # el grupo dejaba de ser subconjunto de la nomina, que es como
            # "Anthony Jumbo" aparecia reportado como si no existiera.
            if t <= gr["semilla"]:
                gr["grafias"].append(g)
                break
        else:
            grupos.append({"semilla": t, "grafias": [g]})
    return grupos


def main():
    env = cargar_env()
    cnx = mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"])
    cur = cnx.cursor(dictionary=True)

    cur.execute("""SELECT tecnico_nombre, zona, fecha_atencion
                   FROM ots
                   WHERE en_cuarentena = 0 AND tecnico_nombre IS NOT NULL
                     AND tecnico_nombre <> '' AND fecha_atencion IS NOT NULL""")
    filas = cur.fetchall()

    cur.execute("""SELECT nombres, apellidos, fecha_ingreso, tipo_tecnico, zona_asignada
                   FROM tecnicos""")
    nomina = [{"nombre": f'{r["nombres"]} {r["apellidos"]}'.strip(),
               "tokens": tokens(f'{r["nombres"]} {r["apellidos"]}'),
               "ingreso": r["fecha_ingreso"], "tipo": r["tipo_tecnico"],
               "zona": r["zona_asignada"]} for r in cur.fetchall()]
    cur.close(); cnx.close()

    # Una orden firmada por dos personas cuenta para las dos: es trabajo que
    # ambas hicieron, y para la ventana de actividad es exactamente lo que
    # queremos saber.
    por_grafia = defaultdict(lambda: {"n": 0, "primera": None, "ultima": None,
                                      "zonas": set(), "compartidas": 0})
    for f in filas:
        partes = [p.strip() for p in RE_SEPARA.split(f["tecnico_nombre"]) if p.strip()]
        for p in partes:
            d = por_grafia[p]
            d["n"] += 1
            fa = f["fecha_atencion"]
            d["primera"] = fa if d["primera"] is None else min(d["primera"], fa)
            d["ultima"] = fa if d["ultima"] is None else max(d["ultima"], fa)
            d["zonas"].add(f["zona"])
            if len(partes) > 1:
                d["compartidas"] += 1

    grupos = agrupar(list(por_grafia))
    hoy = date.today()

    personas = []
    for gr in grupos:
        n = sum(por_grafia[g]["n"] for g in gr["grafias"])
        primera = min(por_grafia[g]["primera"] for g in gr["grafias"])
        ultima = max(por_grafia[g]["ultima"] for g in gr["grafias"])
        zonas = set().union(*(por_grafia[g]["zonas"] for g in gr["grafias"]))
        compartidas = sum(por_grafia[g]["compartidas"] for g in gr["grafias"])
        # La grafia mas frecuente es como se le conoce; la mas larga suele ser
        # la del documento. Se muestra la frecuente y se listan todas.
        principal = max(gr["grafias"], key=lambda g: por_grafia[g]["n"])
        # Una orden fechada despues de hoy es un error de digitacion del tecnico,
        # no una prediccion. Se marca en vez de corregirse: la fecha correcta
        # solo la sabe quien atendio (I-6).
        futura = ultima > hoy
        # Contra cada grafia por separado, no contra el grupo entero: basta que
        # UNA de las formas en que se le escribe encaje con el nombre legal.
        en_nomina = next((x["nombre"] for x in nomina
                          if x["tokens"] and any(tokens(g) and tokens(g) <= x["tokens"]
                                                 for g in gr["grafias"])), None)
        personas.append({
            "nombre": principal, "grafias": sorted(gr["grafias"]),
            "ordenes": n, "primera": primera, "ultima": ultima,
            "dias_sin_firmar": (hoy - ultima).days,
            "zonas": ", ".join(sorted(z for z in zonas if z)),
            "compartidas": compartidas,
            "en_nomina": en_nomina,
            "fecha_futura": futura,
        })
    personas.sort(key=lambda p: p["ultima"], reverse=True)

    # -------------------- Excel para que INDUSTEC lo marque --------------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PADRON"
    cols = [("TECNICO", 26), ("¿SIGUE EN LA EMPRESA?", 22), ("ORDENES", 10),
            ("PRIMERA ORDEN", 15), ("ULTIMA ORDEN", 15), ("DIAS SIN FIRMAR", 16),
            ("ZONAS", 16), ("ORDENES COMPARTIDAS", 20), ("EN LA TABLA DE NOMINA", 30),
            ("AVISO", 26), ("COMO APARECE ESCRITO", 70)]
    ws.append([c[0] for c in cols])
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    for p in personas:
        ws.append([p["nombre"], "", p["ordenes"], p["primera"], p["ultima"],
                   p["dias_sin_firmar"], p["zonas"], p["compartidas"],
                   p["en_nomina"] or "no está",
                   "fecha futura en una orden" if p["fecha_futura"] else "",
                   " · ".join(p["grafias"])])
        for col in (4, 5):
            ws.cell(ws.max_row, col).number_format = "YYYY-MM-DD"

    # La columna que INDUSTEC llena. Lista cerrada para que la respuesta vuelva
    # utilizable y no como texto libre -- que es justo el problema que estamos
    # resolviendo en todo lo demas.
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"SI,NO,NO SE"', allow_blank=True)
    dv.prompt = "Marca si esta persona sigue trabajando en INDUSTEC"
    dv.promptTitle = "Vigencia"
    ws.add_data_validation(dv)
    dv.add(f"B2:B{ws.max_row}")

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / "PADRON TECNICOS (generado agente).xlsx"
    wb.save(destino)

    recientes = [p for p in personas if p["dias_sin_firmar"] <= 60]
    lejanos = [p for p in personas if p["dias_sin_firmar"] > 180]
    print(f"Grafias distintas en las ordenes : {len(por_grafia)}")
    print(f"Personas tras agrupar            : {len(personas)}")
    print(f"En la tabla de nomina            : {sum(1 for p in personas if p['en_nomina'])}")
    print(f"Firmaron en los ultimos 60 dias  : {len(recientes)}")
    print(f"Sin firmar hace mas de 180 dias  : {len(lejanos)}")
    futuras = [p for p in personas if p["fecha_futura"]]
    if futuras:
        print(f"Con alguna orden en fecha FUTURA : {len(futuras)} "
              f"({', '.join(p['nombre'] for p in futuras)})")
    print(f"\n{destino}")
    print("\nLas 12 personas con actividad mas reciente:")
    for p in personas[:12]:
        marca = "" if p["en_nomina"] else "  <- no esta en la tabla de nomina"
        print(f"  {p['ultima']}  {p['ordenes']:>4} ots  {p['nombre']:<24}{marca}")


if __name__ == "__main__":
    main()
