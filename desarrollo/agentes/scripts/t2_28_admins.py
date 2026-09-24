"""
t2_28_admins.py - T2.28.3: el administrador y el correo de cada local, sembrados
desde el histórico (absorbe la acción O de §11b del plan).

POR QUÉ
El formulario va a dejar elegir el administrador y el correo del local de una
lista precargada, como ya se hace con el nombre (H-08, `locales_admin`). Esa
lista tiene que nacer de algo: las 7.410 filas de `ots` que ya traen
`admin_nombre` (y, algunas, `correo_local`) son la única fuente con 99 de 100
locales cubiertos. Pero es texto libre tecleado por técnicos en un formulario:
el mismo local trae hasta 30 variantes distintas del mismo nombre (typos,
mayúsculas sueltas, un punto de más), y también ruido real ("----", ".",
"10288697", nombres de trabajo tecleados donde iba el administrador).

Este script SOLO analiza. La carga a `locales_admin` (`--ejecutar`, con la
cifra exacta) es otra fase que requiere la aprobación de Andrés
(T2_28_OBSERVACIONES_INDUSTEC.md, T2.28.3, tabla de permisos) y a propósito
**no está implementada todavía**: no hay una sola sentencia INSERT/UPDATE en
este archivo. Lo que produce es la propuesta para que alguien la revise.

CÓMO SE ARMA LA PROPUESTA (en este orden, igual que la especificación)

1. CLAVE DE COMPARACIÓN. `admin_nombre` en mayúsculas, sin tildes, con los
   espacios colapsados a uno solo. Sirve solo para AGRUPAR: lo que se muestra
   es la forma tal como se leyó con más frecuencia dentro de cada grupo
   (`elegir_forma`), nunca la clave en mayúsculas.

2. FUSIÓN DENTRO DEL MISMO LOCAL. Dos claves del mismo local se funden si:
     a) la distancia de Levenshtein entre las claves completas es <= 2
        ("Alex Yanez" / "Alex Yanez." -> 1: un punto de más); o
     b) el conjunto de palabras de la clave más corta está contenido en el de
        la más larga y las dos empiezan con la misma primera palabra
        ("LIZ" ⊂ "LIZ TIAMARCA").
   La fusión usa unión-búsqueda (DSU) por local: cada unión que de verdad
   conecta dos componentes distintos queda como una fila en la hoja
   "Fusionados", con el motivo exacto. Es intencional no ocultar nada aquí:
   la distancia de Levenshtein <= 2 sobre nombres de 3-4 letras puede fundir
   dos personas distintas ("ANA"/"ANI"), y esa hoja es donde se atrapa antes
   de cargar nada.

3. RUIDO FUERA (se aplica sobre la forma ya fusionada):
     - la clave tiene algún dígito;
     - alguna palabra de la clave es una palabra de trabajo (REVISION, EQUIPO,
       LOCAL, KFC, MANTENIMIENTO);
     - la clave tiene menos de 3 letras (se cuentan solo las letras, así que
       "----" y "." caen aquí, no en la regla de dígitos);
     - se vio una sola vez y hace más de 12 meses (calendario, no 365 días
       fijos: `relativedelta`). Un par visto 2+ veces NUNCA es ruido por
       fecha, sin importar cuán vieja sea.

4. SIEMBRA. De lo que no es ruido, se siembra lo visto 2+ veces, o lo visto en
   los últimos 6 meses. Lo que queda -visto una sola vez, entre 6 y 12 meses
   atrás- no es "ruido" en el sentido del punto 3 pero tampoco pasa el umbral:
   igual va a la hoja de descartados, con su propio motivo, para que la
   ecuación de abajo cuadre y para que I-7 se cumpla (nada desaparece sin
   quedar escrito en algún lado).

   Por cada grupo sembrado: `veces`, `visto_ultimo`, y el correo más
   frecuente de ese grupo que NO sea `@industec...` (esas son la casilla
   genérica del maestro, no el correo real del local), con su propio
   `correo_veces`.

CONTABILIDAD QUE SE PUEDE VERIFICAR A MANO
    pares_leídos (grupos (local, clave) antes de fundir)
        = filas de "Fusionados" (una por cada unión que de verdad conectó
          dos componentes -- exactamente n-k por local, propiedad estándar de
          un DSU construido solo con uniones que tuvieron éxito)
        + filas de "Se siembra" + filas de "Descartados (motivo)"
          (un grupo final, tras fundir, cae en una sola de las dos)
    Es el criterio de aceptación de T2.28.3 y lo comprueba `main()` antes de
    escribir el Excel (aborta si no cuadra).

QUÉ PRODUCE
    SALIDAS IA\\OTS\\ADMINISTRADORES POR LOCAL (generado agente).xlsx
        hojas: "Se siembra", "Fusionados", "Descartados (motivo)", "Resumen"
    SALIDAS IA\\OTS\\t2_28_admins_siembra.json
        lo mismo que "Se siembra", para que la fase de carga no tenga que
        releer el Excel

Uso (en la estación; NUNCA toca la base, solo hace SELECT sobre `ots`):
    .venv/Scripts/python.exe scripts/t2_28_admins.py --analizar
    .venv/Scripts/python.exe scripts/t2_28_admins.py               # mismo efecto
"""
from __future__ import annotations

import argparse
import json
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import mysql.connector
import openpyxl
from dateutil.relativedelta import relativedelta
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comun import SALIDAS, cargar_env, escribir_json_atomico  # noqa: E402

XLSX = SALIDAS / "ADMINISTRADORES POR LOCAL (generado agente).xlsx"
SALIDA_JSON = SALIDAS / "t2_28_admins_siembra.json"

# Regla 3 de la especificación: palabras que delatan que en el campo del
# administrador se tecleó otra cosa (una nota de trabajo, no un nombre).
PALABRAS_TRABAJO = {"REVISION", "EQUIPO", "LOCAL", "KFC", "MANTENIMIENTO"}

AZUL = "FF2F75B5"
BLANCO = "FFFFFFFF"


# --- Normalización y comparación -----------------------------------------------
def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def clave_normalizada(raw: str) -> str:
    """Mayúsculas, sin tildes, espacios colapsados a uno. Solo para comparar:
    lo que se muestra es `elegir_forma()`, nunca esta clave."""
    return " ".join(sin_acentos(raw).upper().split())


def levenshtein(a: str, b: str) -> int:
    if a == b:
        return 0
    la, lb = len(a), len(b)
    if abs(la - lb) > 2:
        # Cota inferior real de la distancia: no hace falta calcularla entera
        # para saber que no va a dar <= 2.
        return abs(la - lb)
    fila = list(range(lb + 1))
    for i in range(1, la + 1):
        nueva = [i] + [0] * lb
        for j in range(1, lb + 1):
            costo = 0 if a[i - 1] == b[j - 1] else 1
            nueva[j] = min(fila[j] + 1, nueva[j - 1] + 1, fila[j - 1] + costo)
        fila = nueva
    return fila[lb]


def contenida_primera_palabra(clave_a: str, clave_b: str) -> bool:
    """LIZ ⊂ LIZ TIAMARCA: la más corta cabe entera en la más larga (por
    palabras, no como subcadena) y las dos empiezan igual."""
    ta, tb = clave_a.split(), clave_b.split()
    if not ta or not tb or ta[0] != tb[0]:
        return False
    sa, sb = set(ta), set(tb)
    if sa == sb:
        return False  # misma bolsa de palabras -> ya deberían ser la misma clave
    return sa <= sb or sb <= sa


def elegir_forma(contador: Counter) -> str:
    """La forma tal como se leyó, más frecuente. Empate: la más larga
    (más completa); si sigue empatado, la primera en orden alfabético."""
    return sorted(contador.items(), key=lambda kv: (-kv[1], -len(kv[0]), kv[0]))[0][0]


def elegir_correo(contador: Counter) -> tuple[str | None, int]:
    if not contador:
        return None, 0
    correo, veces = sorted(contador.items(), key=lambda kv: (-kv[1], kv[0]))[0]
    return correo, veces


class DSU:
    def __init__(self, n: int):
        self.p = list(range(n))

    def find(self, x: int) -> int:
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, x: int, y: int) -> bool:
        rx, ry = self.find(x), self.find(y)
        if rx == ry:
            return False
        self.p[rx] = ry
        return True


# --- Paso 1: leer ots y agrupar por (local, clave) -----------------------------
def leer_pares(cur) -> dict[tuple[str, str], dict]:
    """SELECT de solo lectura. `correlativo < 90000` descarta registros
    sintéticos de prueba si alguna vez aparecen en esta base (industec-
    agentes-y-entregables, A3); hoy da 0 filas, así que no cambia nada."""
    cur.execute("""
        SELECT local_codigo, admin_nombre, correo_local, fecha_atencion
          FROM ots
         WHERE en_cuarentena = 0
           AND correlativo < 90000
           AND local_codigo IS NOT NULL
           AND admin_nombre IS NOT NULL
           AND TRIM(admin_nombre) <> ''
    """)
    grupos: dict[tuple[str, str], dict] = {}
    for local, admin_nombre, correo_local, fecha_atencion in cur.fetchall():
        raw = admin_nombre.strip()
        clave = clave_normalizada(raw)
        if not clave:
            continue
        g = grupos.setdefault((local, clave), {
            "local_codigo": local, "clave": clave,
            "raw_counts": Counter(), "correo_counter": Counter(),
            "veces": 0, "visto_ultimo": None,
        })
        g["raw_counts"][raw] += 1
        g["veces"] += 1
        if fecha_atencion and (g["visto_ultimo"] is None or fecha_atencion > g["visto_ultimo"]):
            g["visto_ultimo"] = fecha_atencion
        correo = (correo_local or "").strip().lower()
        if correo and "@industec" not in correo:
            g["correo_counter"][correo] += 1
    return grupos


# --- Paso 2: fusionar variantes dentro de cada local ---------------------------
def fusionar(grupos: dict[tuple[str, str], dict]) -> tuple[list[dict], list[dict]]:
    """Devuelve (grupos_finales, fusiones_registradas). Ver el docstring del
    módulo para la contabilidad que debe cuadrar."""
    por_local: dict[str, list[dict]] = defaultdict(list)
    for g in grupos.values():
        por_local[g["local_codigo"]].append(g)

    grupos_finales: list[dict] = []
    fusiones: list[dict] = []

    for local in sorted(por_local):
        lst = sorted(por_local[local], key=lambda g: g["clave"])
        n = len(lst)
        dsu = DSU(n)
        aristas: list[tuple[int, int, str]] = []
        for i in range(n):
            for j in range(i + 1, n):
                a, b = lst[i]["clave"], lst[j]["clave"]
                d = levenshtein(a, b)
                if d <= 2:
                    motivo = f"Levenshtein {d}"
                elif contenida_primera_palabra(a, b):
                    motivo = "clave contenida en la otra, misma primera palabra"
                else:
                    continue
                if dsu.union(i, j):          # solo cuenta si de verdad conectó dos componentes
                    aristas.append((i, j, motivo))

        clusters: dict[int, list[int]] = defaultdict(list)
        for idx in range(n):
            clusters[dsu.find(idx)].append(idx)

        forma_cluster = {}
        for raiz, idxs in clusters.items():
            total = Counter()
            for idx in idxs:
                total.update(lst[idx]["raw_counts"])
            forma_cluster[raiz] = elegir_forma(total)

        for i, j, motivo in aristas:
            fusiones.append({
                "local": local,
                "clave_a": lst[i]["clave"], "forma_a": elegir_forma(lst[i]["raw_counts"]), "veces_a": lst[i]["veces"],
                "clave_b": lst[j]["clave"], "forma_b": elegir_forma(lst[j]["raw_counts"]), "veces_b": lst[j]["veces"],
                "criterio": motivo,
                "se_agrupa_bajo": forma_cluster[dsu.find(i)],
            })

        for raiz, idxs in clusters.items():
            raw_total, correo_total = Counter(), Counter()
            veces_total, visto_total = 0, None
            for idx in idxs:
                m = lst[idx]
                raw_total.update(m["raw_counts"])
                correo_total.update(m["correo_counter"])
                veces_total += m["veces"]
                if m["visto_ultimo"] and (visto_total is None or m["visto_ultimo"] > visto_total):
                    visto_total = m["visto_ultimo"]
            forma = elegir_forma(raw_total)
            correo, correo_veces = elegir_correo(correo_total)
            grupos_finales.append({
                "local_codigo": local,
                "forma": forma,
                "clave_final": clave_normalizada(forma),
                "veces": veces_total,
                "visto_ultimo": visto_total,
                "correo": correo,
                "correo_veces": correo_veces,
                "variantes_fusionadas": len(idxs),
            })

    return grupos_finales, fusiones


# --- Paso 3: ruido y umbral de siembra ------------------------------------------
def clasificar(grupo: dict, hoy: date) -> tuple[str, str | None]:
    """Devuelve ('SIEMBRA', None) o ('DESCARTADO', motivo)."""
    clave = grupo["clave_final"]

    if any(ch.isdigit() for ch in clave):
        return "DESCARTADO", "nombre con dígitos"

    tokens_trabajo = sorted(set(clave.split()) & PALABRAS_TRABAJO)
    if tokens_trabajo:
        return "DESCARTADO", f"palabra de trabajo: {', '.join(tokens_trabajo)}"

    if sum(1 for ch in clave if ch.isalpha()) < 3:
        return "DESCARTADO", "nombre de menos de 3 letras"

    hace_12_meses = hoy - relativedelta(months=12)
    hace_6_meses = hoy - relativedelta(months=6)

    if grupo["veces"] == 1:
        if grupo["visto_ultimo"] is None:
            return "DESCARTADO", "visto una sola vez, sin fecha de atención registrada"
        if grupo["visto_ultimo"] < hace_12_meses:
            return "DESCARTADO", f"visto una sola vez, la última el {grupo['visto_ultimo']:%d/%m/%Y} (hace más de 12 meses)"

    if grupo["veces"] >= 2:
        return "SIEMBRA", None
    if grupo["visto_ultimo"] and grupo["visto_ultimo"] >= hace_6_meses:
        return "SIEMBRA", None
    return "DESCARTADO", (f"no llega al umbral de siembra: visto una sola vez, "
                           f"la última el {grupo['visto_ultimo']:%d/%m/%Y} "
                           f"(hace más de 6 meses, no más de 12)")


# --- Excel -----------------------------------------------------------------------
def _hoja_encabezado(ws, encabezados: list[str], anchos: list[int]):
    for j, texto in enumerate(encabezados, start=1):
        c = ws.cell(row=1, column=j, value=texto)
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", start_color=AZUL)
    for col, w in zip("ABCDEFGHI", anchos):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def escribir_excel(hoy: date, siembra: list[dict], descartados: list[dict],
                    fusiones: list[dict], pares_leidos: int, filas_ots: int,
                    locales_maestro: int) -> None:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Se siembra"
    _hoja_encabezado(ws, ["Local", "Administrador (forma más leída)", "Veces", "Visto por última vez",
                           "Correo propuesto (sin @industec)", "Correo, veces", "Variantes fundidas"],
                      [10, 34, 8, 16, 32, 12, 14])
    for i, g in enumerate(sorted(siembra, key=lambda g: (g["local_codigo"], -g["veces"])), start=2):
        ws.cell(i, 1, g["local_codigo"])
        ws.cell(i, 2, g["forma"])
        ws.cell(i, 3, g["veces"])
        c = ws.cell(i, 4, g["visto_ultimo"])
        if g["visto_ultimo"]:
            c.number_format = "dd/mm/yyyy"
        ws.cell(i, 5, g["correo"] or "")
        ws.cell(i, 6, g["correo_veces"])
        ws.cell(i, 7, g["variantes_fusionadas"])

    hf = wb.create_sheet("Fusionados")
    _hoja_encabezado(hf, ["Local", "Clave A", "Forma A", "Veces A", "Clave B", "Forma B", "Veces B",
                           "Criterio de fusión", "Se agrupa bajo"],
                      [10, 22, 26, 8, 22, 26, 8, 30, 26])
    for i, f in enumerate(sorted(fusiones, key=lambda f: (f["local"], f["clave_a"])), start=2):
        hf.cell(i, 1, f["local"]); hf.cell(i, 2, f["clave_a"]); hf.cell(i, 3, f["forma_a"])
        hf.cell(i, 4, f["veces_a"]); hf.cell(i, 5, f["clave_b"]); hf.cell(i, 6, f["forma_b"])
        hf.cell(i, 7, f["veces_b"]); hf.cell(i, 8, f["criterio"]); hf.cell(i, 9, f["se_agrupa_bajo"])

    hd = wb.create_sheet("Descartados (motivo)")
    _hoja_encabezado(hd, ["Local", "Clave", "Forma", "Veces", "Visto por última vez", "Motivo"],
                      [10, 26, 30, 8, 16, 48])
    for i, g in enumerate(sorted(descartados, key=lambda g: (g["local_codigo"], g["clave_final"])), start=2):
        hd.cell(i, 1, g["local_codigo"]); hd.cell(i, 2, g["clave_final"]); hd.cell(i, 3, g["forma"])
        hd.cell(i, 4, g["veces"])
        c = hd.cell(i, 5, g["visto_ultimo"])
        if g["visto_ultimo"]:
            c.number_format = "dd/mm/yyyy"
        hd.cell(i, 6, g["motivo"])

    locales_sembrados = sorted({g["local_codigo"] for g in siembra})
    motivos = Counter(g["motivo"] for g in descartados)
    hr = wb.create_sheet("Resumen")
    hr["A1"] = "T2.28.3 · Administrador y correo del local — siembra por análisis (t2_28_admins.py --analizar)"
    hr["A1"].font = Font(bold=True, size=13)
    hr["A2"] = f"Generado el {hoy:%d/%m/%Y}. Fuente: tabla ots de la estación, en_cuarentena=0. NO se cargó nada a la base."
    hr["A2"].font = Font(italic=True, color="FFC00000")
    filas = [
        ("", ""),
        ("Filas de ots con local y administrador", filas_ots),
        ("Pares (local, clave) antes de fundir variantes", pares_leidos),
        ("Fusiones registradas (hoja Fusionados)", len(fusiones)),
        ("Grupos finales (tras fundir)", len(siembra) + len(descartados)),
        ("  ... se siembra", len(siembra)),
        ("  ... descartados", len(descartados)),
        ("Cuadre: fusionados + siembra + descartados = pares leídos",
         "OK" if len(fusiones) + len(siembra) + len(descartados) == pares_leidos else "NO CUADRA"),
        ("", ""),
        ("Locales del maestro", locales_maestro),
        ("Locales con al menos un administrador sembrado", len(locales_sembrados)),
        ("Cifra esperada (T2.28.3): >= 90 locales", "CUMPLE" if len(locales_sembrados) >= 90 else "NO CUMPLE — no cargar, revisar el filtro"),
        ("", ""),
        ("Desglose de descartados por motivo", ""),
    ]
    for etiqueta, valor in filas:
        hr.append([etiqueta, valor])
    for motivo, cnt in sorted(motivos.items(), key=lambda kv: -kv[1]):
        hr.append([f"  {motivo}", cnt])
    hr.column_dimensions["A"].width = 62
    hr.column_dimensions["B"].width = 60

    XLSX.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit(f"'{XLSX}' está abierto en otro programa (Excel). Ciérralo y vuelve a correr el script; "
                  f"no se fuerza la escritura (I-4).")


def escribir_json(siembra: list[dict]) -> None:
    registros = [{
        "local_codigo": g["local_codigo"],
        "admin_nombre": g["forma"],
        "veces": g["veces"],
        "visto_ultimo": g["visto_ultimo"].isoformat() if g["visto_ultimo"] else None,
        "correo": g["correo"],
        "correo_veces": g["correo_veces"],
    } for g in sorted(siembra, key=lambda g: (g["local_codigo"], -g["veces"]))]
    escribir_json_atomico(SALIDA_JSON, json.dumps({
        "generado_en": date.today().isoformat(),
        "fuente": "t2_28_admins.py --analizar (tabla ots, en_cuarentena=0)",
        "nota": "Propuesta sin cargar. La carga a locales_admin es --ejecutar, pendiente de aprobación.",
        "registros": registros,
    }, ensure_ascii=False, indent=2))


# --- main -------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--analizar", action="store_true", help="Analiza y escribe el Excel + JSON (comportamiento por omisión)")
    ap.add_argument("--ejecutar", action="store_true", help="Carga a locales_admin — NO implementado en esta fase")
    a = ap.parse_args()

    if a.ejecutar:
        sys.exit("t2_28_admins.py --ejecutar todavía no está implementado a propósito: la carga a "
                 "locales_admin requiere la aprobación de Andrés con la cifra exacta "
                 "(T2_28_OBSERVACIONES_INDUSTEC.md, T2.28.3). Corre --analizar y revisa el Excel primero.")

    env = cargar_env(("DB_HOST", "DB_USER", "DB_PASSWORD", "DB_NAME"))
    cnx = mysql.connector.connect(host=env["DB_HOST"], port=int(env.get("DB_PORT", 3306)),
                                   user=env["DB_USER"], password=env["DB_PASSWORD"], database=env["DB_NAME"])
    cur = cnx.cursor()

    grupos = leer_pares(cur)
    filas_ots = sum(g["veces"] for g in grupos.values())
    pares_leidos = len(grupos)

    cur.execute("SELECT COUNT(*) FROM locales")
    locales_maestro = cur.fetchone()[0]
    cnx.close()

    grupos_finales, fusiones = fusionar(grupos)

    hoy = date.today()
    siembra, descartados = [], []
    for g in grupos_finales:
        categoria, motivo = clasificar(g, hoy)
        if categoria == "SIEMBRA":
            siembra.append(g)
        else:
            g["motivo"] = motivo
            descartados.append(g)

    cuadre = len(fusiones) + len(siembra) + len(descartados)
    if cuadre != pares_leidos:
        sys.exit(f"NO CUADRA (I-10): fusionados({len(fusiones)}) + siembra({len(siembra)}) + "
                 f"descartados({len(descartados)}) = {cuadre}, pero pares leídos = {pares_leidos}. "
                 f"No se escribe el Excel.")

    escribir_excel(hoy, siembra, descartados, fusiones, pares_leidos, filas_ots, locales_maestro)
    escribir_json(siembra)

    locales_sembrados = len({g["local_codigo"] for g in siembra})
    print(f"Filas de ots leídas (local+admin, en_cuarentena=0): {filas_ots}")
    print(f"Pares (local, clave) antes de fundir: {pares_leidos}")
    print(f"Fusiones registradas: {len(fusiones)}")
    print(f"Grupos finales: {len(siembra) + len(descartados)}  (siembra {len(siembra)} · descartados {len(descartados)})")
    print(f"Cuadre fusionados+siembra+descartados = pares leídos: {cuadre} == {pares_leidos} -> OK")
    print(f"Locales con al menos un administrador a sembrar: {locales_sembrados} de {locales_maestro} del maestro "
          f"({'CUMPLE' if locales_sembrados >= 90 else 'NO CUMPLE'} el umbral >= 90)")
    print(f"Excel: {XLSX}")
    print(f"JSON:  {SALIDA_JSON}")


if __name__ == "__main__":
    main()
