"""
Historico de preventivos en el formato de seguimiento de la administracion.

El preventivo no se planifica como el correctivo. No hay aviso SAP por caso ni
dos fases: hay un programa de ingresos por local -- hasta cuatro "vueltas" al
ano, cada una de uno o dos dias -- que la administracion sigue en
`SEGUIMIENTO PREVENTIVOS _ INDUSTEC.xlsx`: una hoja por ano, una fila por
local, y las columnas FECHA INGRESO 1 a 4.

Este script rellena ese mismo formato con los ingresos REALMENTE ejecutados,
reconstruidos desde las 720 ordenes preventivas de la base. Es la base de
partida sobre la que despues se marcaran los ingresos del dia.

Como se reconstruye un ingreso: las ordenes preventivas de un local se agrupan
por cercania de fechas. Dos ordenes a menos de 4 dias son el mismo ingreso -- el
mantenimiento de un local ocupa uno o dos dias seguidos y genera una orden por
dia (sufijo D1, D2...), a veces mas de una por dia cuando entran dos tecnicos.
Verificado contra el archivo real: G001EC figura como "12 y 13 de febrero" y la
base tiene ordenes el 12 y el 13 de febrero de 2026.

Lo que NO se inventa (I-7):
  - CIUDAD sale del propio archivo de la administracion, que es la unica fuente
    que la tiene; el maestro de locales no guarda ciudad. El local que no
    aparezca alli queda con la celda vacia, no con una ciudad supuesta.
  - Un local sin ordenes preventivas ese ano se lista igual, con los ingresos
    vacios y la observacion SIN INGRESO REGISTRADO. Omitirlo escondería
    justamente lo que hay que ver.
  - Un quinto ingreso o posterior no cabe en el formato de cuatro columnas: se
    nombra en OBSERVACIONES en vez de descartarse.

Uso:
    .venv/Scripts/python.exe scripts/t2_historico_preventivos.py
"""
import sys
from collections import defaultdict
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from agente2_consolidador import conectar

PLANTILLA_DIR = Path(r"D:\RESPALDOS\_ORIGEN_DRIVE\Oficina Industec\MTTO PREVENTIVO UIO  2026")
SALIDA = Path(r"D:\INDUSTECH IA\SALIDAS IA\MANTENIMIENTO\MANTENIMIENTOS PREVENTIVOS")
NOMBRE_SALIDA = "SEGUIMIENTO PREVENTIVOS (generado agente).xlsx"

COLS = ["#", "LOCAL", "UBICACIÓN", "CIUDAD", "FECHA INGRESO 1", "FECHA INGRESO 2",
        "FECHA INGRESO 3", "FECHA INGRESO 4", "OBSERVACIONES"]
FILA_ENCABEZADO = 2          # el archivo real deja la fila 1 en blanco
FILA_DATOS = 3
MAX_INGRESOS = 4
DIAS_MISMO_INGRESO = 4       # dos ordenes a menos de 4 dias son la misma visita


def plantilla():
    candidatos = sorted(PLANTILLA_DIR.glob("SEGUIMIENTO PREVENTIVOS*.xlsx"))
    if not candidatos:
        sys.exit(f"ABORTA: no aparece el archivo real de seguimiento en {PLANTILLA_DIR}")
    return candidatos[0]


def ciudades_conocidas(ruta):
    """local -> ciudad, leido del archivo de la administracion (solo lectura)."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    mapa = {}
    for ws in wb.worksheets:
        for r in range(FILA_DATOS, ws.max_row + 1):
            local, ciudad = ws.cell(r, 2).value, ws.cell(r, 4).value
            if local and ciudad:
                mapa.setdefault(str(local).strip().upper(), str(ciudad).strip().upper())
    return mapa


def cargar(cnx):
    cur = cnx.cursor(dictionary=True)
    cur.execute("""SELECT local_codigo, zona, cadena, nombre FROM locales
                   WHERE zona IN ('UIO','LARB','CNLJ') ORDER BY zona, local_codigo""")
    locales = cur.fetchall()
    cur.execute("""SELECT id_industec, local_codigo, zona, fecha_atencion, dia_intervencion,
                          tecnico_nombre, aviso
                   FROM ots
                   WHERE en_cuarentena=0 AND modulo='PREVENTIVO' AND fecha_atencion IS NOT NULL
                   ORDER BY local_codigo, fecha_atencion""")
    ordenes = cur.fetchall()
    cur.close()
    return locales, ordenes


def agrupar_ingresos(ordenes):
    """Ordenes de un local y ano -> lista de ingresos, cada uno con sus ordenes."""
    ingresos = []
    for o in sorted(ordenes, key=lambda x: (x["fecha_atencion"], x["id_industec"])):
        if ingresos and (o["fecha_atencion"] - ingresos[-1][-1]["fecha_atencion"]).days < DIAS_MISMO_INGRESO:
            ingresos[-1].append(o)
        else:
            ingresos.append([o])
    return ingresos


def escribir_hoja(ws, filas):
    """Reescribe los valores conservando el formato de la fila 3 del original."""
    modelo = [copy(ws.cell(FILA_DATOS, j)._style) for j in range(1, len(COLS) + 1)]
    filas_originales = ws.max_row

    for i, fila in enumerate(filas):
        r = FILA_DATOS + i
        for j, nombre in enumerate(COLS, start=1):
            celda = ws.cell(r, j)
            celda.value = fila.get(nombre)
            if r > filas_originales:
                celda._style = copy(modelo[j - 1])
            if nombre.startswith("FECHA") and isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"

    for r in range(FILA_DATOS + len(filas), filas_originales + 1):
        for j in range(1, len(COLS) + 1):
            ws.cell(r, j).value = None


def main():
    ruta_plantilla = plantilla()
    ciudades = ciudades_conocidas(ruta_plantilla)
    cnx = conectar()
    locales, ordenes = cargar(cnx)
    cnx.close()

    por_local_anio = defaultdict(list)
    for o in ordenes:
        por_local_anio[(o["local_codigo"], o["fecha_atencion"].year)].append(o)

    anios = sorted({f.year for f in (o["fecha_atencion"] for o in ordenes)})
    wb = openpyxl.load_workbook(ruta_plantilla)

    detalle = []
    resumen = {}
    for anio in anios:
        nombre_hoja = str(anio)
        ws = wb[nombre_hoja] if nombre_hoja in wb.sheetnames else wb.create_sheet(nombre_hoja)
        if ws.cell(FILA_ENCABEZADO, 1).value is None:
            for j, nombre in enumerate(COLS, start=1):
                ws.cell(FILA_ENCABEZADO, j).value = nombre

        filas, con_ingreso = [], 0
        for n, loc in enumerate(locales, start=1):
            codigo = loc["local_codigo"]
            ingresos = agrupar_ingresos(por_local_anio.get((codigo, anio), []))
            fila = {"#": n, "LOCAL": codigo,
                    "UBICACIÓN": (loc["nombre"] or "").upper(),
                    "CIUDAD": ciudades.get(codigo.upper())}
            for k in range(MAX_INGRESOS):
                fila[f"FECHA INGRESO {k+1}"] = ingresos[k][0]["fecha_atencion"] if k < len(ingresos) else None

            notas = []
            if not ingresos:
                notas.append("SIN INGRESO REGISTRADO")
            else:
                con_ingreso += 1
                notas.append(f"{len(ingresos)} ingreso(s) · {sum(len(i) for i in ingresos)} OT")
                for extra in ingresos[MAX_INGRESOS:]:
                    notas.append("INGRESO ADICIONAL " + extra[0]["fecha_atencion"].isoformat())
            fila["OBSERVACIONES"] = " · ".join(notas)
            filas.append(fila)

            for k, ingreso in enumerate(ingresos, start=1):
                for o in ingreso:
                    detalle.append([codigo, loc["zona"], loc["cadena"], anio, k,
                                    o["fecha_atencion"], o["dia_intervencion"],
                                    o["id_industec"], o["tecnico_nombre"], o["aviso"]])

        escribir_hoja(ws, filas)
        resumen[anio] = (len(filas), con_ingreso)

    # Hoja de trazabilidad: el formato de 4 columnas no puede citar cada orden,
    # y toda cifra tiene que poder rastrearse hasta la OT que la sostiene (I-7).
    if "DETALLE OTS" in wb.sheetnames:
        del wb["DETALLE OTS"]
    ws_det = wb.create_sheet("DETALLE OTS")
    ws_det.append(["LOCAL", "ZONA", "CADENA", "AÑO", "INGRESO", "FECHA", "DIA",
                   "OT INDUSTEC", "TECNICO", "AVISO"])
    for fila in sorted(detalle, key=lambda x: (x[3], x[1], x[0], x[5])):
        ws_det.append(fila)

    SALIDA.mkdir(parents=True, exist_ok=True)
    ruta = SALIDA / NOMBRE_SALIDA
    wb.save(ruta)

    # Verificacion I-10: ninguna orden preventiva puede quedar fuera del detalle
    if len(detalle) != len(ordenes):
        sys.exit(f"ABORTA: {len(ordenes)} ordenes preventivas en la base y "
                 f"{len(detalle)} en el detalle generado")

    print(f"Plantilla leida (solo lectura): {ruta_plantilla.name}")
    print(f"Ciudades conocidas desde el archivo real: {len(ciudades)} locales")
    for anio, (n, con) in resumen.items():
        print(f"  {anio}: {n} locales listados · {con} con ingreso registrado · "
              f"{n - con} sin ingreso")
    print(f"Ordenes preventivas en el detalle: {len(detalle)} de {len(ordenes)} (I-10 OK)")
    print(f"Generado: {ruta}")


if __name__ == "__main__":
    main()
