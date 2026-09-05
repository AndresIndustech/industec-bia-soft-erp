"""
Verificacion cruzada del historico de correctivos (I-10) contra los planes
mensuales que la administracion llevo a mano.

Fuente independiente: los 24 archivos reales `PLANES MENSUALES/{MES}.xlsx` de
las tres zonas, enero a agosto 2026. Nada de lo que genera el agente participa
en esta comparacion: se contrastan filas por CLAVE DE NEGOCIO (el aviso SAP de
la columna "# OT"), nunca por posicion, porque el orden de filas difiere sin
que eso sea un error.

Criterio de aceptacion vigente (T1.11): >=95% de celdas identicas antes de que
la administracion deje de generar el archivo a mano.

Uso:
    .venv/Scripts/python.exe scripts/t2_historico_comparar.py
"""
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ORIGEN = Path(r"D:\RESPALDOS\_ORIGEN_DRIVE\GESTION DE OTS INDUSTEC\2026\PLANES SEMANALES")
GENERADO = Path(r"D:\INDUSTECH IA\SALIDAS IA\MANTENIMIENTO\MANTENIMIENTOS CORRECTIVOS")

ZONAS = {
    "UIO": ("PLAN DE TRABAJO _ ZONA UIO", "ZONA UIO"),
    "LARB": ("PLAN DE TRABAJO _ ZONA LARB", "ZONA LARB"),
    "CNLJ": ("PLAN DE TRABAJO _ ZONA C-L", "ZONA CUENCA LOJA"),
}
MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
         "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11,
         "DICIEMBRE": 12}

COLS = ["#", "# OT", "TECNICO EVALUACION", "TECNICO CIERRE", "LOCAL",
        "FECHA DE INICIO", "EQUIPO", "MARCA", "TRABAJO REALIZADO EVALUACION",
        "REPUESTO", "ESTATUS SAP", "PRESUPUESTO", "TRABAJO REALIZADO CIERRE",
        "OBSERVACIONES", "ESTATUS DEL EQUIPO", "#OT INDUSTEC EVALUACION",
        "FECHA EVALUACION", "#OT INDUSTEC CIERRE", "FECHA CIERRE",
        "REQUERIMIENTO A TIEMPO", "CALIFICACION SATISFACCIÓN", "ESTADO"]


def normalizar(v):
    """Compara contenido, no tipografia: fechas a ISO, espacios colapsados,
    mayusculas unificadas. Lo que queda distinto es diferencia real de dato."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split()).strip().upper()


def leer_filas(path):
    """Devuelve {aviso: fila}. Algunos meses traen una columna 'Columna1'
    intercalada tras LOCAL (enero) que hay que saltar para no correr todo."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    encabezado = [normalizar(c.value) for c in ws[1]]
    salto = 6 if "COLUMNA1" in encabezado else None

    filas = {}
    for r in range(2, ws.max_row + 1):
        aviso = normalizar(ws.cell(r, 2).value)
        if not aviso:
            continue
        fila = {}
        for j, nombre in enumerate(COLS, start=1):
            col = j if (salto is None or j < salto) else j + 1
            fila[nombre] = normalizar(ws.cell(r, col).value)
        filas[aviso] = fila
    return filas


def comparar(real, generado):
    reales, generadas = leer_filas(real), leer_filas(generado)
    comunes = set(reales) & set(generadas)
    total = iguales = 0
    por_columna = {}
    for aviso in comunes:
        for col in COLS:
            if col == "#":       # el numero de fila no es contenido de negocio
                continue
            total += 1
            vr, vg = reales[aviso][col], generadas[aviso][col]
            if vr == vg:
                iguales += 1
            else:
                por_columna.setdefault(col, []).append((aviso, vr[:60], vg[:60]))
    return {
        "filas_real": len(reales), "filas_generado": len(generadas),
        "comunes": len(comunes),
        "solo_real": sorted(set(reales) - set(generadas)),
        "solo_generado": sorted(set(generadas) - set(reales)),
        "celdas": total, "iguales": iguales,
        "pct": (iguales / total * 100) if total else 0.0,
        "por_columna": por_columna,
    }


def main():
    resultados, difs_globales = [], {}
    for zona, (carpeta_real, carpeta_gen) in ZONAS.items():
        for archivo in sorted((ORIGEN / carpeta_real / "PLANES MENSUALES").glob("*.xlsx")):
            mes = MESES.get(archivo.stem.upper())
            if not mes:
                continue
            gen = GENERADO / carpeta_gen / "PLANES MENSUALES" / f"2026-{mes:02d} {archivo.stem.upper()} (generado agente).xlsx"
            if not gen.exists():
                print(f"  falta el generado: {gen.name} ({zona})")
                continue
            r = comparar(archivo, gen)
            resultados.append((zona, archivo.stem.upper(), r))
            for col, difs in r["por_columna"].items():
                difs_globales.setdefault(col, []).extend(difs)

    if not resultados:
        sys.exit("ABORTA: no habia nada que comparar")

    print(f"{'ZONA':<6}{'MES':<11}{'real':>6}{'gen':>6}{'comunes':>9}{'celdas':>8}{'iguales':>9}{'%':>8}")
    for zona, mes, r in resultados:
        print(f"{zona:<6}{mes:<11}{r['filas_real']:>6}{r['filas_generado']:>6}"
              f"{r['comunes']:>9}{r['celdas']:>8}{r['iguales']:>9}{r['pct']:>7.1f}%")

    celdas = sum(r["celdas"] for _, _, r in resultados)
    iguales = sum(r["iguales"] for _, _, r in resultados)
    pct = iguales / celdas * 100 if celdas else 0
    filas_real = sum(r["filas_real"] for _, _, r in resultados)
    comunes = sum(r["comunes"] for _, _, r in resultados)
    print(f"\nTOTAL: {celdas} celdas comparadas · {iguales} identicas · {pct:.1f}%")
    print(f"Cobertura de filas: {comunes} de {filas_real} casos del archivo real "
          f"({comunes / filas_real * 100:.1f}%)")
    print(f"Criterio T1.11 (>=95%): {'CUMPLE' if pct >= 95 else 'NO CUMPLE'}")

    print("\n=== Diferencias por columna (todas las zonas y meses) ===")
    for col, difs in sorted(difs_globales.items(), key=lambda x: -len(x[1])):
        print(f"  {col}: {len(difs)}")
        for aviso, vr, vg in difs[:2]:
            print(f"      aviso {aviso}: real={vr!r}  generado={vg!r}")


if __name__ == "__main__":
    main()
