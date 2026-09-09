"""
T2.8 - Carga el padrón de técnicos que entregó la administración.

FUENTE: G:\\Mi unidad\\INDUSTEC IA\\LISTADO DE TECNICOS ACTUALIZADO.xlsx, hoja
PADRON. G:\\ es de SOLO LECTURA (I-3): se lee y no se toca.

QUE ARREGLA:
La tabla `tecnicos` tenía 19 filas, todas activas, sin fecha de salida. Al cruzar
contra el padrón nuevo aparecieron **9 personas trabajando hoy que la base no
conocía** y 9 que ya no están. Eso venía distorsionando una conclusión del
proyecto: las órdenes firmadas por gente "fuera de la nómina" se explicaban como
rotación, y en realidad **1.362 de ellas tienen autor conocido** -- varios de
ellos trabajando hoy. No era rotación: era un maestro incompleto.

LAS TRES SITUACIONES, que no se pueden colapsar:
    fecha_salida NULL   + sin_registro 0  ->  sigue trabajando
    fecha_salida <fecha>                  ->  salió ese día
    fecha_salida NULL   + sin_registro 1  ->  salió, no consta cuándo

Guardar "SIN REGISTRO" como NULL a secas dejaría a 11 personas pareciendo
activas, que es exactamente el problema que esto viene a arreglar.

ESTRATEGIA DE CARGA: DELETE + INSERT.
Nada referencia a `tecnicos` por clave foránea (se verificó contra
information_schema), la tabla se reconstruye entera desde su documento origen, y
la clave natural -la cédula- solo la traen los 19 vigentes. Con upsert, los 21
sin cédula quedarían huérfanos en cada corrida.

Uso:
    .venv/Scripts/python.exe scripts/t2_8_padron_tecnicos.py            # simula
    .venv/Scripts/python.exe scripts/t2_8_padron_tecnicos.py --ejecutar
"""

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

import mysql.connector
import openpyxl

ENV_PATH = Path(r"D:\INDUSTECH IA\desarrollo\agentes\config\.env")
PADRON = Path(r"G:\Mi unidad\INDUSTEC IA\LISTADO DE TECNICOS ACTUALIZADO.xlsx")
HOJA = "PADRON"

# A NOMBRE | B F.INGRESO | C F.SALIDA | D ZONA | E MANTENIMIENTO | F AFILIADO | G CEDULA
COL_NOMBRE, COL_INGRESO, COL_SALIDA = 0, 1, 2
COL_ZONA, COL_MANT, COL_AFIL, COL_CEDULA = 3, 4, 5, 6
COL_MINIMA = 7

# Lo que el archivo dice y lo que significa. Cualquier otra cosa NO se adivina.
SIN_REGISTRO = {"SIN REGISTRO", "SINREGISTRO"}
ZONAS = {"UIO", "LARB", "CNLJ", "OTRA"}

# Los 19 usuarios que aprobo Andres Basantes el 2026-09-08, en
# PADRON_TECNICOS_REVISION.md. Estan aqui para que el script no los invente: el
# nombre de usuario se CALCULA del padron y se COMPARA contra esta lista. Si el
# archivo cambia un apellido, el calculo se corre y el script aborta en vez de
# crear un usuario que nadie aprobo.
USUARIOS_CONFIRMADOS = {
    'CNLJ': {'dmelendrez': 'JEFE_ZONA', 'abarrozo': 'TECNICO', 'aibarra': 'TECNICO',
             'dcuenca': 'TECNICO', 'ebasantes': 'TECNICO', 'hmelendrez': 'TECNICO',
             'lerazo': 'TECNICO'},
    'LARB': {'lperdomo': 'JEFE_ZONA', 'cfarfan': 'TECNICO', 'ecampos': 'TECNICO',
             'emontoya': 'TECNICO', 'portiz': 'TECNICO', 'storres': 'TECNICO'},
    'UIO':  {'kchimbo': 'JEFE_ZONA', 'amorales': 'TECNICO', 'ajumbo': 'TECNICO',
             'dsisalema': 'TECNICO', 'ftipan': 'TECNICO', 'htaco': 'TECNICO'},
}

# Donde queda la lista que se sube a Hostinger. Fuera de git: son nombres de
# personas. No lleva cedula -- esa se queda en la estacion.
SALIDA_JSON = Path(r"D:\INDUSTECH IA\SALIDAS IA\OTS\padron_usuarios.json")

# Totales esperados, tomados del sondeo del archivo el 2026-09-08. Son la
# compuerta de I-10: si el archivo cambió, se aborta antes de escribir.
ESPERADO_TOTAL = 40
ESPERADO_ACTIVOS = 19
ESPERADO_SIN_REGISTRO = 11
ESPERADO_CON_FECHA_SALIDA = 10


def norma(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.upper()).strip()


def cargar_env():
    env = {}
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


def partir_nombre(completo):
    """'Diego Fernando Melendrez Sinchiguano' -> ('Diego Fernando', 'Melendrez Sinchiguano').

    El padrón escribe nombre(s) y luego apellido(s). Con 4 palabras la partida
    es 2/2; con 3, se asume 1 nombre y 2 apellidos, que es lo habitual en
    Ecuador. Con 2, uno y uno.
    """
    p = str(completo).strip().split()
    if len(p) >= 4:
        return " ".join(p[:2]), " ".join(p[2:])
    if len(p) == 3:
        return p[0], " ".join(p[1:])
    if len(p) == 2:
        return p[0], p[1]
    return p[0] if p else "", ""


def nombre_usuario(nombres, apellidos):
    """'Kevin Omar' + 'Chimbo Amaguana' -> 'kchimbo'.

    Inicial del primer nombre + primer apellido, sin tildes. Es como se
    reconocen entre ellos, y ninguno de los 19 vigentes se repite. El resultado
    se verifica contra USUARIOS_CONFIRMADOS: aqui no se decide nada.
    """
    ini = norma(nombres)[:1].lower()
    ape = norma(apellidos).split()
    return f"{ini}{ape[0].lower()}" if ini and ape else ""


def rol_de(tipo_tecnico):
    """El padron marca 'JEFE TECNICO' en la columna MANTENIMIENTO: uno por zona."""
    return 'JEFE_ZONA' if 'JEFE' in norma(tipo_tecnico) else 'TECNICO'


def leer_padron():
    if not PADRON.exists():
        sys.exit(f"No encuentro {PADRON}")
    wb = openpyxl.load_workbook(PADRON, data_only=True)
    ws = wb[HOJA]
    if ws.max_column < COL_MINIMA:
        sys.exit(f"Cambio de formato: se esperaban {COL_MINIMA} columnas y hay {ws.max_column}")

    gente, avisos = [], []
    for n, f in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row,
                                       max_col=ws.max_column, values_only=True), start=2):
        if not f[COL_NOMBRE]:
            continue
        completo = str(f[COL_NOMBRE]).strip()
        nombres, apellidos = partir_nombre(completo)

        # --- Salida: las tres situaciones -----------------------------------
        crudo = f[COL_SALIDA]
        fecha_salida, sin_registro = None, 0
        if crudo is None or str(crudo).strip() == "":
            pass                                   # sigue trabajando
        elif isinstance(crudo, (dt.date, dt.datetime)):
            fecha_salida = crudo.date() if isinstance(crudo, dt.datetime) else crudo
        elif norma(crudo).replace(" ", "") in {x.replace(" ", "") for x in SIN_REGISTRO}:
            sin_registro = 1
        else:
            # Texto que no se reconoce: no se adivina si salió o no (I-7).
            avisos.append(f"fila {n}: salida ilegible {crudo!r} en {completo!r}")
            sin_registro = 1

        # --- Ingreso: 'SI' no es una fecha, es "por confirmar" --------------
        ci = f[COL_INGRESO]
        fecha_ingreso, por_confirmar = None, 0
        if isinstance(ci, (dt.date, dt.datetime)):
            fecha_ingreso = ci.date() if isinstance(ci, dt.datetime) else ci
        elif ci is not None and str(ci).strip() != "":
            # Regla del cliente 2026-09-08: el 'SI' significa que la
            # administración todavía no confirma la fecha. No es un error.
            por_confirmar = 1

        zona = norma(f[COL_ZONA]) or None
        if zona and zona not in ZONAS:
            avisos.append(f"fila {n}: zona desconocida {zona!r} en {completo!r}")
            zona = None

        ced = re.sub(r"[^0-9]", "", str(f[COL_CEDULA] or "")) or None
        afil = f[COL_AFIL]
        afiliado = 1 if norma(afil) == "SI" else (0 if norma(afil) == "NO" else None)

        activo = 1 if (fecha_salida is None and sin_registro == 0) else 0
        tipo = (str(f[COL_MANT]).strip()[:40] if f[COL_MANT] else None)

        gente.append({
            "nombres": nombres[:80], "apellidos": apellidos[:80], "cedula": ced,
            # Solo los vigentes llevan usuario. Quien ya salio no entra al
            # sistema; su historial se sigue leyendo por nombre.
            "usuario": (nombre_usuario(nombres, apellidos) or None) if activo else None,
            "fecha_ingreso": fecha_ingreso, "fecha_salida": fecha_salida,
            "salida_sin_registro": sin_registro, "ingreso_por_confirmar": por_confirmar,
            "tipo_tecnico": tipo,
            "afiliado": afiliado, "zona_asignada": zona, "activo": activo,
            "_rol": rol_de(tipo) if activo else None,
            "fuente_padron": PADRON.name[:60],
            "_completo": completo,
        })
    return gente, avisos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ejecutar", action="store_true",
                    help="sin esto solo simula y no escribe nada")
    args = ap.parse_args()

    gente, avisos = leer_padron()
    activos = [g for g in gente if g["activo"]]
    sin_reg = [g for g in gente if g["salida_sin_registro"]]
    con_fecha = [g for g in gente if g["fecha_salida"]]
    por_conf = [g for g in gente if g["ingreso_por_confirmar"]]

    print(f"padron leido        : {len(gente)} personas")
    print(f"  activos           : {len(activos)}")
    print(f"  salidos con fecha : {len(con_fecha)}")
    print(f"  salidos sin fecha : {len(sin_reg)}")
    print(f"  ingreso a confirmar:{len(por_conf)}")
    for a in avisos:
        print(f"  AVISO: {a}")

    # ---- Compuerta de verificacion cruzada (I-10) --------------------------
    # Los totales NO se imprimen y se sigue: se comparan y se aborta.
    ok = True
    for etiqueta, real, esperado in [
        ("total", len(gente), ESPERADO_TOTAL),
        ("activos", len(activos), ESPERADO_ACTIVOS),
        ("sin registro", len(sin_reg), ESPERADO_SIN_REGISTRO),
        ("con fecha salida", len(con_fecha), ESPERADO_CON_FECHA_SALIDA),
    ]:
        marca = "OK" if real == esperado else "DIFERENCIA"
        print(f"  {etiqueta:<18} esperado {esperado:>3} / real {real:>3}  {marca}")
        ok &= real == esperado

    # Los 19 usuarios y sus roles salen del padron; aqui se comprueba que sean
    # exactamente los que Andres aprobo. Un apellido corregido en el Excel
    # cambiaria un nombre de usuario en silencio, y alguien quedaria sin poder
    # entrar el dia del arranque.
    calculado = {}
    for g in activos:
        calculado.setdefault(g["zona_asignada"], {})[g["usuario"]] = g["_rol"]
    if calculado != USUARIOS_CONFIRMADOS:
        for z in sorted(set(calculado) | set(USUARIOS_CONFIRMADOS)):
            c, e = calculado.get(z, {}), USUARIOS_CONFIRMADOS.get(z, {})
            for u in sorted(set(c) | set(e)):
                if c.get(u) != e.get(u):
                    print(f"  {z}: usuario {u!r} aprobado={e.get(u)} calculado={c.get(u)}")
        ok = False
    else:
        print("  usuarios calculados         19 / 19 iguales a los aprobados  OK")

    ceds = [g["cedula"] for g in gente if g["cedula"]]
    if len(ceds) != len(set(ceds)):
        print("  cedulas DUPLICADAS en el padron")
        ok = False
    for g in activos:
        if not g["zona_asignada"]:
            print(f"  activo sin zona: {g['_completo']}")
            ok = False

    if not ok:
        print("\nABORTADO: el padron no cuadra con lo verificado. No se escribe nada en la BD.",
              file=sys.stderr)
        sys.exit(1)
    print("\nVerificacion cruzada: OK")

    if not args.ejecutar:
        print("\nSIMULACION. Nada se escribio. Corre con --ejecutar para aplicar.")
        return

    env = cargar_env()
    cnx = mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"])
    try:
        cur = cnx.cursor()
        cur.execute("SELECT COUNT(*) FROM tecnicos")
        antes = cur.fetchone()[0]

        campos = ["nombres", "apellidos", "cedula", "usuario", "fecha_ingreso", "fecha_salida",
                  "salida_sin_registro", "ingreso_por_confirmar", "tipo_tecnico",
                  "afiliado", "zona_asignada", "activo", "fuente_padron"]
        cur.execute("DELETE FROM tecnicos")
        cur.executemany(
            f"INSERT INTO tecnicos ({', '.join(campos)}) "
            f"VALUES ({', '.join(['%s'] * len(campos))})",
            [tuple(g[c] for c in campos) for g in gente])
        cnx.commit()

        cur.execute("""SELECT COUNT(*),
                              SUM(activo=1),
                              SUM(salida_sin_registro=1),
                              SUM(fecha_salida IS NOT NULL),
                              SUM(activo=1 AND (fecha_salida IS NOT NULL OR salida_sin_registro=1))
                         FROM tecnicos""")
        t, act, sr, cf, incoherentes = cur.fetchone()
        print(f"\ntecnicos: {antes} -> {t}")
        print(f"  activos {act} | sin registro {sr} | con fecha de salida {cf}")
        if incoherentes:
            print(f"  ALERTA: {incoherentes} filas activas y con salida a la vez", file=sys.stderr)
            sys.exit(1)
        c2 = cur
        c2.execute("SELECT COUNT(*), COUNT(DISTINCT usuario) FROM tecnicos WHERE usuario IS NOT NULL")
        con_usr, distintos = c2.fetchone()
        print(f"  con usuario {con_usr} | distintos {distintos}")
        if con_usr != len(activos) or distintos != con_usr:
            print("  ALERTA: los usuarios no son uno por persona vigente", file=sys.stderr)
            sys.exit(1)
    finally:
        cnx.close()

    # La lista que se sube a Hostinger como nucleo/padron.json. Va SIN cedula ni
    # fecha: la app solo necesita a quien crear, con que rol y en que zona. Todo
    # lo demas se queda en la estacion (DECISION_ARQUITECTURA_Y_DATOS.md).
    SALIDA_JSON.parent.mkdir(parents=True, exist_ok=True)
    SALIDA_JSON.write_text(json.dumps({
        "generado": dt.date.today().isoformat(),
        "fuente": PADRON.name,
        "personas": [{"usuario": g["usuario"], "nombre": g["_completo"],
                      "rol": g["_rol"], "zona": g["zona_asignada"]}
                     for g in sorted(activos, key=lambda x: (x["zona_asignada"], x["usuario"]))],
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print()
    print(f"lista para Hostinger: {SALIDA_JSON}")
    print("  subela a  public_html/ot/nucleo/padron.json  y abre alta_padron.php")


if __name__ == "__main__":
    main()
