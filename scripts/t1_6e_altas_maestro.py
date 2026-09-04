"""T1.6e - Altas al maestro de locales y alias de codigos mal escritos.

Decisiones del cliente, 2026-09-04, tomadas sobre la evidencia de T1.6b:

  ALTAS (5 locales nuevos)
    G044, G045, G047, G054  Pollo Gus. El correo del documento (gusg44@gus.com.ec...)
                            confirma el codigo exactamente. INDUSTEC los archivo en
                            "ORDENES DE TRABAJO OTROS CLIENTES\\PROYECTO HORNOS GUS":
                            son locales reales atendidos en un proyecto aparte.
    T050                    Tropi Burger Portal Shopping. La zona UIO viene del propio
                            nombre de archivo del sistema (OT-0966-T050-0-UIO.pdf).

  ALIAS (2 codigos que no son locales nuevos sino mal escritos)
    K166 -> K167EC   el PDF trae el correo kfc167@kfc.com.ec y K167EC es "KFC CENTRO
                     HISTORICO"; la Plaza San Francisco esta en el Centro Historico.
    A018 -> A014EC   el PDF trae amca14@americandeli.com.ec y A014EC es el unico
                     American Deli de LARB en el maestro.

SOBRE LA ZONA DE LOS CUATRO GUS: los documentos no declaran ciudad y los tecnicos que
los firman no estan en la nomina de INDUSTEC. Los nombres de ubicacion (Boyaca,
Tungurahua, Alborada, Riocentro El Dorado) no corresponden a ninguna de las tres zonas
contratadas. NO se les inventa una zona: se cargan como 'OTRA' y quedan marcados en el
archivo de propuesta para que la administracion confirme la real.

El maestro de verdad es el Excel de la administracion en Drive, que NO se toca (I-3, I-4).
Este script carga la base para que el sistema funcione y deja en SALIDAS IA un archivo
con las filas listas para que ella las pegue en su Excel cuando quiera.
"""
from pathlib import Path

import mysql.connector
import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = Path(r"D:\INDUSTECH IA\agentes")
PROPUESTA = Path(r"D:\INDUSTECH IA\SALIDAS IA\CALIDAD\PROPUESTA_ALTAS_MAESTRO_LOCALES.xlsx")

ALTAS = [
    # codigo,   zona,   cadena,          nombre,                          correo,                  nota
    ("G044EC", "OTRA", "GUS", "POLLO GUS BOYACA", "gusg44@gus.com.ec",
     "Proyecto hornos Gus, fuera del contrato KFC. ZONA POR CONFIRMAR: el documento no declara ciudad."),
    ("G045EC", "OTRA", "GUS", "POLLO GUS TUNGURAHUA", "gusg45@gus.com.ec",
     "Proyecto hornos Gus, fuera del contrato KFC. ZONA POR CONFIRMAR: el documento no declara ciudad."),
    ("G047EC", "OTRA", "GUS", "POLLO GUS ALBORADA", "gusg47@gus.com.ec",
     "Proyecto hornos Gus, fuera del contrato KFC. ZONA POR CONFIRMAR: el documento no declara ciudad."),
    ("G054EC", "OTRA", "GUS", "POLLO GUS RIOCENTRO EL DORADO", "gusg54@gus.com.ec",
     "Proyecto hornos Gus, fuera del contrato KFC. ZONA POR CONFIRMAR: el documento no declara ciudad."),
    ("T050EC", "UIO", "TROPI BURGER", "TROPI BURGER PORTAL SHOPPING", None,
     "Zona UIO tomada del nombre de archivo que genero el propio sistema (OT-0966-T050-0-UIO.pdf)."),
]

ALIAS = [
    ("K166", "K167EC", "CORREO_DEL_DOCUMENTO (kfc167@kfc.com.ec) + Plaza San Francisco esta en el Centro Historico"),
    ("K166EC", "K167EC", "CORREO_DEL_DOCUMENTO (kfc167@kfc.com.ec) + Plaza San Francisco esta en el Centro Historico"),
    ("A018", "A014EC", "CORREO_DEL_DOCUMENTO (amca14@americandeli.com.ec) + unico American Deli de LARB"),
    ("A018EC", "A014EC", "CORREO_DEL_DOCUMENTO (amca14@americandeli.com.ec) + unico American Deli de LARB"),
]


def conectar():
    env = {}
    for line in (BASE / "config/.env").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"], autocommit=True)


def main():
    cnx = conectar()
    cur = cnx.cursor(dictionary=True)

    # 'OTRA' para locales fuera de las tres zonas contratadas. Se agrega al final del
    # enum para no alterar el orden ni el significado de los valores existentes.
    cur.execute("SHOW COLUMNS FROM locales LIKE 'zona'")
    tipo = cur.fetchone()["Type"]
    if "OTRA" not in tipo:
        cur.execute("ALTER TABLE locales MODIFY zona "
                    "ENUM('UIO','LARB','CNLJ','OTRA') NOT NULL "
                    "COMMENT 'OTRA = local atendido fuera de las tres zonas del contrato KFC'")
        print("Enum de zona ampliado con 'OTRA'")
    else:
        print("El enum de zona ya contempla 'OTRA'")

    print("\n=== ALTAS AL MAESTRO ===")
    for codigo, zona, cadena, nombre, correo, nota in ALTAS:
        cur.execute("SELECT local_codigo FROM locales WHERE local_codigo=%s", (codigo,))
        existia = cur.fetchone() is not None
        cur.execute(
            """INSERT INTO locales (local_codigo, zona, cadena, nombre, correo_local, activo)
               VALUES (%s,%s,%s,%s,%s,1)
               ON DUPLICATE KEY UPDATE zona=VALUES(zona), cadena=VALUES(cadena),
                                       nombre=VALUES(nombre), correo_local=VALUES(correo_local)""",
            (codigo, zona, cadena, nombre, correo))
        print(f"  {'actualizado' if existia else 'creado':11s} {codigo:8s} {zona:5s} {cadena:14s} {nombre}")

    print("\n=== ALIAS (codigos mal escritos) ===")
    for alias, canonico, regla in ALIAS:
        cur.execute(
            """INSERT INTO locales_alias (alias_texto, local_codigo, regla_aplicada, nivel_confianza)
               VALUES (%s,%s,%s,2)
               ON DUPLICATE KEY UPDATE local_codigo=VALUES(local_codigo),
                                       regla_aplicada=VALUES(regla_aplicada)""",
            (alias, canonico, regla[:60]))
        print(f"  {alias:8s} -> {canonico:8s}  ({regla[:56]})")

    cur.execute("SELECT COUNT(*) c FROM locales")
    total = cur.fetchone()["c"]
    cur.execute("SELECT zona, COUNT(*) c FROM locales GROUP BY zona")
    por_zona = {r["zona"]: r["c"] for r in cur.fetchall()}
    print(f"\nTotal de locales en la base: {total}  {por_zona}")

    # ---- archivo de propuesta para la administracion (ella decide cuando pegarlo) ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ALTAS PROPUESTAS"
    ws.append(["COD_TIENDA", "ZONA", "CADENA", "UBICACION / NOMBRE", "CORREO DEL LOCAL",
               "POR QUE SE PROPONE", "CONFIRMA ADMIN"])
    for codigo, zona, cadena, nombre, correo, nota in ALTAS:
        ws.append([codigo, zona, cadena, nombre, correo or "", nota, ""])

    ws2 = wb.create_sheet("CODIGOS MAL ESCRITOS")
    ws2.append(["CODIGO EN EL DOCUMENTO", "LOCAL REAL", "EVIDENCIA", "CONFIRMA ADMIN"])
    vistos = set()
    for alias, canonico, regla in ALIAS:
        if canonico in vistos:
            continue
        vistos.add(canonico)
        ws2.append([alias, canonico, regla, ""])

    for hoja in (ws, ws2):
        hoja.freeze_panes = "A2"
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", start_color="1F4E78")
        for col in hoja.columns:
            ancho = max((len(str(c.value)) for c in col if c.value), default=10)
            hoja.column_dimensions[col[0].column_letter].width = min(ancho + 2, 60)
    PROPUESTA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PROPUESTA)
    print(f"\nPropuesta para el Excel de la administracion: {PROPUESTA}")
    print("(el maestro en Drive no se toca: ella decide cuando incorporarlo)")
    cur.close()
    cnx.close()


if __name__ == "__main__":
    main()
