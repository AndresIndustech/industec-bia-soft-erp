import openpyxl
PATH = r"D:\RESPALDOS\_ORIGEN_DRIVE\KPI´S - INDUSTEC\KPI´S INDUSTEC.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True, read_only=True)
ws = wb["FILTRO"]
print("dim:", "max_row:", ws.max_row, "max_col:", ws.max_column)
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, max_col=ws.max_column)):
    vals = [f"{c.coordinate}={c.value}" for c in row]
    print(vals)
