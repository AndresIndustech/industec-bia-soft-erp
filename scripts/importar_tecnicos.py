"""T1.7 - Importa los 19 empleados desde BASE DE DATOS DE EMPLEADOS INDUSTEC.xlsx"""
from pathlib import Path
import openpyxl
import mysql.connector

ENV_PATH = Path(r"D:\INDUSTECH IA\agentes\config\.env")
env = {}
for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
    line = line.strip()
    if not line or line.startswith("#") or "=" not in line:
        continue
    k, v = line.split("=", 1)
    env[k.strip()] = v.strip()

XLSX_PATH = Path("D:/RESPALDOS/_ORIGEN_DRIVE/INDUSTEC/BASE DE DATOS DE EMPLEADOS INDUSTEC.xlsx")

ZONA_MAP = {
    "ZONA UIO": "UIO",
    "ZONA CUENCA LOJA": "CNLJ",
}


def mapear_zona(texto):
    if not texto:
        return None
    t = str(texto).strip().upper()
    if t in ZONA_MAP:
        return ZONA_MAP[t]
    if t.startswith("ZONA LARB"):
        return "LARB"
    return None


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    filas = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=8):
        n, nombres, apellidos, cedula, fecha_ingreso, tipo, afiliado, zona_txt = [c.value for c in row]
        if not nombres:
            continue
        cedula_str = str(cedula).strip() if cedula is not None else None
        fecha = fecha_ingreso.date() if hasattr(fecha_ingreso, "date") else None
        afiliado_bool = 1 if str(afiliado).strip().upper() == "SI" else 0
        zona = mapear_zona(zona_txt)
        filas.append((str(nombres).strip(), str(apellidos).strip(), cedula_str, fecha,
                       str(tipo).strip() if tipo else None, afiliado_bool, zona))

    print(f"Empleados leidos: {len(filas)} (esperado 19)")

    cnx = mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"],
    )
    cur = cnx.cursor()
    cur.execute("DELETE FROM tecnicos")
    cur.executemany(
        """INSERT INTO tecnicos (nombres, apellidos, cedula, fecha_ingreso, tipo_tecnico, afiliado, zona_asignada, activo)
           VALUES (%s,%s,%s,%s,%s,%s,%s,1)""",
        filas,
    )
    cnx.commit()
    cur.execute("SELECT COUNT(*) FROM tecnicos")
    total = cur.fetchone()[0]
    cur.execute("SELECT zona_asignada, COUNT(*) FROM tecnicos GROUP BY zona_asignada")
    por_zona = cur.fetchall()
    cur.close()
    cnx.close()
    print(f"Total en tecnicos: {total}")
    print("Por zona:", por_zona)


if __name__ == "__main__":
    main()
