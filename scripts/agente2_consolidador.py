"""
Agente 2 (Consolidador) - genera el plan de seguimiento de una zona,
replicando el formato real de PLAN SEGUIMIENTO OTS {ZONA} _ {MES}.xlsx.

Diseño verificado contra el archivo real de UIO de septiembre 2026 (61
filas): el eje de cada fila es el AVISO SAP, no la orden de INDUSTEC. Isabel
arranca de lo que llega de SAP y lo enriquece con la(s) orden(es) de
INDUSTEC vinculadas por ese aviso (0, 1 o 2 -- evaluacion y/o cierre).
Patrones de completitud reales observados:
  A) Aviso + orden de evaluacion (sin cierre aun)              -> ABIERTA
  B) Aviso sin ninguna orden de INDUSTEC vinculada             -> fila minima
  C) Aviso + evaluacion + cierre                               -> CERRADA
  D) Aviso + solo cierre (evaluacion no localizada)            -> caso raro

No se recrea el archivo desde cero: se abre la plantilla real como base
(preserva estilos, tabla, formato condicional, anchos) y solo se
reescriben los VALORES de las filas de datos.
"""
import re
import sys
from pathlib import Path
from datetime import date, datetime
import mysql.connector
import openpyxl
from openpyxl.utils import get_column_letter

ENV_PATH = Path(r"D:\INDUSTECH IA\agentes\config\.env")
env = {}
for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
    line = line.strip()
    if not line or line.startswith("#") or "=" not in line:
        continue
    k, v = line.split("=", 1)
    env[k.strip()] = v.strip()

# Plantilla real de UIO (preserva estilos/tabla/formato condicional exactos)
PLANTILLA = Path(r"D:\RESPALDOS\_ORIGEN_DRIVE\GESTION DE OTS INDUSTEC\2026\PLANES SEMANALES\PLAN DE TRABAJO _ ZONA UIO\PLAN SEGUIMIENTO OTS UIO _ SEPTIEMBRE.xlsx")
OUT_DIR = Path(r"D:\INDUSTECH IA\SALIDAS IA\MANTENIMIENTO\MANTENIMIENTOS CORRECTIVOS\ZONA UIO\PLANES SEMANALES")

# Columnas A..V, en el orden real verificado del archivo
COLS = ["#", "# OT", "TECNICO EVALUACION", "TECNICO CIERRE", "LOCAL",
        "FECHA DE INICIO", "EQUIPO", "MARCA", "TRABAJO REALIZADO EVALUACION",
        "REPUESTO", "ESTATUS SAP", "PRESUPUESTO", "TRABAJO REALIZADO CIERRE",
        "OBSERVACIONES", "ESTATUS DEL EQUIPO", "#OT INDUSTEC EVALUACION",
        "FECHA EVALUACION", "#OT INDUSTEC CIERRE", "FECHA CIERRE",
        "REQUERIMIENTO A TIEMPO", "CALIFICACION SATISFACCIÓN", "ESTADO"]

# Columnas que usan el literal "NINGUNO" cuando no hay dato (verificado en
# el archivo real: SOLO estas, K/L/N; D/M/R/S/T/U simplemente quedan
# vacias cuando la fase de cierre no existe).
COLS_NINGUNO = {"K", "L", "N"}


def conectar():
    cnx = mysql.connector.connect(
        host=env["DB_HOST"], port=int(env["DB_PORT"]), user=env["DB_USER"],
        password=env["DB_PASSWORD"], database=env["DB_NAME"],
    )
    cnx.autocommit = True
    return cnx


def id_industec_original(id_canonico, local_codigo):
    """Reconstruye el ID tal como aparece en el archivo original (sin el
    sufijo EC que T1.6 le agrego al local): 'OT-1627-G015EC-...' ->
    'OT-1627-G015-...'. Se usa asi para maxima fidelidad con lo que la
    administracion ya reconoce; el codigo CANONICO (con EC) es el que vive
    en la base de datos como clave primaria de la orden."""
    if not id_canonico or not local_codigo:
        return id_canonico
    local_sin_ec = re.sub(r"EC$", "", local_codigo)
    return id_canonico.replace(local_codigo, local_sin_ec, 1)


def calcular_semaforo(estado, fecha_inicio, hoy):
    """Las 3 reglas exactas del plan/archivo original: abierta y >=3 dias,
    =2 dias, <2 dias. Solo aplica si la orden sigue ABIERTA; una CERRADA no
    necesita semaforo de antiguedad (ya se resolvio)."""
    if estado != "ABIERTA" or not fecha_inicio:
        return None
    dias = (hoy - fecha_inicio).days
    if dias >= 3:
        return "ROJO"
    if dias == 2:
        return "AMBAR"
    return "VERDE"


def cargar_equipos(cnx, ids):
    if not ids:
        return {}
    cur = cnx.cursor(dictionary=True)
    formato = ",".join(["%s"] * len(ids))
    cur.execute(f"""SELECT id_industec, equipo, marca, estado_equipo, orden
                     FROM ot_equipos WHERE id_industec IN ({formato}) ORDER BY orden""", list(ids))
    resultado = {}
    for r in cur.fetchall():
        resultado.setdefault(r["id_industec"], []).append(r)
    cur.close()
    return resultado


def main():
    zona = sys.argv[1] if len(sys.argv) > 1 else "UIO"
    anio = int(sys.argv[2]) if len(sys.argv) > 2 else 2026
    mes = int(sys.argv[3]) if len(sys.argv) > 3 else 9
    hoy = date.today()

    mes_ini = date(anio, mes, 1)
    mes_fin = date(anio + (1 if mes == 12 else 0), 1 if mes == 12 else mes + 1, 1)

    cnx = conectar()
    cur = cnx.cursor(dictionary=True)
    cur.execute("""SELECT id_industec FROM ots WHERE zona=%s AND aviso IS NOT NULL AND en_cuarentena=0""", (zona,))
    ids_ots = [r["id_industec"] for r in cur.fetchall()]
    equipos_por_ot = cargar_equipos(cnx, ids_ots)
    cur.close()

    # El plan real (verificado sobre el archivo de UIO/septiembre) no es solo
    # "lo notificado este mes": arrastra el backlog de meses anteriores
    # mientras el CASO siga sin resolver. El criterio de "resuelto" correcto
    # es avisos_sap.estatus_general (columna ESTATUS A del export SAP: el
    # estado real segun el cliente -- CERRADO/TRATAMIENTO/ABIERTO), NO si
    # existe una segunda OT de INDUSTEC con estado_ot=CERRADA: se verifico
    # que basarse solo en la OT de INDUSTEC sobre-contaba el backlog varias
    # veces (474 filas vs 61 reales), porque un caso puede darse por resuelto
    # en SAP sin que exista una visita formal de "cierre" de INDUSTEC.
    cur = cnx.cursor(dictionary=True)
    cur.execute("""SELECT a.aviso, a.fecha_notificacion, a.descripcion, a.centro_coste,
                          a.estatus_aviso, a.estatus_general
                   FROM avisos_sap a
                   WHERE (a.centro_coste IN (SELECT local_codigo FROM locales WHERE zona=%s)
                          OR a.aviso IN (SELECT aviso FROM ots WHERE zona=%s AND aviso IS NOT NULL))
                     AND (a.fecha_notificacion >= %s AND a.fecha_notificacion < %s
                          OR a.estatus_general IN ('ABIERTO', 'TRATAMIENTO'))
                   ORDER BY a.fecha_notificacion""", (zona, zona, mes_ini, mes_fin))
    avisos = cur.fetchall()
    # OJO: se elimino "OR a.aviso IS NULL" -- esa condicion colaba TODO el
    # historico de ordenes sin cobertura de catalogo (2025 entero) como si
    # fuera backlog activo de septiembre 2026 (bug real: 394 filas de mas).
    # Sin cobertura de catalogo, el unico criterio disponible es la fecha.
    cur.execute("""SELECT ots.* FROM ots
                   LEFT JOIN avisos_sap a ON a.aviso = ots.aviso
                   WHERE ots.zona=%s AND ots.aviso IS NOT NULL AND ots.en_cuarentena=0
                   AND (ots.fecha_atencion >= %s AND ots.fecha_atencion < %s
                        OR a.estatus_general IN ('ABIERTO', 'TRATAMIENTO'))
                   ORDER BY ots.fecha_atencion""", (zona, mes_ini, mes_fin))
    ots_rows = cur.fetchall()
    cur.close()
    cnx.close()
    print(f"Filtro: notificado/atendido en {mes_ini} a {mes_fin}, o backlog sin cierre. "
          f"Avisos: {len(avisos)}  Ordenes: {len(ots_rows)}")

    for o in ots_rows:
        o["_equipos"] = equipos_por_ot.get(o["id_industec"], [])
    ots_por_aviso = {}
    for o in ots_rows:
        ots_por_aviso.setdefault(o["aviso"], []).append(o)

    filas = []
    avisos_cubiertos = set()
    for a in avisos:
        ordenes = ots_por_aviso.get(a["aviso"], [])
        avisos_cubiertos.add(a["aviso"])
        filas.append(_armar_fila(a, ordenes, hoy))
    for aviso, ordenes in ots_por_aviso.items():
        if aviso not in avisos_cubiertos:
            filas.append(_armar_fila(None, ordenes, hoy))

    print(f"Zona {zona}: {len(filas)} filas construidas (avisos SAP: {len(avisos)}, "
          f"avisos solo-en-ots: {len(ots_por_aviso) - len(avisos_cubiertos & set(ots_por_aviso))})")

    # --- Escribir sobre una COPIA de la plantilla real (nunca el original)
    wb = openpyxl.load_workbook(PLANTILLA)
    ws = wb["Hoja1"]
    fila_inicio = 2
    for i, fila in enumerate(filas):
        r = fila_inicio + i
        fila["#"] = i + 1
        for j, nombre_col in enumerate(COLS, start=1):
            col_letra = get_column_letter(j)
            valor = fila.get(nombre_col)
            if valor in (None, "") and col_letra in COLS_NINGUNO:
                valor = "NINGUNO"
            ws.cell(r, j).value = valor

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    nombre_salida = f"PLAN SEGUIMIENTO OTS {zona} _ SEPTIEMBRE (generado agente).xlsx"
    ruta_salida = OUT_DIR / nombre_salida
    wb.save(ruta_salida)
    print(f"Plan generado: {ruta_salida}")


def mayus(v):
    """Isabel escribe tecnico/marca/equipo en MAYUSCULAS en su plan (verificado
    contra el archivo real: 'NACIONAL', 'PRINCE CASTLE'...). Es la correccion
    de mayor impacto encontrada en la comparacion celda a celda (23 de las
    diferencias eran solo de capitalizacion)."""
    return v.upper() if isinstance(v, str) else v


def _armar_fila(aviso_row, ordenes, hoy):
    fila = {c: None for c in COLS}
    fila["# OT"] = aviso_row["aviso"] if aviso_row else (ordenes[0]["aviso"] if ordenes else None)

    ordenes_ordenadas = sorted(ordenes, key=lambda o: (o["fecha_atencion"] or date.min))
    evaluacion = next((o for o in ordenes_ordenadas if o["estado_ot"] != "CERRADA"), None)
    cierre = next((o for o in reversed(ordenes_ordenadas) if o["estado_ot"] == "CERRADA"), None)

    local_codigo = (evaluacion or cierre)["local_codigo"] if (evaluacion or cierre) else (aviso_row["centro_coste"] if aviso_row else None)
    fila["LOCAL"] = local_codigo

    if evaluacion:
        fila["TECNICO EVALUACION"] = mayus(evaluacion["tecnico_nombre"])
        fila["FECHA DE INICIO"] = evaluacion["fecha_atencion"]
        eqs = evaluacion.get("_equipos") or []
        if eqs:
            fila["EQUIPO"] = mayus(eqs[0].get("equipo"))
            fila["MARCA"] = mayus(eqs[0].get("marca"))
            fila["ESTATUS DEL EQUIPO"] = eqs[0].get("estado_equipo")
        fila["TRABAJO REALIZADO EVALUACION"] = evaluacion["actividades"]
        fila["REPUESTO"] = evaluacion["repuestos"]
        fila["#OT INDUSTEC EVALUACION"] = id_industec_original(evaluacion["id_industec"], evaluacion["local_codigo"])
        fila["FECHA EVALUACION"] = evaluacion["fecha_atencion"]
    elif aviso_row:
        fila["FECHA DE INICIO"] = aviso_row["fecha_notificacion"]
        fila["EQUIPO"] = mayus(aviso_row["descripcion"])

    if cierre:
        fila["TECNICO CIERRE"] = mayus(cierre["tecnico_nombre"])
        fila["TRABAJO REALIZADO CIERRE"] = cierre["actividades"] or cierre["observaciones"]
        fila["#OT INDUSTEC CIERRE"] = id_industec_original(cierre["id_industec"], cierre["local_codigo"])
        fila["FECHA CIERRE"] = cierre["fecha_atencion"]
        fila["REQUERIMIENTO A TIEMPO"] = mayus(cierre["atiempo"])
        fila["CALIFICACION SATISFACCIÓN"] = cierre["satisfaccion"]
        if not fila["ESTATUS DEL EQUIPO"]:
            eqs_c = cierre.get("_equipos") or []
            fila["ESTATUS DEL EQUIPO"] = eqs_c[0].get("estado_equipo") if eqs_c else None

    fila["ESTATUS SAP"] = aviso_row["estatus_aviso"] if aviso_row else None
    fila["OBSERVACIONES"] = (evaluacion["observaciones"] if evaluacion else None) or (cierre["observaciones"] if cierre else None)
    # Fuente principal del estado: el veredicto real de SAP (estatus_general),
    # no si INDUSTEC registro una segunda visita formal de "cierre". Solo si
    # el aviso no esta en el catalogo SAP (caso fuera de cobertura) se cae al
    # criterio secundario basado en la propia OT.
    estatus_sap = aviso_row["estatus_general"] if aviso_row else None
    if estatus_sap == "CERRADO":
        fila["ESTADO"] = "CERRADA"
    elif estatus_sap in ("ABIERTO", "TRATAMIENTO"):
        fila["ESTADO"] = "ABIERTA"
    else:
        fila["ESTADO"] = "CERRADA" if cierre else "ABIERTA"

    return fila


if __name__ == "__main__":
    main()
