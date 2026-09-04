"""Inspeccion completa del maestro de locales, sin asumir estructura de memoria."""
import openpyxl

PATH = r"D:\RESPALDOS\_ORIGEN_DRIVE\LOCALES_INDUSTEC_GENERAL\LOCALES_INDUSTEC_GENERAL.xlsx"

wb = openpyxl.load_workbook(PATH, data_only=True)
print("HOJAS:", wb.sheetnames)
print()

for name in ['UIO', 'CUENCA', 'LARB']:
    ws = wb[name]
    print(f"===== HOJA: {name}  dim={ws.dimensions}  filas={ws.max_row}  cols={ws.max_column} =====")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = str(v).replace("\n", " ")
            vals.append(f"{cell.coordinate}={s}")
        if vals:
            print("  " + " | ".join(vals))
    print()
